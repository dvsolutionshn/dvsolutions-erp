from datetime import date
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Empresa, Modulo, PlanComercial, PlanModulo, Usuario
from facturacion.models import CAI, Cliente, CompraInventario, Factura, LineaCompraInventario, LineaFactura, Producto, TipoImpuesto
from .models import AsientoContable, CuentaContable, CuentaFinanciera, LineaAsientoContable
from .reportes import (
    cuentas_efectivo_ids, datos_balance_comprobacion, datos_balance_general,
    datos_estado_resultados, datos_flujo_efectivo, datos_mayor,
)
from .services import registrar_asiento_factura_emitida
from .views import _fiscal_documento_lempiras


class EstadosFinancierosTests(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(nombre="Empresa A", slug="reportes-a", rtn="08011999111121", estado_licencia="activa")
        self.otra = Empresa.objects.create(nombre="Empresa B", slug="reportes-b", rtn="08011999111122", estado_licencia="activa")
        modulo, _ = Modulo.objects.get_or_create(codigo="contabilidad", defaults={"nombre": "Contabilidad", "es_comercial": True})
        plan = PlanComercial.objects.create(nombre="Plan reportes", codigo="plan-reportes", precio_mensual="99.00")
        PlanModulo.objects.create(plan=plan, modulo=modulo, activo=True)
        self.empresa.plan_comercial = plan
        self.empresa.save(update_fields=["plan_comercial"])
        self.usuario = Usuario.objects.create_user(username="reportes-contador", password="pass12345", empresa=self.empresa, is_superuser=True, is_staff=True)
        self.client.force_login(self.usuario)
        self.caja = self.cuenta("1101", "Caja", "activo")
        self.banco = self.cuenta("110201", "Banco principal", "activo")
        self.clientes = self.cuenta("111001", "Clientes", "activo")
        self.inventario = self.cuenta("112001", "Inventario", "activo")
        self.proveedores = self.cuenta("210101", "Proveedores", "pasivo")
        self.ventas = self.cuenta("410101", "Ventas", "ingreso")
        self.costos = self.cuenta("5101", "Costo de ventas", "costo")
        self.gastos = self.cuenta("610101", "Gasto administrativo", "gasto")
        self.capital = self.cuenta("3101", "Capital", "patrimonio")
        self.utilidades = self.cuenta("3103", "Utilidades acumuladas", "patrimonio")

    def cuenta(self, codigo, nombre, tipo, empresa=None, **kwargs):
        return CuentaContable.objects.create(empresa=empresa or self.empresa, codigo=codigo, nombre=nombre, tipo=tipo, **kwargs)

    def asiento(self, fecha, partidas, empresa=None, evento=None, estado="contabilizado", **kwargs):
        empresa = empresa or self.empresa
        asiento = AsientoContable.objects.create(empresa=empresa, fecha=fecha, descripcion="Operación de prueba", estado=estado, evento=evento, **kwargs)
        for cuenta, debe, haber in partidas:
            LineaAsientoContable.objects.create(asiento=asiento, cuenta=cuenta, debe=Decimal(str(debe)), haber=Decimal(str(haber)))
        return asiento

    def excel(self, nombre, filtros=None):
        response = self.client.get(reverse(nombre, args=[self.empresa.slug]), filtros or {})
        self.assertEqual(response.status_code, 200)
        return list(load_workbook(BytesIO(response.content), data_only=True).active.values)

    def test_balance_incluye_utilidad_y_perdida_no_traspasadas(self):
        self.asiento(date(2026, 1, 5), [(self.clientes, 500, 0), (self.ventas, 0, 500)])
        self.asiento(date(2026, 1, 5), [(self.costos, 200, 0), (self.inventario, 0, 200)])
        datos = datos_balance_general(self.empresa, "2026-01-31")
        self.assertEqual(datos["resumen"]["resultado_no_traspasado"], Decimal("300.00"))
        self.assertEqual(datos["resumen"]["total_patrimonio"], Decimal("300.00"))
        self.assertTrue(datos["resumen"]["cuadrado"])
        self.asiento(date(2026, 2, 1), [(self.gastos, 400, 0), (self.caja, 0, 400)])
        posterior = datos_balance_general(self.empresa, "2026-02-28")
        self.assertEqual(posterior["resumen"]["resultado_no_traspasado"], Decimal("-100.00"))
        self.assertTrue(posterior["resumen"]["cuadrado"])

    def test_cierre_anual_no_duplica_resultado_y_conserva_estado_resultados(self):
        self.asiento(date(2025, 12, 10), [(self.caja, 500, 0), (self.ventas, 0, 500)])
        self.asiento(date(2025, 12, 10), [(self.gastos, 200, 0), (self.caja, 0, 200)])
        self.asiento(date(2025, 12, 31), [(self.ventas, 500, 0), (self.gastos, 0, 200), (self.utilidades, 0, 300)], evento="cierre_anual")
        cerrado = datos_balance_general(self.empresa, "2025-12-31")["resumen"]
        self.assertEqual(cerrado["resultado_no_traspasado"], Decimal("0.00"))
        self.assertEqual(cerrado["patrimonio_contabilizado"], Decimal("300.00"))
        self.assertEqual(cerrado["total_patrimonio"], Decimal("300.00"))
        self.assertTrue(cerrado["cuadrado"])
        resultados = datos_estado_resultados(self.empresa, "2025-01-01", "2025-12-31")["resumen"]
        self.assertEqual(resultados["total_ingresos"], Decimal("500.00"))
        self.assertEqual(resultados["utilidad"], Decimal("300.00"))
        self.asiento(date(2026, 1, 5), [(self.caja, 50, 0), (self.ventas, 0, 50)])
        abierto = datos_balance_general(self.empresa, "2026-01-31")["resumen"]
        self.assertEqual(abierto["resultado_no_traspasado"], Decimal("50.00"))
        self.assertEqual(abierto["total_patrimonio"], Decimal("350.00"))

    def test_balance_pantalla_excel_y_resumen_comparten_resultado(self):
        self.asiento(date(2026, 1, 2), [(self.caja, 100, 0), (self.ventas, 0, 100)])
        self.asiento(date(2026, 2, 2), [(self.caja, 30, 0), (self.ventas, 0, 30)])
        response = self.client.get(reverse("balance_general", args=[self.empresa.slug]), {"fecha_corte": "2026-01-31"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["resumen"]["total_patrimonio"], Decimal("100.00"))
        self.assertContains(response, "Resultado acumulado no traspasado")
        filas = self.excel("exportar_balance_general_excel", {"fecha_corte": "2026-01-31"})
        self.assertEqual(next(fila[2] for fila in filas if fila[0] == "Total patrimonio"), 100)
        resumen = self.client.get(reverse("resumen_ejecutivo_contable", args=[self.empresa.slug]), {"fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"})
        self.assertEqual(resumen.context["resumen"]["total_patrimonio"], Decimal("100.00"))
        self.assertEqual(resumen.context["resumen"]["diferencia_balance"], Decimal("0.00"))

    def test_resultados_pantalla_excel_jerarquia_saldos_y_utilidad_sin_duplicar(self):
        principal = self.cuenta("4100", "Ingresos por ventas", "ingreso", acepta_movimientos=False)
        self.ventas.cuenta_padre = principal
        self.ventas.save(update_fields=["cuenta_padre"])
        otra_venta = self.cuenta("410102", "Ventas de servicios", "ingreso", cuenta_padre=principal)
        self.asiento(date(2026, 1, 2), [(self.caja, 150000, 0), (self.ventas, 0, 100000), (otra_venta, 0, 50000)])
        self.asiento(date(2026, 1, 3), [(self.costos, 20000, 0), (self.gastos, 10000, 0), (self.caja, 0, 30000)])
        filtros = {"fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"}
        respuesta = self.client.get(reverse("estado_resultados", args=[self.empresa.slug]), filtros)
        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, '<th class="ef-codigo-columna" scope="col">Código</th>', count=3, html=True)
        self.assertContains(respuesta, f'class="ef-fila ef-principal" data-cuenta-codigo="{principal.codigo}"')
        self.assertContains(respuesta, f'class="ef-fila" data-cuenta-codigo="{self.ventas.codigo}"')
        self.assertContains(respuesta, "padding-inline-start: 18px;")
        self.assertNotContains(respuesta, "<th>Debe</th>", html=True)
        self.assertNotContains(respuesta, "<th>Haber</th>", html=True)
        self.assertEqual(respuesta.context["resumen"]["utilidad"], Decimal("120000.00"))
        excel = self.client.get(reverse("exportar_estado_resultados_excel", args=[self.empresa.slug]), filtros)
        self.assertEqual(excel.status_code, 200)
        hoja = load_workbook(BytesIO(excel.content), data_only=True).active
        encabezado = next(fila[0].row for fila in hoja if [celda.value for celda in fila] == ["Código", "Cuenta", "Saldo"])
        self.assertEqual([celda.value for celda in hoja[encabezado]], ["Código", "Cuenta", "Saldo"])
        codigos = (principal.codigo, self.ventas.codigo, otra_venta.codigo)
        filas = [fila for fila in hoja.iter_rows(min_row=encabezado + 1) if fila[0].value in codigos]
        self.assertEqual([fila[0].value for fila in filas], list(codigos))
        self.assertEqual([fila[2].value for fila in filas], [150000, 100000, 50000])
        self.assertEqual([fila[1].alignment.indent for fila in filas], [0, 1, 1])
        self.assertTrue(all(celda.font.bold for celda in filas[0]))
        self.assertTrue(all(not celda.font.bold for fila in filas[1:] for celda in fila))
        valores = list(hoja.values)
        self.assertEqual(next(fila[2] for fila in valores if fila[0] == "Total ingresos"), 150000)
        self.assertEqual(next(fila[2] for fila in valores if fila[0] == "Utilidad neta"), 120000)
        self.assertEqual(datos_balance_general(self.empresa, "2026-01-31")["resumen"]["diferencia"], Decimal("0.00"))

    def test_mayor_pasivo_apertura_saldo_natural_y_excel(self):
        self.asiento(date(2026, 1, 5), [(self.inventario, 100, 0), (self.proveedores, 0, 100)])
        self.asiento(date(2026, 2, 5), [(self.proveedores, 25, 0), (self.caja, 0, 25)])
        filtros = {"cuenta": self.proveedores.pk, "fecha_inicio": "2026-02-01", "fecha_fin": "2026-02-28"}
        bloques = datos_mayor(self.empresa, filtros["fecha_inicio"], filtros["fecha_fin"], cuenta=self.proveedores)
        self.assertEqual(bloques[0]["saldo_inicial"], Decimal("100.00"))
        self.assertEqual(bloques[0]["movimientos"][0]["saldo"], Decimal("75.00"))
        response = self.client.get(reverse("mayor_cuenta", args=[self.empresa.slug]), filtros)
        self.assertEqual(response.context["resumen"]["saldo"], Decimal("75.00"))
        filas = self.excel("exportar_mayor_cuenta_excel", filtros)
        self.assertEqual(next(fila[7] for fila in filas if fila[3] == "Saldo inicial anterior al periodo"), 100)
        self.assertEqual(next(fila[7] for fila in filas if fila[3] == "Operación de prueba"), 75)

    def test_mayor_conserva_apertura_individual_y_general_oculta_cuentas_sin_movimiento(self):
        self.asiento(date(2026, 1, 5), [(self.inventario, 100, 0), (self.proveedores, 0, 100)])
        self.inventario.activa = False
        self.inventario.save(update_fields=["activa"])
        bloques = datos_mayor(self.empresa, "2026-02-01", "2026-02-28")
        bloque = next(b for b in bloques if b["cuenta"].pk == self.inventario.pk)
        self.assertEqual(bloque["saldo"], Decimal("100.00"))
        self.assertEqual(bloque["movimientos"], [])
        response = self.client.get(reverse("libro_mayor_general", args=[self.empresa.slug]), {"fecha_inicio": "2026-02-01", "fecha_fin": "2026-02-28"})
        self.assertEqual(response.context["cuentas_mayor"], [])
        individual = self.client.get(reverse("mayor_cuenta", args=[self.empresa.slug]), {
            "cuenta": self.inventario.pk, "fecha_inicio": "2026-02-01", "fecha_fin": "2026-02-28",
        })
        self.assertContains(individual, "Inventario")
        self.assertEqual(individual.context["resumen"]["saldo_inicial"], Decimal("100.00"))

    def test_comprobacion_incluye_apertura_y_movimiento_conciliados(self):
        self.asiento(date(2026, 1, 5), [(self.caja, 100, 0), (self.capital, 0, 100)])
        self.asiento(date(2026, 2, 5), [(self.gastos, 25, 0), (self.caja, 0, 25)])
        datos = datos_balance_comprobacion(self.empresa, "2026-02-01", "2026-02-28")
        fila = next(fila for fila in datos["filas"] if fila["cuenta"].pk == self.caja.pk)
        self.assertEqual(fila["apertura_deudora"], Decimal("100.00"))
        self.assertEqual(fila["haber"], Decimal("25.00"))
        self.assertEqual(fila["saldo_deudor"], Decimal("75.00"))
        self.assertTrue(datos["resumen"]["cuadrado"])
        self.assertEqual(datos["resumen"]["apertura_deudora_total"], datos["resumen"]["apertura_acreedora_total"])
        filas = self.excel("exportar_balance_comprobacion_excel", {"fecha_inicio": "2026-02-01", "fecha_fin": "2026-02-28"})
        caja_excel = next(fila for fila in filas if fila[0] == "1101 - Caja")
        self.assertEqual(caja_excel[2:], (100, 0, 0, 25, 75, 0))

    def test_resultados_cuentas_inactivas_anulaciones_y_borradores(self):
        self.asiento(date(2026, 1, 5), [(self.caja, 100, 0), (self.ventas, 0, 100)])
        self.ventas.activa = False
        self.ventas.save(update_fields=["activa"])
        self.asiento(date(2026, 1, 5), [(self.caja, 999, 0), (self.ventas, 0, 999)], estado="borrador")
        self.asiento(date(2026, 1, 5), [(self.caja, 999, 0), (self.ventas, 0, 999)], estado="anulado")
        self.assertEqual(datos_estado_resultados(self.empresa)["resumen"]["utilidad"], Decimal("100.00"))

    def test_empresas_aisladas_con_codigos_repetidos(self):
        caja_b = self.cuenta("1101", "Caja B", "activo", self.otra)
        ventas_b = self.cuenta("410101", "Ventas B", "ingreso", self.otra)
        self.asiento(date(2026, 1, 5), [(self.caja, 100, 0), (self.ventas, 0, 100)])
        self.asiento(date(2026, 1, 5), [(caja_b, 900, 0), (ventas_b, 0, 900)], empresa=self.otra)
        self.assertEqual(datos_balance_general(self.empresa)["resumen"]["total_activos"], Decimal("100.00"))
        self.assertEqual(datos_estado_resultados(self.empresa)["resumen"]["utilidad"], Decimal("100.00"))
        self.assertEqual(datos_flujo_efectivo(self.empresa)["resumen"]["total_neto"], Decimal("100.00"))
        self.assertNotIn(caja_b.pk, cuentas_efectivo_ids(self.empresa))
        response = self.client.get(reverse("mayor_cuenta", args=[self.empresa.slug]), {"cuenta": caja_b.pk})
        self.assertEqual(response.status_code, 404)

    def test_flujo_clientes_y_proveedores_son_operacion(self):
        self.asiento(date(2026, 1, 5), [(self.caja, 115, 0), (self.clientes, 0, 115)])
        self.asiento(date(2026, 1, 6), [(self.proveedores, 40, 0), (self.banco, 0, 40)])
        datos = datos_flujo_efectivo(self.empresa)
        self.assertEqual(datos["flujo"]["operacion"]["entradas"], Decimal("115.00"))
        self.assertEqual(datos["flujo"]["operacion"]["salidas"], Decimal("40.00"))
        self.assertEqual(datos["resumen"]["financiamiento"], Decimal("0.00"))
        self.assertEqual(datos["resumen"]["saldo_final"], Decimal("75.00"))
        self.assertTrue(datos["resumen"]["conciliado"])

    def test_flujo_deuda_capital_activos_fijos_y_transferencia(self):
        prestamo = self.cuenta("2201", "Préstamo bancario", "pasivo")
        equipo = self.cuenta("120104", "Equipo de cómputo", "activo")
        self.asiento(date(2026, 1, 1), [(self.caja, 1000, 0), (self.capital, 0, 1000)])
        self.asiento(date(2026, 2, 1), [(self.banco, 500, 0), (prestamo, 0, 500)])
        self.asiento(date(2026, 2, 2), [(equipo, 200, 0), (self.banco, 0, 200)])
        self.asiento(date(2026, 2, 3), [(self.caja, 100, 0), (self.banco, 0, 100)])
        datos = datos_flujo_efectivo(self.empresa, "2026-02-01", "2026-02-28")
        self.assertEqual(datos["resumen"]["saldo_inicial"], Decimal("1000.00"))
        self.assertEqual(datos["resumen"]["financiamiento"], Decimal("500.00"))
        self.assertEqual(datos["resumen"]["inversion"], Decimal("-200.00"))
        self.assertEqual(datos["resumen"]["total_neto"], Decimal("300.00"))
        self.assertEqual(datos["resumen"]["movimientos"], 2)
        self.assertEqual(datos["resumen"]["saldo_final"], Decimal("1300.00"))
        self.assertTrue(datos["resumen"]["conciliado"])
        filas = self.excel("exportar_estado_flujo_efectivo_excel", {"fecha_inicio": "2026-02-01", "fecha_fin": "2026-02-28"})
        self.assertEqual(next(fila[7] for fila in filas if fila[0] == "Efectivo al inicio"), 1000)
        self.assertEqual(next(fila[7] for fila in filas if fila[0] == "Efectivo al cierre"), 1300)

    def test_flujo_separa_principal_e_intereses_en_un_asiento(self):
        prestamo = self.cuenta("210401", "Préstamo a corto plazo", "pasivo")
        self.asiento(date(2026, 2, 2), [(prestamo, 100, 0), (self.gastos, 10, 0), (self.banco, 0, 110)])
        datos = datos_flujo_efectivo(self.empresa)
        self.assertEqual(datos["resumen"]["operacion"], Decimal("-10.00"))
        self.assertEqual(datos["resumen"]["financiamiento"], Decimal("-100.00"))
        self.assertTrue(datos["resumen"]["conciliado"])

    def test_flujo_categoria_explicita_y_heredada_en_cuentas_personalizadas(self):
        padre = self.cuenta("9900", "Obligaciones del proyecto", "pasivo", acepta_movimientos=False, actividad_flujo="financiamiento")
        deuda = self.cuenta("990001", "Obligación especial", "pasivo", cuenta_padre=padre)
        aporte = self.cuenta("990002", "Obligación operativa", "pasivo", cuenta_padre=padre, actividad_flujo="operacion")
        self.asiento(date(2026, 1, 1), [(self.banco, 100, 0), (deuda, 0, 100)])
        self.asiento(date(2026, 1, 2), [(self.banco, 25, 0), (aporte, 0, 25)])
        datos = datos_flujo_efectivo(self.empresa)
        self.assertEqual(datos["resumen"]["financiamiento"], Decimal("100.00"))
        self.assertEqual(datos["resumen"]["operacion"], Decimal("25.00"))
        self.assertTrue(datos["resumen"]["conciliado"])

    def test_banco_con_enlace_legacy_tarjeta_conserva_efectivo_y_avisa(self):
        CuentaFinanciera.objects.create(empresa=self.empresa, nombre="Banco correcto", tipo="banco", cuenta_contable=self.banco)
        enlace_antiguo = CuentaFinanciera.objects.create(empresa=self.empresa, nombre="Enlace antiguo", tipo="banco", cuenta_contable=self.banco)
        CuentaFinanciera.objects.filter(pk=enlace_antiguo.pk).update(tipo="tarjeta_credito")
        self.asiento(date(2026, 1, 1), [(self.banco, 100, 0), (self.capital, 0, 100)])
        datos = datos_flujo_efectivo(self.empresa)
        self.assertIn(self.banco.pk, cuentas_efectivo_ids(self.empresa))
        self.assertEqual(datos["resumen"]["saldo_final"], Decimal("100.00"))
        self.assertEqual(datos["resumen"]["cuentas_financieras_incompatibles"], 1)
        self.assertTrue(datos["resumen"]["conciliado"])

    def test_tarjeta_credito_no_es_efectivo_ni_compra_a_credito_es_flujo(self):
        tarjeta = self.cuenta("210402", "Tarjeta de crédito", "pasivo")
        CuentaFinanciera.objects.create(empresa=self.empresa, nombre="Tarjeta", tipo="tarjeta_credito", cuenta_contable=tarjeta)
        self.asiento(date(2026, 2, 1), [(self.inventario, 115, 0), (tarjeta, 0, 115)])
        self.assertNotIn(tarjeta.pk, cuentas_efectivo_ids(self.empresa))
        datos = datos_flujo_efectivo(self.empresa)
        self.assertEqual(datos["resumen"]["total_neto"], Decimal("0.00"))
        self.assertEqual(datos["resumen"]["saldo_final"], Decimal("0.00"))
        self.assertEqual(datos["resumen"]["movimientos"], 0)
        self.asiento(date(2026, 2, 2), [(tarjeta, 100, 0), (self.caja, 0, 100)])
        self.assertEqual(datos_flujo_efectivo(self.empresa)["resumen"]["financiamiento"], Decimal("-100.00"))

    def test_cuenta_caja_inactiva_conserva_saldo_flujo(self):
        self.asiento(date(2026, 1, 1), [(self.caja, 100, 0), (self.capital, 0, 100)])
        self.caja.activa = False
        self.caja.save(update_fields=["activa"])
        self.assertEqual(datos_flujo_efectivo(self.empresa, "2026-02-01", "2026-02-28")["resumen"]["saldo_inicial"], Decimal("100.00"))

    def test_diferencias_historicas_no_se_ocultan(self):
        self.asiento(date(2026, 2, 1), [(self.caja, 50, 0), (self.ventas, 0, 40)])
        balance = datos_balance_general(self.empresa)["resumen"]
        flujo = datos_flujo_efectivo(self.empresa)["resumen"]
        self.assertEqual(balance["diferencia"], Decimal("10.00"))
        self.assertFalse(balance["cuadrado"])
        self.assertEqual(flujo["diferencia"], Decimal("10.00"))
        self.assertFalse(flujo["conciliado"])

    def test_comprobacion_tipo_indica_vista_parcial(self):
        self.asiento(date(2026, 1, 1), [(self.caja, 100, 0), (self.capital, 0, 100)])
        response = self.client.get(reverse("balance_comprobacion", args=[self.empresa.slug]), {"tipo": "activo"})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Vista parcial")
        self.assertContains(response, "Selecciona todos los tipos")

    def test_ganancia_cambiaria_alimenta_resultados_balance_y_flujo(self):
        ganancia = self.cuenta("420101", "Ganancia por diferencia cambiaria", "ingreso")
        self.asiento(date(2026, 1, 1), [(self.clientes, 2500, 0), (self.ventas, 0, 2500)], moneda_origen="USD", tipo_cambio=25)
        self.asiento(date(2026, 2, 1), [(self.banco, 2600, 0), (self.clientes, 0, 2500), (ganancia, 0, 100)], moneda_origen="USD", tipo_cambio=26)
        resultados = datos_estado_resultados(self.empresa, "2026-02-01", "2026-02-28")["resumen"]
        self.assertEqual(resultados["utilidad"], Decimal("100.00"))
        self.assertTrue(datos_balance_general(self.empresa)["resumen"]["cuadrado"])
        flujo = datos_flujo_efectivo(self.empresa, "2026-02-01", "2026-02-28")["resumen"]
        self.assertEqual(flujo["total_neto"], Decimal("2600.00"))
        self.assertTrue(flujo["conciliado"])

    def test_fiscal_usd_convierte_componentes_coherentes_a_hnl(self):
        documento = SimpleNamespace(moneda="USD", tipo_cambio=Decimal("25.00"), total=Decimal("115.00"), resumen_fiscal=lambda: {
            "base_15": Decimal("100.00"), "isv_15": Decimal("15.00"), "base_18": 0, "isv_18": 0,
            "base_exento": 0, "base_exonerado": 0, "descuento_total": Decimal("5.00"),
        })
        fiscal = _fiscal_documento_lempiras(documento)
        self.assertEqual(fiscal["subtotal"], Decimal("2500.00"))
        self.assertEqual(fiscal["impuesto"], Decimal("375.00"))
        self.assertEqual(fiscal["total"], Decimal("2875.00"))
        self.assertEqual(fiscal["descuento_total"], Decimal("125.00"))

    def test_fiscal_redondeo_componentes_suma_total(self):
        documento = SimpleNamespace(moneda="USD", tipo_cambio=Decimal("25.1234"), total=Decimal("0.03"), resumen_fiscal=lambda: {
            "base_15": Decimal("0.01"), "isv_15": Decimal("0.01"), "base_exento": Decimal("0.01"),
        })
        fiscal = _fiscal_documento_lempiras(documento)
        self.assertEqual(fiscal["subtotal"] + fiscal["impuesto"], fiscal["total"])
        self.assertEqual(fiscal["total"], Decimal("0.75"))

    def test_factura_usd_enlaza_fiscal_asiento_resultados_y_balance(self):
        cliente = Cliente.objects.create(empresa=self.empresa, nombre="Cliente en dólares")
        impuesto = TipoImpuesto.objects.create(nombre="ISV reportes 15%", porcentaje=Decimal("15.00"))
        producto = Producto.objects.create(empresa=self.empresa, nombre="Servicio reportes", tipo_item="servicio", unidad_medida="servicio", precio=100, controla_inventario=False, impuesto_predeterminado=impuesto)
        CAI.objects.create(empresa=self.empresa, numero_cai="CAI-REPORTES-USD", uso_documento="factura", establecimiento="001", punto_emision="001", tipo_documento="01", rango_inicial=1, rango_final=99999999, correlativo_actual=0, fecha_limite=date(2026, 12, 31), activo=True)
        factura = Factura.objects.create(empresa=self.empresa, cliente=cliente, vendedor=self.usuario, fecha_emision=date(2026, 1, 10), fecha_vencimiento=date(2026, 1, 31), moneda="USD", tipo_cambio=25, estado="borrador")
        LineaFactura.objects.create(factura=factura, producto=producto, cantidad=1, precio_unitario=100, impuesto=impuesto)
        factura.calcular_totales()
        factura.estado = "emitida"
        factura.save()
        asiento = registrar_asiento_factura_emitida(factura)
        self.assertEqual(asiento.total_debe, Decimal("2875.00"))
        self.assertEqual(datos_estado_resultados(self.empresa)["resumen"]["utilidad"], Decimal("2500.00"))
        self.assertTrue(datos_balance_general(self.empresa)["resumen"]["cuadrado"])
        response = self.client.get(reverse("reporte_impuestos", args=[self.empresa.slug]), {"fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["ventas"]["base_15"], Decimal("2500.00"))
        self.assertEqual(response.context["ventas"]["impuesto"], Decimal("375.00"))
        self.assertEqual(response.context["ventas"]["total"], Decimal("2875.00"))
        filas = self.excel("exportar_reporte_impuestos_excel", {"fecha_inicio": "2026-01-01", "fecha_fin": "2026-01-31"})
        venta_excel = next(fila for fila in filas if fila[0] == "Facturas emitidas" and fila[1] == 1)
        self.assertEqual(venta_excel[2], 2500)
        self.assertEqual(venta_excel[8], 375)
        self.assertEqual(venta_excel[9], 2875)

    def test_panel_compras_excluye_borradores_y_otra_empresa(self):
        producto = Producto.objects.create(empresa=self.empresa, nombre="Servicio compra", tipo_item="servicio", unidad_medida="servicio", precio=100, controla_inventario=False)
        aplicada = CompraInventario.objects.create(empresa=self.empresa, proveedor_nombre="Proveedor aplicado", fecha_documento=date(2026, 1, 10), estado="aplicada")
        borrador = CompraInventario.objects.create(empresa=self.empresa, proveedor_nombre="Proveedor borrador", fecha_documento=date(2026, 1, 10), estado="borrador")
        LineaCompraInventario.objects.create(compra=aplicada, producto=producto, cantidad=1, costo_unitario=50)
        LineaCompraInventario.objects.create(compra=borrador, producto=producto, cantidad=1, costo_unitario=999)
        CompraInventario.objects.create(empresa=self.otra, proveedor_nombre="Proveedor B", fecha_documento=date(2026, 1, 10), estado="aplicada")
        response = self.client.get(reverse("dashboard_bi_financiero", args=[self.empresa.slug]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual([compra.pk for compra in response.context["compras"]], [aplicada.pk])
        self.assertEqual(response.context["compras_borrador"], 1)
        self.assertContains(response, "Panel financiero")
        self.assertNotContains(response, "Proveedor borrador")
        self.assertNotContains(response, "Proveedor B")

    def test_alerta_conversion_historica_solo_empresa_y_corte(self):
        self.asiento(date(2026, 1, 1), [(self.caja, 100, 0), (self.ventas, 0, 100)], moneda_origen="USD", tipo_cambio=25, conversion_verificada=False)
        self.asiento(date(2026, 2, 1), [(self.caja, 100, 0), (self.ventas, 0, 100)], moneda_origen="USD", tipo_cambio=25, conversion_verificada=False)
        response = self.client.get(reverse("balance_general", args=[self.empresa.slug]), {"fecha_corte": "2026-01-31"})
        self.assertEqual(response.context["conversiones_pendientes"], 1)
        self.assertContains(response, "revisión de conversión")
        self.assertEqual(response.context["resumen"]["total_activos"], Decimal("100.00"))
