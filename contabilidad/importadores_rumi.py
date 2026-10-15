from decimal import Decimal, InvalidOperation
from collections import deque
from copy import deepcopy
import hashlib
import json
import re
import unicodedata
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import CuentaContable, bloquear_empresas_contables
from .alcance_catalogo import es_empresa_rumi


def _exigir_empresa_rumi(empresa):
    if not es_empresa_rumi(empresa):
        raise ValidationError("Este importador de catálogos está disponible únicamente para la empresa RUMI.")


TIPOS_CUENTA = {clave: clave for clave, _ in CuentaContable.TIPO_CHOICES}
TIPOS_CUENTA.update({
    "activos": "activo",
    "assets": "activo",
    "pasivos": "pasivo",
    "liability": "pasivo",
    "liabilities": "pasivo",
    "patrimonio": "patrimonio",
    "capital": "patrimonio",
    "ingresos": "ingreso",
    "income": "ingreso",
    "revenue": "ingreso",
    "costo": "costo",
    "costos": "costo",
    "cost": "costo",
    "costs": "costo",
    "gastos": "gasto",
    "expenses": "gasto",
    "orden": "orden",
    "cuentas_de_orden": "orden",
    "control_extracontable": "orden",
})

COLUMNAS = {
    "codigo": {"codigo", "code", "cuenta", "numero", "numero_cuenta"},
    "nombre": {"nombre", "name", "descripcion_cuenta", "cuenta_nombre"},
    "tipo": {"tipo", "type", "clase", "categoria"},
    "codigo_padre": {"codigo_padre", "padre", "cuenta_padre", "parent", "parent_code"},
    "acepta_movimientos": {"acepta_movimientos", "movimiento", "acepta_movimiento", "posting", "movimientos"},
    "activa": {"activa", "activo", "estado", "active"},
    "descripcion": {"descripcion", "description", "detalle"},
}


def _normalizar_texto(valor):
    if valor is None:
        return ""
    return str(valor).strip()


def _normalizar_encabezado(valor):
    texto = unicodedata.normalize("NFKD", _normalizar_texto(valor))
    texto = "".join(caracter for caracter in texto if not unicodedata.combining(caracter))
    return texto.lower().replace(" ", "_").replace("-", "_")


def _normalizar_booleano(valor, por_defecto=True):
    texto = _normalizar_texto(valor).lower()
    if not texto:
        return por_defecto
    if texto in {"1", "si", "sí", "s", "true", "verdadero", "activo", "activa", "yes", "y"}:
        return True
    if texto in {"0", "no", "n", "false", "falso", "inactivo", "inactiva"}:
        return False
    return por_defecto


def _normalizar_tipo(valor):
    tipo = _normalizar_encabezado(valor)
    return TIPOS_CUENTA.get(tipo)


def _inferir_tipo_desde_codigo(codigo):
    primer_digito = _normalizar_texto(codigo)[:1]
    return {
        "1": "activo",
        "2": "pasivo",
        "3": "patrimonio",
        "4": "ingreso",
        "5": "costo",
        "6": "gasto",
    }.get(primer_digito)


def _mapear_columnas(encabezados):
    indices = {}
    encabezados_normalizados = [_normalizar_encabezado(valor) for valor in encabezados]
    for campo, alias in COLUMNAS.items():
        for index, encabezado in enumerate(encabezados_normalizados):
            if encabezado in alias:
                indices[campo] = index
                break
    return indices


def _valor(fila, indices, campo):
    index = indices.get(campo)
    if index is None:
        return ""
    return _normalizar_texto(fila[index])


def _es_codigo_contable(valor):
    texto = _normalizar_texto(valor)
    return bool(texto) and texto.replace(".", "").replace("-", "").isdigit()


def _detectar_formato_jerarquico(sheet):
    coincidencias = 0
    for row in range(1, min(sheet.max_row, 40) + 1):
        codigo = sheet.cell(row, 8).value
        nombre = sheet.cell(row, 9).value
        if _es_codigo_contable(codigo) and _normalizar_texto(nombre):
            coincidencias += 1
    return coincidencias >= 3


RUMI_COLUMNAS = {
    "codigo": "codigo_rumi",
    "nombre": "cuenta_rumi",
    "nivel": "nivel",
    "codigo_padre": "cuenta_padre",
    "acepta_movimientos": "permite_asiento",
    "naturaleza": "naturaleza",
    "activacion": "activacion",
    "funcion": "tipo_funcional_erp",
    "conciliar": "conciliar",
    "referencia": "referencia_cnbs",
    "alcance": "alcance_de_la_referencia",
    "auxiliar": "auxiliar_minimo",
    "fuente": "fuente_de_referencia_(url)",
    "autoria": "autoria_del_nombre",
}
RUMI_TIPOS_RAIZ = {
    "activo": "activo", "activos": "activo",
    "pasivo": "pasivo", "pasivos": "pasivo",
    "patrimonio": "patrimonio",
    "ingreso": "ingreso", "ingresos": "ingreso",
    "gasto": "gasto", "gastos": "gasto",
    "activos_y_pasivos_contingentes": "orden",
    "cuentas_de_orden": "orden",
}
CAMPOS_CATALOGO = {
    "codigo", "nombre", "tipo", "codigo_padre", "descripcion",
    "acepta_movimientos", "activa", "hoja", "fila",
}
CAMPOS_ACTUALIZABLES = ("nombre", "tipo", "descripcion", "acepta_movimientos", "activa")


def _ubicacion(item):
    hoja = item.get("hoja", "")
    fila = item.get("fila", "")
    origen = f"Hoja «{hoja}», fila {fila}" if hoja else f"Fila {fila}"
    codigo = item.get("codigo", "")
    return f"{origen}, cuenta {codigo}" if codigo else origen


def _texto_celda(celda, *, clave=False):
    if clave and celda.data_type == "f":
        raise ValidationError("contiene una fórmula; sustituya la fórmula por su valor antes de importar.")
    return _normalizar_texto(celda.value)


def _codigo_celda(celda):
    texto = _texto_celda(celda, clave=True)
    valor = celda.value
    if isinstance(valor, bool):
        raise ValidationError("el código debe ser texto o un número entero.")
    if isinstance(valor, (int, float, Decimal)):
        try:
            numero = Decimal(str(valor))
            if not numero.is_finite() or numero != numero.to_integral_value():
                raise ValidationError("el código numérico debe ser entero; use texto para códigos con puntos o guiones.")
            entero = str(int(numero))
            formato = celda.number_format or ""
            if re.fullmatch(r"0+", formato) and numero >= 0:
                return entero.zfill(len(formato))
            return entero
        except (InvalidOperation, ValueError, OverflowError) as exc:
            raise ValidationError("el código numérico no es válido.") from exc
    return texto


def _booleano_celda(celda):
    texto = _normalizar_encabezado(_texto_celda(celda, clave=True))
    if texto in {"1", "si", "s", "true", "verdadero", "activo", "activa", "yes", "y"}:
        return True
    if texto in {"0", "no", "n", "false", "falso", "inactivo", "inactiva"}:
        return False
    raise ValidationError("use sí/no, verdadero/falso o 1/0.")


def _detectar_catalogo(workbook, hoja=None):
    visibles = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
    if hoja:
        visibles = [sheet for sheet in visibles if sheet.title == hoja]
        if not visibles:
            raise ValidationError(f"No existe una hoja visible llamada «{hoja}» en este archivo.")
    candidatos = []
    for sheet in visibles:
        encontrado = None
        for fila in range(1, min(sheet.max_row, 40) + 1):
            celdas = list(sheet[fila])
            encabezados = [_normalizar_encabezado(cell.value) for cell in celdas]
            if {"codigo_rumi", "cuenta_rumi"}.issubset(encabezados):
                indices = {
                    campo: encabezados.index(nombre)
                    for campo, nombre in RUMI_COLUMNAS.items() if nombre in encabezados
                }
                encontrado = (sheet, fila, "rumi", indices)
                break
            indices = _mapear_columnas([cell.value for cell in celdas])
            if {"codigo", "nombre"}.issubset(indices):
                encontrado = (sheet, fila, "estandar", indices)
                break
        if encontrado:
            candidatos.append(encontrado)
        elif _detectar_formato_jerarquico(sheet):
            candidatos.append((sheet, 0, "jerarquico", {}))
    if not candidatos:
        raise ValidationError(
            "No se encontró un catálogo válido en las primeras 40 filas de las hojas visibles. "
            "Use código, nombre y tipo; el formato RUMI; o el formato jerárquico con código en H y nombre en I."
        )
    if len(candidatos) > 1:
        nombres = ", ".join(f"«{item[0].title}»" for item in candidatos)
        raise ValidationError(f"El archivo contiene varios catálogos ({nombres}). Indique la hoja exacta que desea importar.")
    return candidatos[0]


def _extraer_catalogo_jerarquico(sheet):
    datos, errores = [], []
    codigos_en_archivo = {}
    ultimo_codigo_por_nivel = {}
    for row in range(1, sheet.max_row + 1):
        origen = {"hoja": sheet.title, "fila": row}
        try:
            codigo = _codigo_celda(sheet.cell(row, 8))
            nombre = _texto_celda(sheet.cell(row, 9), clave=True)
        except ValidationError as exc:
            errores.extend(f"{_ubicacion(origen)}: {mensaje}" for mensaje in exc.messages)
            continue
        if not _es_codigo_contable(codigo) or not nombre:
            continue
        origen["codigo"] = codigo
        tipo = _inferir_tipo_desde_codigo(codigo)
        if not tipo:
            errores.append(f"{_ubicacion(origen)}: no se pudo inferir el tipo contable.")
            continue
        if codigo in codigos_en_archivo:
            errores.append(
                f"{_ubicacion(origen)}: codigo duplicado {codigo}. "
                f"Ya existe en fila {codigos_en_archivo[codigo]}."
            )
            continue
        nivel = sum(1 for col in range(1, 8) if _normalizar_texto(sheet.cell(row, col).value)) or 1
        codigo_padre = ultimo_codigo_por_nivel.get(nivel - 1, "")
        if nivel > 1 and not codigo_padre:
            errores.append(f"{_ubicacion(origen)}: falta la cuenta padre del nivel {nivel - 1}.")
        codigos_en_archivo[codigo] = row
        ultimo_codigo_por_nivel[nivel] = codigo
        ultimo_codigo_por_nivel = {n: c for n, c in ultimo_codigo_por_nivel.items() if n <= nivel}
        datos.append({
            **origen, "nombre": nombre, "tipo": tipo, "codigo_padre": codigo_padre,
            "descripcion": f"Importado desde formato jerárquico. Nivel {nivel}.",
            "acepta_movimientos": True, "activa": True,
        })
    padres = {item["codigo_padre"] for item in datos if item["codigo_padre"]}
    for item in datos:
        if item["codigo"] in padres:
            item["acepta_movimientos"] = False
    return datos, errores, {"leidas": len(datos), "omitidas": 0, "controles": 0, "condicionadas": 0}


def _extraer_catalogo_estandar(sheet, encabezado, indices):
    faltantes = sorted({"codigo", "nombre", "tipo"} - set(indices))
    if faltantes:
        raise ValidationError(f"Hoja «{sheet.title}», fila {encabezado}: faltan columnas obligatorias: {', '.join(faltantes)}.")
    datos, errores = [], []
    leidas = 0
    for fila in range(encabezado + 1, sheet.max_row + 1):
        if not any(sheet.cell(fila, indice + 1).value not in (None, "") for indice in indices.values()):
            continue
        leidas += 1
        item = {"hoja": sheet.title, "fila": fila}
        try:
            for campo, indice in indices.items():
                celda = sheet.cell(fila, indice + 1)
                if campo in ("codigo", "codigo_padre"):
                    item[campo] = _codigo_celda(celda)
                elif campo in ("acepta_movimientos", "activa"):
                    # Las celdas vacías conservan el valor existente, igual que una columna ausente.
                    if celda.value not in (None, ""):
                        item[campo] = _booleano_celda(celda)
                elif campo == "tipo":
                    item[campo] = _normalizar_tipo(_texto_celda(celda, clave=True)) or ""
                else:
                    item[campo] = _texto_celda(celda, clave=campo == "nombre")
            datos.append(item)
        except ValidationError as exc:
            errores.extend(f"{_ubicacion(item)}: {campo}: {mensaje}" for mensaje in exc.messages)
    return datos, errores, {"leidas": leidas, "omitidas": 0, "controles": 0, "condicionadas": 0}


def _extraer_catalogo_rumi(sheet, encabezado, indices, existentes, incluir_controles):
    faltantes = sorted(set(RUMI_COLUMNAS) - set(indices))
    if faltantes:
        nombres = [RUMI_COLUMNAS[campo] for campo in faltantes]
        raise ValidationError(f"Hoja «{sheet.title}», fila {encabezado}: faltan columnas RUMI: {', '.join(nombres)}.")
    filas, errores = [], []
    for fila in range(encabezado + 1, sheet.max_row + 1):
        if not any(sheet.cell(fila, indice + 1).value not in (None, "") for indice in indices.values()):
            continue
        item = {"hoja": sheet.title, "fila": fila}
        try:
            item["codigo"] = _codigo_celda(sheet.cell(fila, indices["codigo"] + 1))
            item["nombre"] = _texto_celda(sheet.cell(fila, indices["nombre"] + 1), clave=True)
            item["codigo_padre"] = _codigo_celda(sheet.cell(fila, indices["codigo_padre"] + 1))
            nivel = _texto_celda(sheet.cell(fila, indices["nivel"] + 1), clave=True)
            nivel_numerico = {"clase": 1, "grupo": 2, "subgrupo": 3, "detalle": 4}.get(
                _normalizar_encabezado(nivel)
            )
            if nivel_numerico is None and nivel.isdigit():
                nivel_numerico = int(nivel)
            if not nivel_numerico or nivel_numerico < 1:
                raise ValidationError("el nivel debe indicar Clase, Grupo, Subgrupo, Detalle o un número entero positivo.")
            item["nivel"] = nivel_numerico
            item["acepta_movimientos"] = _booleano_celda(sheet.cell(fila, indices["acepta_movimientos"] + 1))
            activacion = _normalizar_encabezado(_texto_celda(sheet.cell(fila, indices["activacion"] + 1), clave=True))
            if activacion not in {"estructura", "base_propuesta", "condicionada", "control_extracontable", "no_cargar"}:
                raise ValidationError("la activación debe indicar Estructura, Base propuesta, Condicionada, Control extracontable o No cargar.")
            item["activacion"] = activacion
            item["activa"] = activacion in {"estructura", "base_propuesta"}
            lineas = ["Importado del catálogo RUMI. Las referencias se conservan como información del archivo."]
            for campo, indice in indices.items():
                valor = _texto_celda(sheet.cell(fila, indice + 1))
                if valor:
                    etiqueta = _normalizar_texto(sheet.cell(encabezado, indice + 1).value)
                    lineas.append(f"{etiqueta}: {valor}")
            item["descripcion"] = "\n".join(lineas)
            filas.append(item)
        except ValidationError as exc:
            errores.extend(f"{_ubicacion(item)}: {mensaje}" for mensaje in exc.messages)
    por_codigo = {}
    for item in filas:
        if item["codigo"] in por_codigo:
            errores.append(f"{_ubicacion(item)}: código duplicado; ya aparece en fila {por_codigo[item['codigo']]['fila']}.")
        por_codigo[item["codigo"]] = item
    tipos = {}
    for item in filas:
        if not item["codigo_padre"]:
            tipo = RUMI_TIPOS_RAIZ.get(_normalizar_encabezado(item["nombre"]))
            if tipo:
                tipos[item["codigo"]] = tipo
            elif item["activacion"] != "no_cargar":
                errores.append(f"{_ubicacion(item)}: la clase raíz «{item['nombre']}» no tiene una clasificación reconocida.")
    for item in filas:
        if item["activacion"] == "no_cargar":
            continue
        pendientes, visitados = [], set()
        actual = item
        while actual["codigo"] not in tipos:
            codigo = actual["codigo"]
            if codigo in visitados:
                errores.append(f"{_ubicacion(item)}: la jerarquía tiene un ciclo.")
                break
            visitados.add(codigo)
            pendientes.append(codigo)
            padre = actual["codigo_padre"]
            if padre in tipos:
                tipo = tipos[padre]
                tipos.update({pendiente: tipo for pendiente in pendientes})
                break
            if padre in por_codigo:
                actual = por_codigo[padre]
            elif padre in existentes:
                tipos.update({pendiente: existentes[padre].tipo for pendiente in pendientes})
                break
            else:
                errores.append(f"{_ubicacion(item)}: la cuenta padre {padre or '(sin padre)'} no tiene una clase reconocida en esta empresa.")
                break
        if actual["codigo"] in tipos:
            tipos.update({pendiente: tipos[actual["codigo"]] for pendiente in pendientes})
    datos, omitidas, controles, condicionadas = [], 0, 0, 0
    for item in filas:
        tipo = tipos.get(item["codigo"])
        if tipo == "orden":
            controles += 1
        if item["activacion"] == "condicionada":
            condicionadas += 1
        if item["activacion"] == "no_cargar" or (tipo == "orden" and not incluir_controles):
            omitidas += 1
            continue
        if not tipo:
            continue
        if tipo == "orden" or item["activacion"] == "control_extracontable":
            tipo = "orden"
            item["activa"] = False
            item["acepta_movimientos"] = False
        if tipo not in {"activo", "pasivo", "patrimonio", "ingreso", "costo", "gasto", "orden"}:
            errores.append(f"{_ubicacion(item)}: el tipo heredado no es válido.")
            continue
        padre = por_codigo.get(item["codigo_padre"])
        if padre and item["nivel"] != padre["nivel"] + 1:
            errores.append(f"{_ubicacion(item)}: el nivel no coincide con el de su cuenta padre.")
        if not item["codigo_padre"] and item["nivel"] != 1:
            errores.append(f"{_ubicacion(item)}: una cuenta raíz debe tener nivel 1.")
        datos.append({campo: valor for campo, valor in item.items() if campo in CAMPOS_CATALOGO} | {"tipo": tipo})
    padres = {item["codigo_padre"] for item in datos if item["codigo_padre"]}
    for item in datos:
        if item["codigo"] in padres:
            item["acepta_movimientos"] = False
    return datos, errores, {
        "leidas": len(filas), "omitidas": omitidas, "controles": controles, "condicionadas": condicionadas,
    }


def _referencias_cuentas(cuentas):
    """Incluye partidas, configuraciones, auxiliares y respaldos; excluye la relación padre."""
    ids = [cuenta.pk for cuenta in cuentas.values()]
    usadas, referencias = set(), {}
    if not ids:
        return usadas, referencias
    for relacion in CuentaContable._meta.related_objects:
        campo = relacion.field
        if relacion.related_model is CuentaContable or campo.many_to_many:
            continue
        nombre = f"{relacion.related_model._meta.label_lower}.{campo.name}"
        relacionados = list(relacion.related_model.objects.filter(
            **{f"{campo.attname}__in": ids}
        ).values_list(campo.attname, flat=True).order_by(campo.attname))
        if relacionados:
            referencias[nombre] = relacionados
            usadas.update(relacionados)
    # Un padre determina la estructura y parte de la clasificación de sus auxiliares.
    # Conserve también los ancestros de las cuentas que ya están conectadas al sistema.
    por_id = {cuenta.pk: cuenta for cuenta in cuentas.values()}
    pendientes = list(usadas)
    while pendientes:
        cuenta = por_id.get(pendientes.pop())
        padre_id = cuenta.cuenta_padre_id if cuenta else None
        if padre_id in por_id and padre_id not in usadas:
            usadas.add(padre_id)
            pendientes.append(padre_id)
    return usadas, referencias


def _firma_cuentas(cuentas, referencias):
    valores = []
    for cuenta in sorted(cuentas.values(), key=lambda item: item.pk):
        valores.append({
            "id": cuenta.pk, "codigo": cuenta.codigo, "nombre": cuenta.nombre,
            "tipo": cuenta.tipo, "padre": cuenta.cuenta_padre_id, "descripcion": cuenta.descripcion,
            "activa": cuenta.activa, "acepta_movimientos": cuenta.acepta_movimientos,
            "actividad_flujo": cuenta.actividad_flujo,
        })
    contenido = json.dumps({"cuentas": valores, "referencias": referencias}, sort_keys=True, ensure_ascii=False)
    return hashlib.sha256(contenido.encode("utf-8")).hexdigest()


def _validar_propuesta(empresa, datos, existentes, usadas, actualizar_existentes, *, validar_unicidad=True):
    _exigir_empresa_rumi(empresa)
    if not isinstance(datos, list) or not datos:
        raise ValidationError("No se encontraron cuentas para importar con las opciones elegidas.")
    errores, preparadas, por_codigo = [], [], {}
    tipos_validos = {tipo for tipo, _nombre in CuentaContable.TIPO_CHOICES}
    codigos_equivalentes = {codigo.casefold(): codigo for codigo in existentes}
    for original in datos:
        if not isinstance(original, dict) or set(original) - CAMPOS_CATALOGO:
            errores.append("El catálogo contiene campos no admitidos; vuelva a analizar el archivo.")
            continue
        item = dict(original)
        for campo in ("codigo", "nombre", "tipo"):
            if not isinstance(item.get(campo), str) or not item[campo].strip():
                errores.append(f"{_ubicacion(item)}: falta {campo} o su valor no es válido.")
            else:
                item[campo] = item[campo].strip()
        codigo = item.get("codigo")
        if not isinstance(codigo, str) or not codigo:
            continue
        if len(codigo) > 30:
            errores.append(f"{_ubicacion(item)}: el código supera los 30 caracteres.")
        if isinstance(item.get("nombre"), str) and len(item["nombre"]) > 200:
            errores.append(f"{_ubicacion(item)}: el nombre supera los 200 caracteres.")
        if item.get("tipo") not in tipos_validos:
            errores.append(f"{_ubicacion(item)}: tipo inválido. Use activo, pasivo, patrimonio, ingreso, costo, gasto u orden.")
        equivalente = codigos_equivalentes.get(codigo.casefold())
        if equivalente and equivalente != codigo:
            errores.append(f"{_ubicacion(item)}: el código ya existe como {equivalente}; conserve sus mayúsculas y minúsculas.")
        if codigo.casefold() in {clave.casefold() for clave in por_codigo}:
            errores.append(f"{_ubicacion(item)}: código duplicado en el archivo.")
        por_codigo[codigo] = item
        cuenta = existentes.get(codigo)
        if cuenta and cuenta.cuenta_padre_id and cuenta.cuenta_padre.empresa_id != empresa.pk:
            errores.append(f"{_ubicacion(item)}: la relación padre existente pertenece a otra empresa; revise primero esa inconsistencia.")
        for campo in ("activa", "acepta_movimientos"):
            if campo in item and not isinstance(item[campo], bool):
                errores.append(f"{_ubicacion(item)}: {campo} debe indicar sí o no.")
        if "descripcion" in item and item["descripcion"] is not None and not isinstance(item["descripcion"], str):
            errores.append(f"{_ubicacion(item)}: la descripción debe ser texto.")
        if "codigo_padre" in item and not isinstance(item["codigo_padre"], str):
            errores.append(f"{_ubicacion(item)}: el código padre debe ser texto.")
        if "codigo_padre" not in item:
            padre = existentes.get(cuenta.cuenta_padre.codigo) if cuenta and cuenta.cuenta_padre_id else None
            item["codigo_padre"] = padre.codigo if padre else ""
        else:
            item["codigo_padre"] = item["codigo_padre"].strip() if isinstance(item["codigo_padre"], str) else ""
        if "descripcion" not in item:
            item["descripcion"] = cuenta.descripcion if cuenta else ""
        for campo in ("activa", "acepta_movimientos"):
            if campo not in item:
                item[campo] = getattr(cuenta, campo) if cuenta else item.get("tipo") != "orden"
        if item.get("tipo") == "orden" and (item["activa"] or item["acepta_movimientos"]):
            errores.append(f"{_ubicacion(item)}: las cuentas de control extracontable deben permanecer inactivas y sin asientos.")
        if cuenta and not actualizar_existentes:
            padre_actual = cuenta.cuenta_padre.codigo if cuenta.cuenta_padre_id else ""
            if item["codigo_padre"] != padre_actual or any(
                item[campo] != getattr(cuenta, campo) for campo in CAMPOS_ACTUALIZABLES
            ):
                errores.append(f"{_ubicacion(item)}: ya existe con datos diferentes y la actualización está desactivada.")
        if cuenta and cuenta.pk in usadas:
            padre_actual = cuenta.cuenta_padre.codigo if cuenta.cuenta_padre_id else ""
            cambios = [
                campo for campo in ("tipo", "activa", "acepta_movimientos")
                if item.get(campo) != getattr(cuenta, campo)
            ]
            if item["codigo_padre"] != padre_actual:
                cambios.append("cuenta padre")
            if cambios:
                errores.append(
                    f"{_ubicacion(item)}: tiene movimientos o enlaces operativos. "
                    f"No se puede cambiar {', '.join(cambios)}; conserve la cuenta o cree una cuenta nueva."
                )
        preparadas.append(item)
    if errores:
        raise ValidationError(errores)
    padres = {
        codigo: cuenta.cuenta_padre.codigo if cuenta.cuenta_padre_id else ""
        for codigo, cuenta in existentes.items()
    }
    tipos = {codigo: cuenta.tipo for codigo, cuenta in existentes.items()}
    padres.update({item["codigo"]: item["codigo_padre"] for item in preparadas})
    tipos.update({item["codigo"]: item["tipo"] for item in preparadas})
    for item in preparadas:
        padre = item["codigo_padre"]
        if padre and padre not in padres:
            errores.append(f"{_ubicacion(item)}: la cuenta padre {padre} no existe en el archivo ni en esta empresa.")
        elif padre and tipos[padre] != item["tipo"]:
            errores.append(f"{_ubicacion(item)}: el tipo debe coincidir con su cuenta padre {padre}.")
        visitados, actual = set(), item["codigo"]
        while actual:
            if actual in visitados:
                errores.append(f"{_ubicacion(item)}: la jerarquía tiene un ciclo o la cuenta es padre de sí misma.")
                break
            visitados.add(actual)
            actual = padres.get(actual, "")
        for codigo_hijo, codigo_padre in padres.items():
            if codigo_padre == item["codigo"] and tipos[codigo_hijo] != item["tipo"]:
                errores.append(f"{_ubicacion(item)}: la subcuenta {codigo_hijo} tiene un tipo diferente; conserve una jerarquía coherente.")
    if errores:
        raise ValidationError(errores)
    # Todos los objetos se validan antes de iniciar la primera escritura.
    objetos = {codigo: deepcopy(cuenta) for codigo, cuenta in existentes.items()}
    for item in preparadas:
        codigo = item["codigo"]
        cuenta = objetos.get(codigo) or CuentaContable(empresa=empresa, codigo=codigo)
        for campo in CAMPOS_ACTUALIZABLES:
            setattr(cuenta, campo, item[campo])
        cuenta.cuenta_padre = None
        objetos[codigo] = cuenta
    for item in preparadas:
        cuenta = objetos[item["codigo"]]
        # Las relaciones se enlazan tras validar el grafo; todavía no hay padres nuevos guardados.
        cuenta.cuenta_padre = objetos.get(item["codigo_padre"])
        try:
            cuenta.full_clean(exclude=["cuenta_padre"], validate_unique=validar_unicidad)
        except ValidationError as exc:
            errores.extend(f"{_ubicacion(item)}: {mensaje}" for mensaje in exc.messages)
    if errores:
        raise ValidationError(errores)
    conteos = {"creadas": 0, "actualizadas": 0, "sin_cambios": 0}
    for item in preparadas:
        cuenta = existentes.get(item["codigo"])
        if not cuenta:
            conteos["creadas"] += 1
        else:
            padre_actual = cuenta.cuenta_padre.codigo if cuenta.cuenta_padre_id else ""
            cambio = item["codigo_padre"] != padre_actual or any(
                item[campo] != getattr(cuenta, campo) for campo in CAMPOS_ACTUALIZABLES
            )
            conteos["actualizadas" if cambio else "sin_cambios"] += 1
    return preparadas, conteos


def analizar_catalogo_cuentas_desde_excel(
    empresa, archivo, *, actualizar_existentes=True, incluir_controles=True, hoja=None, prefijo_codigos="",
):
    _exigir_empresa_rumi(empresa)
    try:
        if hasattr(archivo, "seek"):
            archivo.seek(0)
        workbook = load_workbook(archivo, data_only=False)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise ValidationError("No se pudo leer el archivo. Suba un Excel válido en formato .xlsx o .xlsm.") from exc
    try:
        sheet, encabezado, formato, indices = _detectar_catalogo(workbook, hoja)
        existentes = {
            cuenta.codigo: cuenta for cuenta in CuentaContable.objects.filter(empresa=empresa).select_related("cuenta_padre")
        }
        if formato == "rumi":
            datos, errores, estadisticas = _extraer_catalogo_rumi(
                sheet, encabezado, indices, existentes, incluir_controles,
            )
        elif formato == "jerarquico":
            datos, errores, estadisticas = _extraer_catalogo_jerarquico(sheet)
        else:
            datos, errores, estadisticas = _extraer_catalogo_estandar(sheet, encabezado, indices)
        if errores:
            raise ValidationError(errores)
        if not isinstance(prefijo_codigos, str):
            raise ValidationError("El prefijo de códigos debe ser texto.")
        prefijo_codigos = prefijo_codigos.strip()
        if prefijo_codigos:
            codigos_originales = {item["codigo"] for item in datos}
            for item in datos:
                original = item["codigo"]
                item["codigo"] = prefijo_codigos + original
                if item.get("codigo_padre") in codigos_originales:
                    item["codigo_padre"] = prefijo_codigos + item["codigo_padre"]
                if formato != "rumi":
                    descripcion = item.get("descripcion") or ""
                    item["descripcion"] = f"Código original del archivo: {original}\n{descripcion}".strip()
        usadas, referencias = _referencias_cuentas(existentes)
        # La vista previa conserva los campos opcionales ausentes; aplicación los resuelve otra vez.
        _preparadas, conteos = _validar_propuesta(empresa, datos, existentes, usadas, actualizar_existentes)
        return {
            "datos": datos, "hoja": sheet.title, "fila_encabezado": encabezado, "formato": formato,
            **estadisticas, **conteos, "firma_existentes": _firma_cuentas(existentes, referencias),
        }
    finally:
        workbook.close()


def aplicar_catalogo_cuentas(empresa, datos, *, actualizar_existentes=True, firma_existentes=None):
    _exigir_empresa_rumi(empresa)
    with transaction.atomic():
        bloquear_empresas_contables([empresa.pk])
        # El bloqueo no usa una unión externa con la cuenta padre: PostgreSQL no permite
        # bloquear el lado nullable de esa unión. Todas las cuentas de la empresa se bloquean aquí.
        list(CuentaContable.objects.select_for_update().filter(empresa=empresa).order_by("pk"))
        existentes = {
            cuenta.codigo: cuenta for cuenta in CuentaContable.objects
            .filter(empresa=empresa).select_related("cuenta_padre").order_by("pk")
        }
        usadas, referencias = _referencias_cuentas(existentes)
        if firma_existentes is not None and firma_existentes != _firma_cuentas(existentes, referencias):
            raise ValidationError(
                "El catálogo o sus enlaces cambiaron después de la vista previa. Vuelva a analizar el archivo antes de confirmar."
            )
        preparadas, conteos = _validar_propuesta(empresa, datos, existentes, usadas, actualizar_existentes)
        por_codigo = {item["codigo"]: item for item in preparadas}
        hijos = {}
        pendientes = {}
        for item in preparadas:
            padre = item["codigo_padre"]
            pendientes[item["codigo"]] = int(padre in por_codigo)
            hijos.setdefault(padre, []).append(item["codigo"])
        disponibles = deque(codigo for codigo, cantidad in pendientes.items() if not cantidad)
        while disponibles:
            codigo = disponibles.popleft()
            item = por_codigo[codigo]
            original = existentes.get(codigo)
            cuenta = original or CuentaContable(empresa=empresa, codigo=codigo)
            cambios = any(item[campo] != getattr(cuenta, campo) for campo in CAMPOS_ACTUALIZABLES)
            padre = existentes.get(item["codigo_padre"])
            cambios = cambios or cuenta.cuenta_padre_id != (padre.pk if padre else None)
            if not original or cambios:
                for campo in CAMPOS_ACTUALIZABLES:
                    setattr(cuenta, campo, item[campo])
                cuenta.cuenta_padre = padre
                cuenta.full_clean()
                cuenta.save()
            existentes[codigo] = cuenta
            for hijo in hijos.get(codigo, []):
                pendientes[hijo] -= 1
                if not pendientes[hijo]:
                    disponibles.append(hijo)
        return {**conteos, "total": len(preparadas)}


def importar_catalogo_cuentas_desde_excel(
    empresa, archivo, *, actualizar_existentes=True, incluir_controles=True, hoja=None, prefijo_codigos="",
):
    _exigir_empresa_rumi(empresa)
    analisis = analizar_catalogo_cuentas_desde_excel(
        empresa, archivo, actualizar_existentes=actualizar_existentes,
        incluir_controles=incluir_controles, hoja=hoja, prefijo_codigos=prefijo_codigos,
    )
    resultado = aplicar_catalogo_cuentas(
        empresa, analisis["datos"], actualizar_existentes=actualizar_existentes,
        firma_existentes=analisis["firma_existentes"],
    )
    return {
        **resultado, "hoja": analisis["hoja"], "formato": analisis["formato"],
        "leidas": analisis["leidas"], "omitidas": analisis["omitidas"],
        "controles": analisis["controles"], "condicionadas": analisis["condicionadas"],
    }
