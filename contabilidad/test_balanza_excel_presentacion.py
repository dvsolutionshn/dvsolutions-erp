"""El encabezado y el formato de la Balanza conservan los importes contables."""

from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Empresa, Modulo, PlanComercial, PlanModulo, RolSistema, Usuario
from .models import AsientoContable, CuentaContable, LineaAsientoContable
from .reportes import datos_balance_comprobacion


class BalanzaExcelPresentacionTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        modulo, _ = Modulo.objects.get_or_create(
            codigo="contabilidad", defaults={"nombre": "Contabilidad", "es_comercial": True},
        )
        plan = PlanComercial.objects.create(
            nombre="Plan balanza", codigo="plan-balanza-presentacion", precio_mensual="99.00",
        )
        PlanModulo.objects.create(plan=plan, modulo=modulo, activo=True)
        cls.empresa = Empresa.objects.create(
            nombre="Empresa de la balanza", slug="balanza-presentacion", rtn="08011999410201",
            estado_licencia="activa", plan_comercial=plan,
        )
        cls.otra = Empresa.objects.create(
            nombre="Empresa privada de la balanza", slug="balanza-privada", rtn="08011999410202",
            estado_licencia="activa", plan_comercial=plan,
        )
        rol = RolSistema.objects.create(
            nombre="Consulta de balanza", codigo="consulta-balanza",
            puede_contabilidad=True, puede_reportes_contables=True,
        )
        cls.usuario = Usuario.objects.create_user(
            username="consulta-balanza", empresa=cls.empresa, rol_sistema=rol,
        )
        cls.caja = cls.cuenta("1101", "Caja", "activo")
        cls.banco = cls.cuenta("1102", "Banco", "activo")
        cls.capital = cls.cuenta("3101", "Capital", "patrimonio")
        cls.ventas = cls.cuenta("4101", "Ventas", "ingreso")
        cls.gastos = cls.cuenta("6101", "Gastos", "gasto")
        cls.sin_movimiento = cls.cuenta("1199", "Cuenta sin movimiento", "activo")
        cls.asiento(date(2025, 12, 31), [(cls.caja, "1000000.50", 0), (cls.capital, 0, "1000000.50")])
        cls.asiento(date(2026, 9, 5), [(cls.caja, "1250.10", 0), (cls.ventas, 0, "1250.10")])
        cls.asiento(date(2026, 9, 6), [(cls.banco, "1250000.50", 0), (cls.capital, 0, "1250000.50")])
        cls.asiento(date(2026, 9, 20), [(cls.gastos, "2.35", 0), (cls.caja, 0, "2.35")])
        cls.asiento(date(2026, 10, 1), [(cls.banco, "999999.99", 0), (cls.ventas, 0, "999999.99")])
        for estado in ("borrador", "anulado"):
            cls.asiento(date(2026, 9, 21), [(cls.banco, "9999.99", 0), (cls.ventas, 0, "9999.99")], estado=estado)
        privada = cls.cuenta("1101", "Cuenta confidencial", "activo", empresa=cls.otra)
        venta_privada = cls.cuenta("4101", "Venta confidencial", "ingreso", empresa=cls.otra)
        cls.asiento(date(2026, 9, 2), [(privada, "9000000.99", 0), (venta_privada, 0, "9000000.99")], empresa=cls.otra)

    @classmethod
    def cuenta(cls, codigo, nombre, tipo, empresa=None):
        return CuentaContable.objects.create(empresa=empresa or cls.empresa, codigo=codigo, nombre=nombre, tipo=tipo)

    @classmethod
    def asiento(cls, fecha, partidas, *, empresa=None, estado="contabilizado"):
        registro = AsientoContable.objects.create(
            empresa=empresa or cls.empresa, fecha=fecha, estado=estado, descripcion="Prueba de presentación de balanza",
        )
        for cuenta, debe, haber in partidas:
            LineaAsientoContable.objects.create(
                asiento=registro, cuenta=cuenta, debe=Decimal(str(debe)), haber=Decimal(str(haber)),
            )

    def setUp(self):
        self.client.force_login(self.usuario)
        self.filtros = {"fecha_inicio": "2026-09-01", "fecha_fin": "2026-09-30"}

    def excel(self, filtros=None):
        respuesta = self.client.get(
            reverse("exportar_balance_comprobacion_excel", args=[self.empresa.slug]),
            self.filtros if filtros is None else filtros,
        )
        self.assertEqual(respuesta.status_code, 200)
        return load_workbook(BytesIO(respuesta.content)).active

    def test_titulo_empresa_moneda_y_periodo_aparecen_antes_de_la_tabla(self):
        for filtros in (self.filtros, {}):
            with self.subTest(filtros=filtros):
                hoja = self.excel(filtros)
                self.assertEqual(hoja["A1"].value, "Balance de Comprobación · HNL")
                self.assertTrue(hoja["A1"].font.bold)
                encabezado = next(celda.row for celda in hoja["A"] if celda.value == "Cuenta")
                esperados = {"Empresa": self.empresa.nombre, "Moneda funcional": "Lempiras (HNL)"}
                if filtros:
                    esperados.update(Desde=filtros["fecha_inicio"], Hasta=filtros["fecha_fin"])
                for etiqueta, valor in esperados.items():
                    filas = [fila for fila in hoja if fila[0].value == etiqueta]
                    self.assertEqual(len(filas), 1)
                    self.assertLess(filas[0][0].row, encabezado)
                    self.assertEqual(filas[0][1].value, valor)
                    self.assertTrue(filas[0][0].font.bold)
                self.assertTrue(hoja.cell(encabezado, 1).font.bold)
                self.assertEqual(hoja.freeze_panes, f"A{encabezado + 1}")
                self.assertEqual(hoja.cell(hoja.max_row, 1).value, "Totales")
                self.assertTrue(all(celda.font.bold for celda in hoja[hoja.max_row]))

    def test_negrita_solo_en_principales_reales_sin_agregar_filas_ni_cambiar_importes(self):
        raiz = self.cuenta("RAIZ", "Agrupadora sin movimientos", "activo")
        principal = CuentaContable.objects.create(
            empresa=self.empresa, codigo="GRUPO", nombre="Principal con movimientos", tipo="activo", cuenta_padre=raiz,
        )
        intermedia = CuentaContable.objects.create(
            empresa=self.empresa, codigo="Z-AUX", nombre="Subcuenta agrupadora", tipo="activo", cuenta_padre=principal,
        )
        CuentaContable.objects.filter(pk=self.caja.pk).update(cuenta_padre=principal)
        CuentaContable.objects.filter(pk=self.banco.pk).update(cuenta_padre=intermedia)
        self.asiento(date(2026, 9, 22), [(principal, "7.05", 0), (self.capital, 0, "7.05")])
        self.asiento(date(2026, 9, 23), [(intermedia, "11.15", 0), (self.capital, 0, "11.15")])
        datos = datos_balance_comprobacion(self.empresa, **self.filtros)
        hoja = self.excel()
        encabezado = next(celda.row for celda in hoja["A"] if celda.value == "Cuenta")
        filas_excel = list(hoja.iter_rows(min_row=encabezado + 1, max_row=hoja.max_row - 1))
        self.assertEqual(
            [fila[0].value for fila in filas_excel],
            [f"{fila['cuenta'].codigo} - {fila['cuenta'].nombre}" for fila in datos["filas"]],
        )
        self.assertNotIn("RAIZ - Agrupadora sin movimientos", {fila[0].value for fila in filas_excel})
        for registro, fila in zip(filas_excel, datos["filas"]):
            agrupadora = fila["cuenta"].pk in {principal.pk, intermedia.pk}
            self.assertTrue(all(bool(celda.font.bold) == agrupadora for celda in registro))
            for celda, clave in zip(registro[2:8], ("apertura_deudora", "apertura_acreedora", "debe", "haber", "saldo_deudor", "saldo_acreedor")):
                self.assertEqual(Decimal(str(celda.value)), fila[clave])
                self.assertEqual(celda.data_type, "n")
                self.assertEqual(celda.number_format, '#,##0.00;[Red]-#,##0.00')

    def test_cuenta_sin_hijos_propios_no_se_marca_por_codigo_ni_relacion_ajena(self):
        hoja_raiz = self.cuenta("1", "Hoja con código de principal", "activo")
        self.asiento(date(2026, 9, 24), [(hoja_raiz, "5.10", 0), (self.capital, 0, "5.10")])
        # Un vínculo histórico ajeno no convierte la cuenta local en agrupadora.
        CuentaContable.objects.create(
            empresa=self.otra, codigo="HIJA-AJENA", nombre="Hija de otra empresa", tipo="activo", cuenta_padre=self.caja,
        )
        hoja = self.excel()
        for etiqueta in ("1 - Hoja con código de principal", "1101 - Caja", "1102 - Banco", "3101 - Capital"):
            registro = next(fila for fila in hoja if fila[0].value == etiqueta)
            self.assertTrue(all(not celda.font.bold for celda in registro))

    def test_detalles_y_totales_siguen_siendo_numericos_con_miles_y_dos_decimales(self):
        datos = datos_balance_comprobacion(self.empresa, **self.filtros)
        hoja = self.excel()
        por_etiqueta = {fila[0].value: fila for fila in hoja}
        for fila in datos["filas"]:
            celdas = por_etiqueta[f"{fila['cuenta'].codigo} - {fila['cuenta'].nombre}"]
            valores = [fila[clave] for clave in ("apertura_deudora", "apertura_acreedora", "debe", "haber", "saldo_deudor", "saldo_acreedor")]
            for celda, esperado in zip(celdas[2:8], valores):
                self.assertEqual(celda.data_type, "n")
                self.assertEqual(Decimal(str(celda.value)), esperado)
                self.assertEqual(celda.number_format, '#,##0.00;[Red]-#,##0.00')
        totales = por_etiqueta["Totales"]
        claves = ("apertura_deudora_total", "apertura_acreedora_total", "debe_total", "haber_total", "saldo_deudor_total", "saldo_acreedor_total")
        for celda, clave in zip(totales[2:8], claves):
            self.assertEqual(celda.data_type, "n")
            self.assertEqual(Decimal(str(celda.value)), datos["resumen"][clave])
            self.assertEqual(celda.number_format, '#,##0.00;[Red]-#,##0.00')
        self.assertTrue(datos["resumen"]["cuadrado"])
        self.assertEqual(totales[2].value, totales[3].value)
        self.assertEqual(totales[4].value, totales[5].value)
        self.assertEqual(totales[6].value, totales[7].value)

    def test_pantalla_muestra_miles_millones_y_centavos_sin_cambiar_resumen(self):
        respuesta = self.client.get(reverse("balance_comprobacion", args=[self.empresa.slug]), self.filtros)
        self.assertContains(respuesta, "1,250,000.50")
        self.assertContains(respuesta, "1,250.10")
        self.assertContains(respuesta, "2.35")
        self.assertContains(respuesta, "0.00")
        self.assertEqual(respuesta.context["resumen"], datos_balance_comprobacion(self.empresa, **self.filtros)["resumen"])

    def test_advertencia_historica_se_conserva_al_inicio_antes_del_encabezado(self):
        with patch("contabilidad.views._conversiones_pendientes_reporte", return_value=2):
            hoja = self.excel()
        aviso = next(fila for fila in hoja if fila[0].value == "Conversión histórica pendiente")
        encabezado = next(celda.row for celda in hoja["A"] if celda.value == "Cuenta")
        self.assertLess(aviso[0].row, encabezado)
        self.assertIn("2 asientos en dólares", aviso[1].value)

    def test_exportacion_respeta_filtros_empresa_y_conserva_datos(self):
        anteriores = {modelo: list(modelo.objects.order_by("pk").values()) for modelo in (CuentaContable, AsientoContable, LineaAsientoContable)}
        filtros = {**self.filtros, "tipo": "activo", "incluir_ceros": "1"}
        hoja = self.excel(filtros)
        etiquetas = {celda.value for celda in hoja["A"]}
        self.assertIn("1199 - Cuenta sin movimiento", etiquetas)
        self.assertNotIn("4101 - Ventas", etiquetas)
        self.assertNotIn("1101 - Cuenta confidencial", etiquetas)
        datos = datos_balance_comprobacion(self.empresa, **self.filtros, tipo="activo", incluir_ceros=True)
        totales = next(fila for fila in hoja if fila[0].value == "Totales")
        self.assertEqual(Decimal(str(totales[4].value)), datos["resumen"]["debe_total"])
        respuesta_ajena = self.client.get(reverse("exportar_balance_comprobacion_excel", args=[self.otra.slug]), self.filtros)
        self.assertIn(respuesta_ajena.status_code, (302, 403))
        for modelo, valores in anteriores.items():
            self.assertEqual(list(modelo.objects.order_by("pk").values()), valores)
