"""El Mayor conserva el alcance del Estado de Resultados sin alterar su vista normal."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Empresa, Modulo, PlanComercial, PlanModulo, Usuario
from .models import AsientoContable, CuentaContable, LineaAsientoContable
from .reportes import EVENTOS_CIERRE_RESULTADOS, datos_mayor


class MayorOrigenResultadosTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        modulo, _ = Modulo.objects.get_or_create(codigo="contabilidad", defaults={"nombre": "Contabilidad", "es_comercial": True})
        plan = PlanComercial.objects.create(nombre="Plan mayor resultados", codigo="plan-mayor-resultados", precio_mensual="99.00")
        PlanModulo.objects.create(plan=plan, modulo=modulo, activo=True)
        cls.empresa = Empresa.objects.create(nombre="Mayor desde resultados", slug="mayor-resultados", rtn="08011999288241", estado_licencia="activa", plan_comercial=plan)
        cls.otra_empresa = Empresa.objects.create(nombre="Otra empresa", slug="otro-mayor-resultados", rtn="08011999288242", estado_licencia="activa", plan_comercial=plan)
        cls.usuario = Usuario.objects.create_user(username="contador-mayor-resultados", empresa=cls.empresa, is_superuser=True, is_staff=True)
        cls.ingreso = cls.cuenta("4.1", "Ingresos", "ingreso")
        cls.otro_ingreso = cls.cuenta("4.2", "Otros ingresos", "ingreso")
        cls.costo = cls.cuenta("5.1", "Costos", "costo")
        cls.gasto = cls.cuenta("6.1", "Gastos", "gasto")
        cls.banco = cls.cuenta("1.1", "Banco", "activo")
        cls.capital = cls.cuenta("3.1", "Capital", "patrimonio")
        cls.ajena = cls.cuenta("4.1", "Ingreso privado", "ingreso", empresa=cls.otra_empresa)
        cls.anterior = cls.linea(cls.ingreso, "100.00", fecha=date(2025, 10, 1), detalle="Ingreso del ejercicio anterior")
        cls.cierre_anterior = cls.linea(cls.ingreso, "-100.00", fecha=date(2025, 12, 31), evento="cierre_anual", detalle="Cierre del ejercicio anterior")
        cls.seleccionada = cls.linea(cls.ingreso, "50.25", detalle="Ingreso del período")
        cls.devolucion = cls.linea(cls.ingreso, "-5.20", detalle="Devolución del período")
        cls.cierres = [
            cls.linea(cls.ingreso, importe, evento=evento, detalle=f"Cierre excluido {evento}")
            for evento, importe in zip(EVENTOS_CIERRE_RESULTADOS, ("-10.00", "-20.00", "-15.05"))
        ]
        cls.linea(cls.otro_ingreso, "999.00", detalle="Movimiento de otra cuenta")
        cls.linea(cls.ingreso, "888.00", fecha=date(2026, 11, 1), detalle="Movimiento posterior")
        cls.linea(cls.ingreso, "777.00", estado="borrador", detalle="Movimiento en borrador")

    @classmethod
    def cuenta(cls, codigo, nombre, tipo, empresa=None):
        return CuentaContable.objects.create(empresa=empresa or cls.empresa, codigo=codigo, nombre=nombre, tipo=tipo)

    @classmethod
    def linea(cls, cuenta, importe, *, fecha=date(2026, 10, 5), evento=None, detalle="Detalle de prueba", estado="contabilizado"):
        importe = Decimal(importe)
        asiento = AsientoContable.objects.create(empresa=cuenta.empresa, fecha=fecha, estado=estado, evento=evento, descripcion=detalle)
        deudora = cuenta.tipo in ("activo", "costo", "gasto")
        debe = max(importe if deudora else -importe, Decimal("0.00"))
        haber = max(-importe if deudora else importe, Decimal("0.00"))
        linea = LineaAsientoContable.objects.create(asiento=asiento, cuenta=cuenta, detalle=detalle, debe=debe, haber=haber)
        LineaAsientoContable.objects.create(asiento=asiento, cuenta=cls.capital, debe=haber, haber=debe)
        return linea

    def setUp(self):
        self.client.force_login(self.usuario)
        self.url = reverse("mayor_cuenta", args=[self.empresa.slug])

    def parametros(self, **cambios):
        return {"cuenta": self.ingreso.pk, "fecha_inicio": "2026-10-01", "fecha_fin": "2026-10-31", **cambios}

    def consultar(self, **cambios):
        return self.client.get(self.url, self.parametros(**cambios))

    def test_mayor_normal_conserva_cierres_y_saldo_real(self):
        bloque = datos_mayor(self.empresa, "2026-10-01", "2026-10-31", cuenta=self.ingreso)[0]
        self.assertEqual(bloque["saldo_inicial"], Decimal("0.00"))
        self.assertEqual(bloque["saldo"], Decimal("0.00"))
        self.assertEqual(len(bloque["movimientos"]), 5)
        self.assertEqual({m["linea"].pk for m in bloque["movimientos"]}, {self.seleccionada.pk, self.devolucion.pk, *(linea.pk for linea in self.cierres)})
        respuesta = self.consultar()
        self.assertFalse(respuesta.context["desde_estado_resultados"])
        self.assertEqual(respuesta.context["resumen"]["saldo"], Decimal("0.00"))
        self.assertContains(respuesta, "Cierre excluido cierre_anual")

    def test_filtro_omite_tres_cierres_del_periodo_y_conserva_apertura_real(self):
        bloque = datos_mayor(self.empresa, "2026-10-01", "2026-10-31", cuenta=self.ingreso, excluir_cierres=True)[0]
        self.assertEqual(bloque["saldo_inicial"], Decimal("0.00"))
        self.assertEqual(bloque["saldo"], Decimal("45.05"))
        self.assertEqual(bloque["debe"], Decimal("5.20"))
        self.assertEqual(bloque["haber"], Decimal("50.25"))
        self.assertEqual([m["linea"].pk for m in bloque["movimientos"]], [self.seleccionada.pk, self.devolucion.pk])

    def test_navegacion_preserva_filtros_seleccion_y_regreso_a_resultados(self):
        respuesta = self.consultar(origen="estado_resultados", movimiento=self.seleccionada.pk)
        self.assertEqual(respuesta.status_code, 200)
        self.assertTrue(respuesta.context["desde_estado_resultados"])
        self.assertEqual(respuesta.context["movimiento_seleccionado"], self.seleccionada.pk)
        self.assertEqual(respuesta.context["resumen"]["saldo"], Decimal("45.05"))
        self.assertContains(respuesta, '<input type="hidden" name="origen" value="estado_resultados">', html=True)
        self.assertContains(respuesta, "se omiten los asientos de cierre del período")
        self.assertContains(respuesta, "el saldo del Mayor incluye la apertura")
        self.assertContains(respuesta, f'{reverse("estado_resultados", args=[self.empresa.slug])}?fecha_inicio=2026-10-01&amp;fecha_fin=2026-10-31')
        self.assertNotContains(respuesta, "Volver al Balance")
        for texto in ("Cierre excluido", "Ingreso del ejercicio anterior", "Movimiento de otra cuenta", "Movimiento posterior", "Movimiento en borrador"):
            self.assertNotContains(respuesta, texto)

    def test_id_de_cierre_omitido_no_queda_resaltado(self):
        respuesta = self.consultar(origen="estado_resultados", movimiento=self.cierres[0].pk)
        self.assertIsNone(respuesta.context["movimiento_seleccionado"])
        self.assertNotContains(respuesta, "Movimiento seleccionado")
        self.assertContains(respuesta, "Volver al Estado de Resultados")

    def test_vista_completa_conserva_cuenta_fechas_seleccion_y_recupera_cierres(self):
        respuesta = self.consultar(origen="estado_resultados", movimiento=self.seleccionada.pk)
        self.assertContains(respuesta, f'href="{self.url}?cuenta={self.ingreso.pk}&amp;fecha_inicio=2026-10-01&amp;fecha_fin=2026-10-31&amp;movimiento={self.seleccionada.pk}#movimiento-{self.seleccionada.pk}"')
        completa = self.consultar(movimiento=self.seleccionada.pk)
        self.assertFalse(completa.context["desde_estado_resultados"])
        self.assertEqual(completa.context["movimiento_seleccionado"], self.seleccionada.pk)
        self.assertContains(completa, "Cierre excluido")

    def test_origen_solo_se_activa_en_cuentas_de_resultados(self):
        for cuenta in (self.ingreso, self.costo, self.gasto):
            with self.subTest(tipo=cuenta.tipo):
                self.assertTrue(self.consultar(cuenta=cuenta.pk, origen="estado_resultados").context["desde_estado_resultados"])
        for cuenta, origen in ((self.banco, "estado_resultados"), (self.capital, "estado_resultados"), (self.ingreso, "otro"), (self.ingreso, "")):
            with self.subTest(cuenta=cuenta.pk, origen=origen):
                respuesta = self.consultar(cuenta=cuenta.pk, origen=origen)
                self.assertFalse(respuesta.context["desde_estado_resultados"])
                self.assertNotContains(respuesta, 'name="origen"')
        self.assertEqual(self.consultar(cuenta=self.ajena.pk, origen="estado_resultados").status_code, 404)

    def test_excel_comparte_filtro_y_apertura_con_pantalla(self):
        url = reverse("exportar_mayor_cuenta_excel", args=[self.empresa.slug])
        filtrado = self.client.get(url, self.parametros(origen="estado_resultados"))
        completo = self.client.get(url, self.parametros())
        self.assertEqual(filtrado.status_code, 200)
        filas = list(load_workbook(BytesIO(filtrado.content), data_only=True).active.values)
        filas_completas = list(load_workbook(BytesIO(completo.content), data_only=True).active.values)
        detalles = [fila[3] for fila in filas if len(fila) > 3]
        detalles_completos = [fila[3] for fila in filas_completas if len(fila) > 3]
        self.assertIn("Ingreso del período", detalles)
        self.assertIn("Devolución del período", detalles)
        self.assertFalse(any(isinstance(detalle, str) and detalle.startswith("Cierre excluido") for detalle in detalles))
        self.assertTrue(any(isinstance(detalle, str) and detalle.startswith("Cierre excluido") for detalle in detalles_completos))
        self.assertEqual(next(fila[7] for fila in filas if fila[3] == "Saldo inicial anterior al periodo"), 0)
        self.assertEqual(next(fila[7] for fila in filas if fila[3] == "Totales del periodo / saldo final"), 45.05)
