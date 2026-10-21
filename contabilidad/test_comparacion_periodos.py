"""Cada columna comparativa conserva el alcance del reporte individual."""

from calendar import monthrange
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Empresa, Modulo, PlanComercial, PlanModulo, RolSistema, Usuario

from .models import AsientoContable, CuentaContable, LineaAsientoContable
from .reportes import datos_balance_general, datos_estado_resultados


CERO = Decimal("0.00")


class ComparacionPeriodosTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        modulo, _ = Modulo.objects.get_or_create(
            codigo="contabilidad", defaults={"nombre": "Contabilidad", "es_comercial": True},
        )
        plan = PlanComercial.objects.create(
            nombre="Plan comparación", codigo="plan-comparacion-periodos", precio_mensual="99.00",
        )
        PlanModulo.objects.create(plan=plan, modulo=modulo, activo=True)
        cls.empresa = Empresa.objects.create(
            nombre="Empresa comparativa", slug="comparacion-periodos", rtn="08011999410101",
            estado_licencia="activa", plan_comercial=plan,
        )
        cls.otra = Empresa.objects.create(
            nombre="Empresa privada comparativa", slug="comparacion-privada", rtn="08011999410102",
            estado_licencia="activa", plan_comercial=plan,
        )
        rol = RolSistema.objects.create(
            nombre="Consulta comparativa", codigo="consulta-comparativa",
            puede_contabilidad=True, puede_reportes_contables=True,
        )
        cls.usuario = Usuario.objects.create_user(
            username="contador-comparativo", empresa=cls.empresa, rol_sistema=rol,
        )
        cls.sin_permiso = Usuario.objects.create_user(username="sin-comparacion", empresa=cls.empresa)
        cls.usuario_ajeno = Usuario.objects.create_user(
            username="contador-comparativo-ajeno", empresa=cls.otra, rol_sistema=rol,
        )
        cls.activo = cls.cuenta("1", "Activos", "activo", acepta_movimientos=False)
        cls.bancos = cls.cuenta("1.1", "Bancos", "activo", cuenta_padre=cls.activo)
        cls.atlantida = cls.cuenta("ZZ-ATL", "Banco Atlántida", "activo", cuenta_padre=cls.bancos)
        cls.bac = cls.cuenta("1.1.02", "BAC", "activo", cuenta_padre=cls.bancos)
        cls.pasivo = cls.cuenta("2", "Pasivos", "pasivo", acepta_movimientos=False)
        cls.deuda = cls.cuenta("2.1", "Préstamo bancario", "pasivo", cuenta_padre=cls.pasivo)
        cls.patrimonio = cls.cuenta("3", "Patrimonio", "patrimonio", acepta_movimientos=False)
        cls.capital = cls.cuenta("3.1", "Capital", "patrimonio", cuenta_padre=cls.patrimonio)
        cls.ingresos = cls.cuenta("4", "Ingresos", "ingreso", acepta_movimientos=False)
        cls.ventas = cls.cuenta("4.1", "Ventas", "ingreso", cuenta_padre=cls.ingresos)
        cls.servicios = cls.cuenta("4.1.01", "Servicios profesionales", "ingreso", cuenta_padre=cls.ventas)
        cls.productos = cls.cuenta("4.1.02", "Venta de productos", "ingreso", cuenta_padre=cls.ventas)
        cls.costos = cls.cuenta("5", "Costos", "costo", acepta_movimientos=False)
        cls.costo = cls.cuenta("5.1", "Costo de ventas", "costo", cuenta_padre=cls.costos)
        cls.gastos = cls.cuenta("6", "Gastos", "gasto", acepta_movimientos=False)
        cls.gasto = cls.cuenta("6.1", "Servicios administrativos", "gasto", cuenta_padre=cls.gastos)
        cls.sin_movimientos = cls.cuenta("6.9", "Sin movimientos", "gasto", cuenta_padre=cls.gastos)

        cls.asiento(date(2025, 12, 1), [(cls.atlantida, "1000000.10", 0), (cls.capital, 0, "1000000.10")])
        cls.asiento(date(2025, 12, 15), [(cls.atlantida, "80.00", 0), (cls.servicios, 0, "80.00")])
        cls.asiento(
            date(2025, 12, 31), [(cls.servicios, "80.00", 0), (cls.capital, 0, "80.00")],
            evento="cierre_anual",
        )
        for mes in range(1, 10):
            if mes == 4:
                continue  # Saldo conservado en Balance; movimiento cero en Resultados.
            ingreso = Decimal(mes * 1000) + Decimal("0.10")
            costo = Decimal(mes * 100) + Decimal("0.20")
            gasto = Decimal(mes * 10) + Decimal("0.30")
            cls.asiento(date(2026, mes, 5), [
                (cls.atlantida, ingreso, 0), (cls.ventas, 0, "100.00"),
                (cls.servicios, 0, ingreso - Decimal("500.00")), (cls.productos, 0, "400.00"),
            ])
            cls.asiento(date(2026, mes, 10), [(cls.costo, costo, 0), (cls.atlantida, 0, costo)])
            cls.asiento(date(2026, mes, 15), [(cls.gasto, gasto, 0), (cls.atlantida, 0, gasto)])
            cls.asiento(date(2026, mes, 20), [(cls.bancos, "5.05", 0), (cls.capital, 0, "5.05")])
            if mes == 3:
                cls.asiento(date(2026, mes, 21), [(cls.bac, "100.50", 0), (cls.deuda, 0, "100.50")])
            if mes == 6:
                # Un cierre real ya queda incluido en patrimonio por la lógica actual.
                datos = datos_estado_resultados(cls.empresa, fecha_fin="2026-06-30", excluir_cierres=False)
                cierre = [
                    (fila["cuenta"], fila["saldo"], CERO) for fila in datos["ingresos"]
                ] + [
                    (fila["cuenta"], CERO, fila["saldo"]) for fila in [*datos["costos"], *datos["gastos"]]
                ] + [(cls.capital, CERO, datos["resumen"]["utilidad"])]
                cls.asiento(date(2026, 6, 30), cierre, evento="traspaso_resultados")

        cls.asiento(date(2026, 9, 25), [(cls.servicios, "5.25", 0), (cls.atlantida, 0, "5.25")])
        for estado in ("borrador", "anulado"):
            cls.asiento(date(2026, 9, 26), [(cls.bac, "999999.99", 0), (cls.productos, 0, "999999.99")], estado=estado)
        cls.asiento(date(2026, 10, 1), [(cls.bac, "99999.99", 0), (cls.productos, 0, "99999.99")])
        CuentaContable.objects.filter(pk=cls.atlantida.pk).update(activa=False)
        cls.ajena = cls.cuenta("ZZ-ATL", "Banco confidencial", "activo", empresa=cls.otra)
        cls.ingreso_ajeno = cls.cuenta("4.1.01", "Ingreso confidencial", "ingreso", empresa=cls.otra)
        cls.asiento(
            date(2026, 2, 1), [(cls.ajena, "9000000.99", 0), (cls.ingreso_ajeno, 0, "9000000.99")],
            empresa=cls.otra,
        )

    @classmethod
    def cuenta(cls, codigo, nombre, tipo, empresa=None, **campos):
        return CuentaContable.objects.create(
            empresa=empresa or cls.empresa, codigo=codigo, nombre=nombre, tipo=tipo, **campos,
        )

    @classmethod
    def asiento(cls, fecha, partidas, *, estado="contabilizado", evento=None, empresa=None):
        asiento = AsientoContable.objects.create(
            empresa=empresa or cls.empresa, fecha=fecha, estado=estado, evento=evento,
            descripcion="Operación comparativa de prueba",
        )
        for cuenta, debe, haber in partidas:
            LineaAsientoContable.objects.create(
                asiento=asiento, cuenta=cuenta, debe=Decimal(str(debe)), haber=Decimal(str(haber)),
            )
        return asiento

    def setUp(self):
        self.client.force_login(self.usuario)

    def consultar(self, reporte, **filtros):
        respuesta = self.client.get(reverse(reporte, args=[self.empresa.slug]), {
            "comparar": "1", "mes_inicio": "2026-01", "mes_fin": "2026-09", **filtros,
        })
        self.assertEqual(respuesta.status_code, 200)
        return respuesta

    def comparacion(self, reporte, **filtros):
        comparacion = self.consultar(reporte, **filtros).context["comparacion"]
        self.assertTrue(comparacion)
        return comparacion

    def grupos(self, comparacion):
        return {grupo["clave"]: grupo for grupo in comparacion["grupos"]}

    def resumen(self, comparacion):
        return {fila["clave"]: fila["saldos"] for fila in comparacion["resumen_filas"]}

    def comprobar_columnas(self, comparacion, datos_mensuales, grupos):
        for grupo in comparacion["grupos"]:
            clave = grupo["clave"]
            filas = {fila["cuenta"].pk: fila for fila in grupo["filas"] if fila.get("cuenta")}
            for indice, datos in enumerate(datos_mensuales):
                individuales = {fila["cuenta"].pk: fila for fila in datos[f"{clave}_jerarquia"]}
                for pk, fila in filas.items():
                    with self.subTest(grupo=clave, cuenta=pk, mes=indice):
                        self.assertEqual(fila["saldos"][indice], individuales.get(pk, {}).get("saldo", CERO))
                        self.assertIsInstance(fila["saldos"][indice], Decimal)
                self.assertEqual(grupo["totales"][indice], datos["resumen"][grupos[clave]])

    def test_balance_enero_septiembre_coincide_con_cada_cierre_individual(self):
        comparacion = self.comparacion("balance_general")
        self.assertEqual(len(comparacion["columnas"]), 9)
        self.assertFalse(any(columna["es_acumulado"] for columna in comparacion["columnas"]))
        datos = [datos_balance_general(self.empresa, date(2026, mes, monthrange(2026, mes)[1])) for mes in range(1, 10)]
        self.comprobar_columnas(comparacion, datos, {
            "activos": "total_activos", "pasivos": "total_pasivos", "patrimonio": "total_patrimonio",
        })
        resumen = self.resumen(comparacion)
        for indice, individual in enumerate(datos):
            self.assertEqual(resumen["diferencia"][indice], CERO)
            self.assertEqual(resumen["total_activos"][indice], resumen["total_pasivo_patrimonio"][indice])
            self.assertTrue(individual["resumen"]["cuadrado"])
        activos = self.grupos(comparacion)["activos"]
        self.assertEqual(activos["totales"][2], activos["totales"][3])
        self.assertNotEqual(activos["totales"][-1], sum(activos["totales"], CERO))

    def test_resultados_enero_septiembre_coincide_con_mes_y_acumulado_individual(self):
        comparacion = self.comparacion("estado_resultados")
        self.assertEqual(len(comparacion["columnas"]), 10)
        self.assertTrue(comparacion["columnas"][-1]["es_acumulado"])
        datos = [
            datos_estado_resultados(self.empresa, date(2026, mes, 1), date(2026, mes, monthrange(2026, mes)[1]))
            for mes in range(1, 10)
        ]
        acumulado = datos_estado_resultados(self.empresa, "2026-01-01", "2026-09-30")
        self.comprobar_columnas(comparacion, [*datos, acumulado], {
            "ingresos": "total_ingresos", "costos": "total_costos", "gastos": "total_gastos",
        })
        for clave, saldos in self.resumen(comparacion).items():
            for indice, individual in enumerate([*datos, acumulado]):
                self.assertEqual(saldos[indice], individual["resumen"][clave])
        for grupo in comparacion["grupos"]:
            self.assertEqual(grupo["totales"][3], CERO)
            self.assertEqual(grupo["totales"][-1], sum(grupo["totales"][:-1], CERO))
        self.assertGreater(self.resumen(comparacion)["utilidad"][5], CERO)

    def test_union_conserva_jerarquia_principales_y_cuentas_inactivas_sin_duplicar_totales(self):
        for reporte, clave, padre, hijo, nieto in (
            ("balance_general", "activos", self.activo, self.bancos, self.atlantida),
            ("estado_resultados", "ingresos", self.ingresos, self.ventas, self.servicios),
        ):
            comparacion = self.comparacion(reporte)
            grupo = self.grupos(comparacion)[clave]
            filas = {fila["cuenta"].pk: fila for fila in grupo["filas"] if fila.get("cuenta")}
            orden = list(filas)
            self.assertLess(orden.index(padre.pk), orden.index(hijo.pk))
            self.assertLess(orden.index(hijo.pk), orden.index(nieto.pk))
            self.assertTrue(filas[padre.pk]["es_agrupadora"])
            self.assertTrue(filas[hijo.pk]["es_agrupadora"])
            self.assertFalse(filas[nieto.pk]["es_agrupadora"])
            self.assertEqual(filas[nieto.pk]["nivel"], 2)
            self.assertEqual(filas[nieto.pk]["sangria_px"], 36)
            self.assertEqual(filas[padre.pk]["saldos"], grupo["totales"])
            self.assertNotIn(self.ajena.pk, filas)
            self.assertNotIn(self.ingreso_ajeno.pk, filas)

    def test_modo_desactivado_conserva_contexto_y_filtros_del_reporte_actual(self):
        for reporte, filtros, esperado in (
            ("balance_general", {"fecha_corte": "2026-05-12"}, datos_balance_general(self.empresa, "2026-05-12")),
            ("estado_resultados", {"fecha_inicio": "2026-02-05", "fecha_fin": "2026-07-17"},
             datos_estado_resultados(self.empresa, "2026-02-05", "2026-07-17")),
        ):
            for comparar in (None, "0"):
                parametros = {**filtros, "mes_inicio": "invalido", "mes_fin": "invalido"}
                if comparar is not None:
                    parametros["comparar"] = comparar
                respuesta = self.client.get(reverse(reporte, args=[self.empresa.slug]), parametros)
                self.assertEqual(respuesta.status_code, 200)
                self.assertFalse(respuesta.context.get("comparacion"))
                self.assertEqual(respuesta.context["resumen"], esperado["resumen"])
                self.assertEqual(respuesta.context["filtros"], filtros)

    def test_cambio_de_anio_y_febrero_bisiesto_toman_el_ultimo_dia_real(self):
        for reporte in ("balance_general", "estado_resultados"):
            comparacion = self.comparacion(reporte, mes_inicio="2023-12", mes_fin="2024-03")
            columnas = [columna for columna in comparacion["columnas"] if not columna["es_acumulado"]]
            self.assertEqual([str(columna["inicio"]) for columna in columnas], [
                "2023-12-01", "2024-01-01", "2024-02-01", "2024-03-01",
            ])
            self.assertEqual([str(columna["fin"]) for columna in columnas], [
                "2023-12-31", "2024-01-31", "2024-02-29", "2024-03-31",
            ])
            self.assertEqual([columna["anio"] for columna in columnas], [2023, 2024, 2024, 2024])

    def test_rango_invalido_muestra_error_y_excel_lo_rechaza(self):
        invalidos = (
            ("", "2026-09"), ("2026-01", ""), ("texto", "2026-09"),
            ("2026-1", "2026-09"), ("2026-13", "2026-09"),
            ("2026-09", "2026-01"), ("0000-01", "2026-01"),
            ("2010-01", "2026-09"),
        )
        for reporte in ("balance_general", "estado_resultados"):
            for inicio, fin in invalidos:
                with self.subTest(reporte=reporte, inicio=inicio, fin=fin):
                    parametros = {"comparar": "1", "mes_inicio": inicio, "mes_fin": fin}
                    respuesta = self.client.get(reverse(reporte, args=[self.empresa.slug]), parametros)
                    self.assertEqual(respuesta.status_code, 200)
                    self.assertTrue(respuesta.context["comparacion_error"])
                    self.assertFalse(respuesta.context.get("comparacion"))
                    self.assertContains(respuesta, 'role="alert"')
                    self.assertNotContains(respuesta, 'class="erp-table ef-tabla')
                    self.assertEqual(respuesta.context["comparacion_form"]["mes_inicio"], inicio)
                    self.assertEqual(respuesta.context["comparacion_form"]["mes_fin"], fin)
                    excel = self.client.get(reverse(f"exportar_{reporte}_excel", args=[self.empresa.slug]), parametros)
                    self.assertEqual(excel.status_code, 400)

    def test_rango_de_un_mes_no_incluye_mes_adyacente(self):
        balance = self.comparacion("balance_general", mes_inicio="2026-09", mes_fin="2026-09")
        self.assertEqual(len(balance["columnas"]), 1)
        resultados = self.comparacion("estado_resultados", mes_inicio="2026-09", mes_fin="2026-09")
        self.assertEqual(len(resultados["columnas"]), 2)
        for grupo in resultados["grupos"]:
            self.assertEqual(grupo["totales"][0], grupo["totales"][1])

    def test_consulta_y_exportacion_respetan_empresa_y_permiso(self):
        for usuario in (self.sin_permiso, self.usuario_ajeno):
            self.client.force_login(usuario)
            for reporte in ("balance_general", "estado_resultados", "exportar_balance_general_excel", "exportar_estado_resultados_excel"):
                respuesta = self.client.get(reverse(reporte, args=[self.empresa.slug]), {
                    "comparar": "1", "mes_inicio": "2026-01", "mes_fin": "2026-09",
                })
                self.assertIn(respuesta.status_code, (302, 403))
                self.assertNotIn(self.atlantida.nombre.encode(), respuesta.content)
        self.client.force_login(self.usuario)
        for reporte in ("balance_general", "estado_resultados"):
            respuesta = self.consultar(reporte, empresa=self.otra.pk)
            self.assertNotContains(respuesta, self.ajena.nombre)
            self.assertNotContains(respuesta, self.ingreso_ajeno.nombre)
            self.assertEqual(respuesta.context["empresa"].pk, self.empresa.pk)

    def test_reportes_no_modifican_catalogo_asientos_ni_partidas(self):
        anteriores = {
            modelo: list(modelo.objects.order_by("pk").values())
            for modelo in (CuentaContable, AsientoContable, LineaAsientoContable)
        }
        for reporte in ("balance_general", "estado_resultados", "exportar_balance_general_excel", "exportar_estado_resultados_excel"):
            self.consultar(reporte)
        for modelo, valores in anteriores.items():
            self.assertEqual(list(modelo.objects.order_by("pk").values()), valores)

    def test_excel_conserva_columnas_importes_decimales_negrita_e_indentacion(self):
        for reporte in ("balance_general", "estado_resultados"):
            comparacion = self.comparacion(reporte)
            respuesta = self.consultar(f"exportar_{reporte}_excel")
            hoja = load_workbook(BytesIO(respuesta.content)).active
            encabezado = next(fila for fila in hoja if any(celda.value == "Cuenta" for celda in fila))
            columna_cuenta = next(celda.column for celda in encabezado if celda.value == "Cuenta")
            self.assertEqual(hoja.max_column, columna_cuenta + len(comparacion["columnas"]))
            self.assertTrue(hoja.freeze_panes)
            self.assertGreaterEqual(hoja[hoja.freeze_panes].column, columna_cuenta + 1)
            for grupo in comparacion["grupos"]:
                for fila in grupo["filas"]:
                    if not fila.get("cuenta"):
                        continue
                    cuenta = fila["cuenta"]
                    fila_excel = next(
                        registro for registro in hoja
                        if registro[columna_cuenta - 1].value == cuenta.nombre
                    )
                    self.assertEqual(fila_excel[columna_cuenta - 1].alignment.indent, fila["nivel"])
                    self.assertEqual(bool(fila_excel[columna_cuenta - 1].font.bold), fila["es_agrupadora"])
                    for indice, esperado in enumerate(fila["saldos"]):
                        celda = fila_excel[columna_cuenta + indice]
                        self.assertEqual(celda.data_type, "n")
                        self.assertEqual(Decimal(str(celda.value)), esperado)
                        self.assertIn("#,##0.00", celda.number_format)
                        self.assertEqual(bool(celda.font.bold), fila["es_agrupadora"])
