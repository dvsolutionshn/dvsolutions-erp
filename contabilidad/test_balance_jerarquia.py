"""El balance acumula sus principales sin duplicar movimientos ni resultados."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import Empresa, Modulo, PlanComercial, PlanModulo, Usuario
from .models import AsientoContable, CuentaContable, LineaAsientoContable
from .reportes import datos_balance_general, datos_estado_resultados, saldos_financieros


class BalanceJerarquiaTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        modulo, _ = Modulo.objects.get_or_create(
            codigo="contabilidad", defaults={"nombre": "Contabilidad", "es_comercial": True},
        )
        plan = PlanComercial.objects.create(
            nombre="Plan balance jerárquico", codigo="plan-balance-jerarquia", precio_mensual="99.00",
        )
        PlanModulo.objects.create(plan=plan, modulo=modulo, activo=True)
        cls.empresa = Empresa.objects.create(
            nombre="Empresa del balance", slug="balance-jerarquia", rtn="08011999188221",
            estado_licencia="activa", plan_comercial=plan,
        )
        cls.otra = Empresa.objects.create(
            nombre="Otra empresa del balance", slug="otro-balance-jerarquia", rtn="08011999188222",
            estado_licencia="activa", plan_comercial=plan,
        )
        cls.usuario = Usuario.objects.create_user(
            username="contador-balance-jerarquia", empresa=cls.empresa, is_superuser=True, is_staff=True,
        )
        cls.activo = cls.cuenta("1", "Activos", acepta_movimientos=False)
        cls.bancos = cls.cuenta("1.1.01", "Bancos", cuenta_padre=cls.activo)
        cls.atlantida = cls.cuenta("1.1.01.01", "Banco Atlántida", cuenta_padre=cls.bancos)
        cls.bac = cls.cuenta("1.1.01.02", "BAC", cuenta_padre=cls.bancos)
        cls.deposito = cls.cuenta("1.1.01.01.01", "Depósito de Atlántida", cuenta_padre=cls.atlantida)
        cls.pasivo = cls.cuenta("2", "Pasivos", "pasivo", acepta_movimientos=False)
        cls.deuda = cls.cuenta("2.1", "Préstamo", "pasivo", cuenta_padre=cls.pasivo)
        cls.patrimonio = cls.cuenta("3", "Patrimonio", "patrimonio", acepta_movimientos=False)
        cls.capital = cls.cuenta("3.1", "Capital", "patrimonio", cuenta_padre=cls.patrimonio)
        cls.ingreso = cls.cuenta("4.1", "Ingresos", "ingreso")
        cls.gasto = cls.cuenta("5.1", "Gastos", "gasto")

    @classmethod
    def cuenta(cls, codigo, nombre, tipo="activo", empresa=None, **kwargs):
        return CuentaContable.objects.create(
            empresa=empresa or cls.empresa, codigo=codigo, nombre=nombre, tipo=tipo, **kwargs,
        )

    def asiento(self, partidas, *, fecha=date(2026, 10, 1), empresa=None, estado="contabilizado", evento=None):
        asiento = AsientoContable.objects.create(
            empresa=empresa or self.empresa, fecha=fecha, descripcion="Prueba del balance jerárquico",
            estado=estado, evento=evento,
        )
        for cuenta, debe, haber in partidas:
            LineaAsientoContable.objects.create(
                asiento=asiento, cuenta=cuenta, debe=Decimal(str(debe)), haber=Decimal(str(haber)),
            )
        return asiento

    def filas(self, datos, grupo="activos"):
        return {fila["cuenta"].pk: fila for fila in datos[f"{grupo}_jerarquia"]}

    def test_bancos_sin_movimiento_directo_acumula_varias_subcuentas_y_total_una_vez(self):
        self.asiento([(self.atlantida, 100000, 0), (self.bac, 50000, 0), (self.capital, 0, 150000)])
        datos = datos_balance_general(self.empresa)
        filas = self.filas(datos)
        self.assertEqual(filas[self.bancos.pk]["saldo"], Decimal("150000.00"))
        self.assertEqual(filas[self.bancos.pk]["saldo_directo"], Decimal("0.00"))
        self.assertEqual(filas[self.activo.pk]["saldo"], Decimal("150000.00"))
        self.assertEqual(datos["resumen"]["total_activos"], Decimal("150000.00"))
        self.assertEqual(datos["resumen"]["total_patrimonio"], Decimal("150000.00"))
        self.assertTrue(datos["resumen"]["cuadrado"])
        self.assertEqual([fila["cuenta"].pk for fila in datos["activos"]], [self.atlantida.pk, self.bac.pk])
        self.assertEqual(
            [(fila["cuenta"].pk, fila["nivel"]) for fila in datos["activos_jerarquia"]],
            [(self.activo.pk, 0), (self.bancos.pk, 1), (self.atlantida.pk, 2), (self.bac.pk, 2)],
        )

    def test_suma_propio_nietos_y_contracuenta_con_decimales_exactos(self):
        contra = self.cuenta("1.1.01.03", "Ajuste acreedor del banco", cuenta_padre=self.bancos)
        self.asiento([
            (self.bancos, "10.10", 0), (self.atlantida, "100.20", 0),
            (self.deposito, "50.30", 0), (self.bac, "40.40", 0), (self.capital, 0, "201.00"),
        ])
        self.asiento([(self.gasto, "5.25", 0), (contra, 0, "5.25")])
        datos = datos_balance_general(self.empresa)
        filas = self.filas(datos)
        self.assertEqual(filas[self.bancos.pk]["saldo"], Decimal("195.75"))
        self.assertEqual(filas[self.bancos.pk]["saldo_directo"], Decimal("10.10"))
        self.assertEqual(filas[self.atlantida.pk]["saldo"], Decimal("150.50"))
        self.assertEqual(filas[self.deposito.pk]["nivel"], 3)
        self.assertEqual(filas[contra.pk]["saldo"], Decimal("-5.25"))
        self.assertEqual(datos["resumen"]["total_activos"], Decimal("195.75"))
        self.assertEqual(datos["resumen"]["resultado_no_traspasado"], Decimal("-5.25"))
        self.assertTrue(datos["resumen"]["cuadrado"])

    def test_pasivos_y_patrimonio_acumulan_y_resultado_no_se_duplica_al_cerrar(self):
        self.asiento([(self.atlantida, 150, 0), (self.deuda, 0, 50), (self.capital, 0, 100)])
        self.asiento([(self.bac, 30, 0), (self.ingreso, 0, 30)])
        anterior = datos_balance_general(self.empresa)
        self.assertEqual(self.filas(anterior, "pasivos")[self.pasivo.pk]["saldo"], Decimal("50.00"))
        self.assertEqual(self.filas(anterior, "patrimonio")[self.patrimonio.pk]["saldo"], Decimal("100.00"))
        self.assertEqual(anterior["resumen"]["total_patrimonio"], Decimal("130.00"))
        self.assertTrue(anterior["resumen"]["cuadrado"])
        self.asiento([(self.ingreso, 30, 0), (self.capital, 0, 30)], evento="cierre_anual")
        cerrado = datos_balance_general(self.empresa)
        self.assertEqual(cerrado["resumen"]["resultado_no_traspasado"], Decimal("0.00"))
        self.assertEqual(self.filas(cerrado, "patrimonio")[self.patrimonio.pk]["saldo"], Decimal("130.00"))
        self.assertEqual(cerrado["resumen"]["total_patrimonio"], Decimal("130.00"))
        self.assertTrue(cerrado["resumen"]["cuadrado"])

    def test_conserva_corte_historico_cuentas_inactivas_y_excluye_borrador_anulado(self):
        self.asiento([(self.atlantida, 100, 0), (self.capital, 0, 100)], fecha=date(2026, 9, 30))
        self.asiento([(self.bac, 50, 0), (self.capital, 0, 50)], fecha=date(2026, 10, 1))
        self.asiento([(self.bac, 999, 0), (self.capital, 0, 999)], estado="borrador")
        self.asiento([(self.bac, 999, 0), (self.capital, 0, 999)], estado="anulado")
        CuentaContable.objects.filter(pk__in=[self.bancos.pk, self.atlantida.pk]).update(activa=False)
        datos = datos_balance_general(self.empresa, "2026-09-30")
        self.assertEqual(self.filas(datos)[self.bancos.pk]["saldo"], Decimal("100.00"))
        self.assertNotIn(self.bac.pk, self.filas(datos))
        self.assertEqual(datos["resumen"]["total_activos"], Decimal("100.00"))
        self.assertTrue(datos["resumen"]["cuadrado"])

    def test_usa_fk_no_prefijos_y_omite_ramas_sin_movimientos(self):
        enlazada = self.cuenta("CODIGO-DIFERENTE", "Subcuenta enlazada", cuenta_padre=self.bancos)
        no_enlazada = self.cuenta("1.1.01.99", "Código similar sin padre")
        self.asiento([(enlazada, 100, 0), (no_enlazada, 50, 0), (self.capital, 0, 150)])
        datos = datos_balance_general(self.empresa)
        filas = self.filas(datos)
        self.assertEqual(filas[self.bancos.pk]["saldo"], Decimal("100.00"))
        self.assertEqual(filas[enlazada.pk]["nivel"], 2)
        self.assertEqual(filas[no_enlazada.pk]["nivel"], 0)
        self.assertNotIn(self.atlantida.pk, filas)
        self.assertNotIn(self.deposito.pk, filas)
        self.assertEqual(datos["resumen"]["total_activos"], Decimal("150.00"))

    def test_cuentas_netas_cero_con_movimientos_siguen_visibles_con_su_principal(self):
        self.asiento([(self.atlantida, 10, 0), (self.capital, 0, 10)])
        self.asiento([(self.capital, 10, 0), (self.atlantida, 0, 10)])
        datos = datos_balance_general(self.empresa)
        filas = self.filas(datos)
        self.assertEqual(filas[self.bancos.pk]["debe"], Decimal("10.00"))
        self.assertEqual(filas[self.bancos.pk]["haber"], Decimal("10.00"))
        self.assertEqual(filas[self.bancos.pk]["saldo"], Decimal("0.00"))
        self.assertTrue(datos["resumen"]["cuadrado"])

    def test_empresas_aisladas_incluso_con_relacion_padre_historica_incorrecta(self):
        padre_ajeno = self.cuenta("1.1.01", "Banco de otra empresa", empresa=self.otra)
        hoja_ajena = self.cuenta("1.1.01.01", "Hoja privada", empresa=self.otra, cuenta_padre=padre_ajeno)
        capital_ajeno = self.cuenta("3", "Capital privado", "patrimonio", empresa=self.otra)
        self.asiento([(self.atlantida, 100, 0), (self.capital, 0, 100)])
        self.asiento([(hoja_ajena, 900, 0), (capital_ajeno, 0, 900)], empresa=self.otra)
        CuentaContable.objects.filter(pk=self.bancos.pk).update(cuenta_padre=padre_ajeno)
        datos = datos_balance_general(self.empresa)
        filas = self.filas(datos)
        self.assertEqual(filas[self.bancos.pk]["saldo"], Decimal("100.00"))
        self.assertEqual(filas[self.bancos.pk]["nivel"], 0)
        self.assertNotIn(padre_ajeno.pk, filas)
        self.assertNotIn(hoja_ajena.pk, filas)
        self.assertEqual(datos["resumen"]["total_activos"], Decimal("100.00"))
        self.assertTrue(datos["resumen"]["cuadrado"])
        self.bancos.refresh_from_db()
        self.assertEqual(self.bancos.cuenta_padre_id, padre_ajeno.pk)

    def test_tipos_incompatibles_y_ciclos_antiguos_no_mezclan_totales_ni_modifican_catalogo(self):
        self.asiento([(self.atlantida, 100, 0), (self.deuda, 0, 100)])
        CuentaContable.objects.filter(pk=self.deuda.pk).update(cuenta_padre=self.bancos)
        CuentaContable.objects.filter(pk=self.bancos.pk).update(cuenta_padre=self.atlantida)
        datos = datos_balance_general(self.empresa)
        filas = self.filas(datos)
        self.assertEqual(len(filas), 2)
        self.assertEqual(filas[self.bancos.pk]["saldo"], Decimal("100.00"))
        self.assertEqual(self.filas(datos, "pasivos")[self.deuda.pk]["nivel"], 0)
        self.assertEqual(datos["resumen"]["total_activos"], Decimal("100.00"))
        self.assertEqual(datos["resumen"]["total_pasivos"], Decimal("100.00"))
        self.assertTrue(datos["resumen"]["cuadrado"])
        self.bancos.refresh_from_db()
        self.deuda.refresh_from_db()
        self.assertEqual(self.bancos.cuenta_padre_id, self.atlantida.pk)
        self.assertEqual(self.deuda.cuenta_padre_id, self.bancos.pk)

    def test_saldos_directos_y_estado_resultados_conservan_su_contrato(self):
        self.asiento([(self.atlantida, "0.10", 0), (self.bac, "0.20", 0), (self.ingreso, 0, "0.30")])
        directas = saldos_financieros(self.empresa, ("activo", "pasivo", "patrimonio"))
        datos = datos_balance_general(self.empresa)
        self.assertEqual(datos["activos"], directas)
        self.assertEqual(set(directas[0]), {"cuenta", "debe", "haber", "saldo"})
        self.assertEqual(datos_estado_resultados(self.empresa)["resumen"]["utilidad"], Decimal("0.30"))
        self.assertEqual(datos["resumen"]["total_patrimonio"], Decimal("0.30"))
        self.assertTrue(datos["resumen"]["cuadrado"])

    def test_pantalla_y_excel_muestran_preorden_sangria_y_mismos_totales(self):
        self.asiento([(self.atlantida, 100000, 0), (self.deposito, "0.50", 0), (self.bac, 50000, 0), (self.capital, 0, "150000.50")])
        self.client.force_login(self.usuario)
        respuesta = self.client.get(reverse("balance_general", args=[self.empresa.slug]))
        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, '<th class="ef-codigo-columna" scope="col">Código</th>', count=3, html=True)
        self.assertNotContains(respuesta, "<th>Debe</th>", html=True)
        self.assertNotContains(respuesta, "<th>Haber</th>", html=True)
        self.assertContains(respuesta, f'class="ef-fila ef-principal" data-cuenta-codigo="{self.bancos.codigo}"')
        self.assertContains(respuesta, f'class="ef-fila" data-cuenta-codigo="{self.bac.codigo}"')
        self.assertContains(respuesta, "padding-inline-start: 54px;")
        self.assertTrue(respuesta.context["resumen"]["cuadrado"])
        excel = self.client.get(reverse("exportar_balance_general_excel", args=[self.empresa.slug]))
        self.assertEqual(excel.status_code, 200)
        hoja = load_workbook(BytesIO(excel.content), data_only=True).active
        self.assertEqual([celda.value for celda in hoja[3]], ["Código", "Cuenta", "Saldo"])
        cuentas = (self.activo, self.bancos, self.atlantida, self.deposito, self.bac)
        codigos = [cuenta.codigo for cuenta in cuentas]
        filas = [fila for fila in hoja.iter_rows(min_row=4) if fila[0].value in codigos]
        self.assertEqual([fila[0].value for fila in filas], codigos)
        self.assertEqual([fila[1].value for fila in filas], [cuenta.nombre for cuenta in cuentas])
        self.assertEqual([fila[1].alignment.indent for fila in filas], [0, 1, 2, 3, 2])
        self.assertTrue(filas[1][1].font.bold)
        self.assertTrue(filas[1][2].font.bold)
        self.assertFalse(filas[4][1].font.bold)
        self.assertFalse(filas[4][2].font.bold)
        self.assertEqual(filas[1][2].value, 150000.5)
        self.assertEqual(filas[1][2].number_format, '#,##0.00;[Red]-#,##0.00')
        valores = list(hoja.values)
        self.assertEqual(next(fila[2] for fila in valores if fila[0] == "Total activos"), 150000.5)
        self.assertEqual(next(fila[2] for fila in valores if fila[0] == "Pasivo + patrimonio"), 150000.5)
        self.assertEqual(next(fila[2] for fila in valores if fila[0] == "Diferencia"), 0)
