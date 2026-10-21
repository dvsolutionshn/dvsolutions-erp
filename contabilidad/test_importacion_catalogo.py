"""Contratos del importador de catálogos, sin depender de un archivo personal."""
from datetime import date
from decimal import Decimal
from io import BytesIO
import json

from django.core.exceptions import ValidationError
from django.test import TestCase
from openpyxl import Workbook

from core.models import Empresa
from facturacion.models import Cliente, Proveedor
from .importadores import (
    analizar_catalogo_cuentas_desde_excel,
    aplicar_catalogo_cuentas,
    importar_catalogo_cuentas_desde_excel,
)
from .models import (
    AsientoContable, ConfiguracionContableEmpresa, CuentaContable,
    CuentaFinanciera, LineaAsientoContable,
)


ENCABEZADOS_RUMI = [
    "Código RUMI", "Cuenta RUMI", "Nivel", "Cuenta padre", "Permite asiento",
    "Naturaleza", "Activación", "Tipo funcional ERP", "Conciliar", "Referencia CNBS",
    "Alcance de la referencia", "Auxiliar mínimo", "Fuente de referencia (URL)",
    "Autoría del nombre",
]


def archivo_excel(workbook):
    archivo = BytesIO()
    workbook.save(archivo)
    archivo.seek(0)
    return archivo


def fila_rumi(codigo, nombre, nivel, padre="", permite="No", activacion="Estructura", funcion=""):
    return [
        codigo, nombre, {1: "Clase", 2: "Grupo", 3: "Subgrupo", 4: "Detalle"}.get(nivel, nivel),
        padre, permite, "Deudora", activacion, funcion, "No",
        "Referencia propuesta", "Información de la propuesta", "Documento",
        "https://ejemplo.invalid/catalogo", "Nombre propuesto",
    ]


def libro_rumi():
    workbook = Workbook()
    workbook.active.title = "Portada"
    workbook.active.append(["Propuesta de catálogo"])
    sheet = workbook.create_sheet("Catalogo_RUMI")
    sheet.append(["Catálogo RUMI"])
    for _fila in range(4):
        sheet.append([])
    sheet.append(ENCABEZADOS_RUMI)
    for fila in [
        fila_rumi("1", "Activo", 1),
        fila_rumi("11", "Activo corriente", 2, "1"),
        fila_rumi("1101", "Caja", 3, "11", "Sí", "Base propuesta", "Caja"),
        fila_rumi("1102", "Bancos por evaluar", 3, "11", "Sí", "Condicionada", "Banco"),
        fila_rumi("2", "Pasivo", 1),
        fila_rumi("21", "Proveedores", 2, "2", "Sí", "Base propuesta", "Proveedor"),
        fila_rumi("3", "Patrimonio", 1),
        fila_rumi("4", "Activos y Pasivos Contingentes", 1, activacion="Control extracontable"),
        fila_rumi("41", "Garantías", 2, "4", "Sí", "Control extracontable", "Garantía"),
        fila_rumi("5", "Ingresos", 1),
        fila_rumi("51", "Servicios", 2, "5", "Sí", "Base propuesta", "Ingreso por servicio"),
        fila_rumi("6", "Gastos", 1),
        fila_rumi("61", "Gastos operativos", 2, "6", "Sí", "Base propuesta", "Gasto operativo"),
        fila_rumi("7", "Cuentas de Orden", 1, activacion="Control extracontable"),
        fila_rumi("99", "Referencia que no se carga", 1, activacion="No cargar"),
    ]:
        sheet.append(fila)
    return workbook


class ImportacionCatalogoTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.empresa = Empresa.objects.create(
            nombre="RUMI", slug="rumi", rtn="08011999111221",
            estado_licencia="activa",
        )
        cls.otra = Empresa.objects.create(
            nombre="Catálogo B", slug="catalogo-b", rtn="08011999111222",
            estado_licencia="activa",
        )

    def estandar(self, filas, encabezados=None):
        workbook = Workbook()
        workbook.active.title = "Cuentas"
        workbook.active.append(encabezados or ["codigo", "nombre", "tipo", "codigo_padre", "acepta_movimientos", "activa"])
        for fila in filas:
            workbook.active.append(fila)
        return archivo_excel(workbook)

    def cuenta(self, codigo="1101", **kwargs):
        return CuentaContable.objects.create(
            empresa=self.empresa, codigo=codigo, nombre=kwargs.pop("nombre", "Caja"),
            tipo=kwargs.pop("tipo", "activo"), **kwargs,
        )

    def test_rumi_detecta_hoja_y_encabezado_tras_portada_sin_escribir(self):
        previa = analizar_catalogo_cuentas_desde_excel(self.empresa, archivo_excel(libro_rumi()))
        self.assertEqual(previa["hoja"], "Catalogo_RUMI")
        self.assertEqual(previa["fila_encabezado"], 6)
        self.assertEqual(previa["formato"], "rumi")
        self.assertEqual(previa["leidas"], 15)
        self.assertEqual(previa["omitidas"], 1)
        self.assertEqual(previa["controles"], 3)
        self.assertEqual(previa["condicionadas"], 1)
        self.assertEqual(previa["creadas"], 14)
        self.assertEqual(len(previa["firma_existentes"]), 64)
        json.dumps(previa)
        self.assertFalse(CuentaContable.objects.filter(empresa=self.empresa).exists())

    def test_rumi_clasifica_por_raiz_y_conserva_metadatos_sin_configurar_enlaces(self):
        resultado = importar_catalogo_cuentas_desde_excel(self.empresa, archivo_excel(libro_rumi()))
        self.assertEqual(resultado["creadas"], 14)
        caja = CuentaContable.objects.get(empresa=self.empresa, codigo="1101")
        self.assertEqual(caja.tipo, "activo")
        self.assertEqual(caja.cuenta_padre.codigo, "11")
        self.assertTrue(caja.acepta_movimientos)
        self.assertIn("Tipo funcional ERP: Caja", caja.descripcion)
        self.assertIn("Referencia CNBS: Referencia propuesta", caja.descripcion)
        self.assertIn("Fuente de referencia (URL): https://ejemplo.invalid/catalogo", caja.descripcion)
        self.assertFalse(CuentaContable.objects.get(empresa=self.empresa, codigo="11").acepta_movimientos)
        self.assertFalse(CuentaContable.objects.get(empresa=self.empresa, codigo="1102").activa)
        self.assertEqual(CuentaContable.objects.get(empresa=self.empresa, codigo="5").tipo, "ingreso")
        self.assertEqual(CuentaContable.objects.get(empresa=self.empresa, codigo="6").tipo, "gasto")
        for codigo in ("4", "41", "7"):
            control = CuentaContable.objects.get(empresa=self.empresa, codigo=codigo)
            self.assertEqual(control.tipo, "orden")
            self.assertFalse(control.activa)
            self.assertFalse(control.acepta_movimientos)
        self.assertFalse(CuentaContable.objects.filter(empresa=self.empresa, codigo="99").exists())
        self.assertFalse(ConfiguracionContableEmpresa.objects.filter(empresa=self.empresa).exists())
        self.assertFalse(CuentaFinanciera.objects.filter(empresa=self.empresa).exists())

    def test_rumi_puede_excluir_toda_la_rama_extracontable(self):
        previa = analizar_catalogo_cuentas_desde_excel(
            self.empresa, archivo_excel(libro_rumi()), incluir_controles=False,
        )
        self.assertEqual(previa["omitidas"], 4)
        self.assertEqual(previa["controles"], 3)
        self.assertEqual(previa["creadas"], 11)
        self.assertFalse(any(item["tipo"] == "orden" for item in previa["datos"]))

    def test_importacion_repetida_es_idempotente_aun_sin_actualizar(self):
        importar_catalogo_cuentas_desde_excel(self.empresa, archivo_excel(libro_rumi()))
        resultado = importar_catalogo_cuentas_desde_excel(
            self.empresa, archivo_excel(libro_rumi()), actualizar_existentes=False,
        )
        self.assertEqual(resultado["creadas"], 0)
        self.assertEqual(resultado["actualizadas"], 0)
        self.assertEqual(resultado["sin_cambios"], 14)
        self.assertEqual(CuentaContable.objects.filter(empresa=self.empresa).count(), 14)

    def test_prefijo_evitar_colision_preserva_codigo_original_y_jerarquia(self):
        vieja = self.cuenta("5", tipo="costo", nombre="Costos actuales")
        ConfiguracionContableEmpresa.objects.create(empresa=self.empresa, cuenta_entradas_inventario=vieja)
        previa = analizar_catalogo_cuentas_desde_excel(
            self.empresa, archivo_excel(libro_rumi()), prefijo_codigos="RUMI-",
        )
        resultado = aplicar_catalogo_cuentas(
            self.empresa, previa["datos"], firma_existentes=previa["firma_existentes"],
        )
        self.assertEqual(resultado["creadas"], 14)
        servicios = CuentaContable.objects.get(empresa=self.empresa, codigo="RUMI-51")
        self.assertEqual(servicios.cuenta_padre.codigo, "RUMI-5")
        self.assertEqual(servicios.tipo, "ingreso")
        self.assertIn("Código RUMI: 51", servicios.descripcion)
        vieja.refresh_from_db()
        self.assertEqual(vieja.tipo, "costo")
        self.assertEqual(vieja.nombre, "Costos actuales")

    def test_cuenta_usada_no_reclasifica_saldos_y_no_importa_parcialmente(self):
        usada = self.cuenta()
        asiento = AsientoContable.objects.create(empresa=self.empresa, fecha=date(2026, 10, 1), descripcion="Saldo anterior")
        LineaAsientoContable.objects.create(asiento=asiento, cuenta=usada, debe=Decimal("100.00"))
        archivo = self.estandar([
            ["999", "Cuenta nueva", "activo", "", "sí", "sí"],
            ["1101", "Cuenta reclasificada", "pasivo", "", "sí", "sí"],
        ])
        with self.assertRaisesMessage(ValidationError, "movimientos o enlaces operativos"):
            importar_catalogo_cuentas_desde_excel(self.empresa, archivo)
        usada.refresh_from_db()
        self.assertEqual(usada.tipo, "activo")
        self.assertFalse(CuentaContable.objects.filter(empresa=self.empresa, codigo="999").exists())

    def test_cuenta_con_configuracion_no_puede_desactivarse_por_excel(self):
        usada = self.cuenta()
        ConfiguracionContableEmpresa.objects.create(empresa=self.empresa, cuenta_caja=usada)
        with self.assertRaisesMessage(ValidationError, "movimientos o enlaces operativos"):
            analizar_catalogo_cuentas_desde_excel(
                self.empresa, self.estandar([["1101", "Caja", "activo", "", "sí", "no"]]),
            )
        usada.refresh_from_db()
        self.assertTrue(usada.activa)

    def test_padre_de_cuenta_con_movimientos_conserva_su_estructura(self):
        padre = self.cuenta("1", nombre="Activo", acepta_movimientos=False)
        usada = self.cuenta(cuenta_padre=padre)
        asiento = AsientoContable.objects.create(empresa=self.empresa, fecha=date(2026, 10, 1), descripcion="Auxiliar usado")
        LineaAsientoContable.objects.create(asiento=asiento, cuenta=usada, debe=Decimal("100.00"))
        with self.assertRaisesMessage(ValidationError, "movimientos o enlaces operativos"):
            analizar_catalogo_cuentas_desde_excel(
                self.empresa, self.estandar([["1", "Activo", "activo", "", "no", "no"]]),
            )
        padre.refresh_from_db()
        self.assertTrue(padre.activa)

    def test_cliente_proveedor_y_banco_protegen_estructura_de_cuentas(self):
        casos = ("cliente", "proveedor", "banco")
        for index, clase in enumerate(casos):
            with self.subTest(clase=clase):
                cuenta = self.cuenta(f"112{index}")
                if clase == "cliente":
                    Cliente.objects.create(empresa=self.empresa, nombre="Auxiliar", cuenta_contable=cuenta)
                elif clase == "proveedor":
                    Proveedor.objects.create(empresa=self.empresa, nombre="Auxiliar", cuenta_contable=cuenta)
                else:
                    CuentaFinanciera.objects.create(empresa=self.empresa, nombre="Banco", cuenta_contable=cuenta)
                with self.assertRaisesMessage(ValidationError, "movimientos o enlaces operativos"):
                    analizar_catalogo_cuentas_desde_excel(
                        self.empresa, self.estandar([[cuenta.codigo, "Auxiliar", "activo", "", "no", "sí"]]),
                    )

    def test_nombre_y_descripcion_pueden_actualizarse_en_cuenta_configurada(self):
        usada = self.cuenta(descripcion="Anterior")
        ConfiguracionContableEmpresa.objects.create(empresa=self.empresa, cuenta_caja=usada)
        resultado = importar_catalogo_cuentas_desde_excel(
            self.empresa, self.estandar(
                [["1101", "Caja operativa", "activo", "Descripción revisada"]],
                ["codigo", "nombre", "tipo", "descripcion"],
            ),
        )
        self.assertEqual(resultado["actualizadas"], 1)
        usada.refresh_from_db()
        self.assertEqual(usada.nombre, "Caja operativa")
        self.assertEqual(usada.descripcion, "Descripción revisada")
        self.assertTrue(usada.activa)

    def test_columnas_opcionales_ausentes_preservan_padre_y_estado(self):
        padre = self.cuenta("11", nombre="Activo corriente", acepta_movimientos=False)
        hija = self.cuenta(cuenta_padre=padre, activa=False, acepta_movimientos=False, descripcion="Se conserva")
        importar_catalogo_cuentas_desde_excel(
            self.empresa, self.estandar([["1101", "Nombre corregido", "activo"]], ["codigo", "nombre", "tipo"]),
        )
        hija.refresh_from_db()
        self.assertEqual(hija.cuenta_padre_id, padre.pk)
        self.assertFalse(hija.activa)
        self.assertFalse(hija.acepta_movimientos)
        self.assertEqual(hija.descripcion, "Se conserva")

    def test_codigos_de_otras_empresas_no_se_modifican_ni_sirven_de_padres(self):
        otra_cuenta = CuentaContable.objects.create(empresa=self.otra, codigo="1101", nombre="Otra Caja", tipo="activo")
        importar_catalogo_cuentas_desde_excel(self.empresa, self.estandar([["1101", "Caja A", "activo", "", "sí", "sí"]]))
        otra_cuenta.refresh_from_db()
        self.assertEqual(otra_cuenta.nombre, "Otra Caja")
        CuentaContable.objects.create(empresa=self.otra, codigo="999", nombre="Padre ajeno", tipo="activo")
        with self.assertRaisesMessage(ValidationError, "no existe en el archivo ni en esta empresa"):
            analizar_catalogo_cuentas_desde_excel(self.empresa, self.estandar([["9991", "Hija", "activo", "999", "sí", "sí"]]))

    def test_firma_impide_aplicar_catalogo_cambiado_despues_de_vista_previa(self):
        cuenta = self.cuenta()
        previa = analizar_catalogo_cuentas_desde_excel(
            self.empresa, self.estandar([["1101", "Nuevo nombre", "activo", "", "sí", "sí"]]),
        )
        cuenta.nombre = "Cambio concurrente"
        cuenta.save()
        with self.assertRaisesMessage(ValidationError, "cambiaron después de la vista previa"):
            aplicar_catalogo_cuentas(self.empresa, previa["datos"], firma_existentes=previa["firma_existentes"])
        cuenta.refresh_from_db()
        self.assertEqual(cuenta.nombre, "Cambio concurrente")

    def test_nuevos_enlaces_invalidan_firma_y_revalidan_guardias_sin_firma(self):
        cuenta = self.cuenta()
        previa = analizar_catalogo_cuentas_desde_excel(
            self.empresa, self.estandar([["1101", "Caja", "pasivo", "", "sí", "sí"]]),
        )
        Cliente.objects.create(empresa=self.empresa, nombre="Cliente posterior", cuenta_contable=cuenta)
        with self.assertRaisesMessage(ValidationError, "cambiaron después de la vista previa"):
            aplicar_catalogo_cuentas(self.empresa, previa["datos"], firma_existentes=previa["firma_existentes"])
        with self.assertRaisesMessage(ValidationError, "movimientos o enlaces operativos"):
            aplicar_catalogo_cuentas(self.empresa, previa["datos"])
        cuenta.refresh_from_db()
        self.assertEqual(cuenta.tipo, "activo")

    def test_numeros_integrales_y_formato_ceros_iniciales_se_conservan(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["codigo", "nombre", "tipo", "codigo_padre"])
        sheet.append([12, "Padre", "activo", ""])
        sheet["A2"].number_format = "0000"
        sheet.append([123.0, "Hija", "activo", 12])
        sheet["A3"].number_format = "00000"
        sheet["D3"].number_format = "0000"
        previa = analizar_catalogo_cuentas_desde_excel(self.empresa, archivo_excel(workbook))
        self.assertEqual(previa["datos"][0]["codigo"], "0012")
        self.assertEqual(previa["datos"][1]["codigo"], "00123")
        self.assertEqual(previa["datos"][1]["codigo_padre"], "0012")

    def test_formula_en_campos_clave_se_rechaza_con_hoja_y_fila(self):
        workbook = Workbook()
        workbook.active.title = "Datos"
        workbook.active.append(["codigo", "nombre", "tipo"])
        workbook.active.append(['="1101"', "Caja", "activo"])
        with self.assertRaises(ValidationError) as error:
            analizar_catalogo_cuentas_desde_excel(self.empresa, archivo_excel(workbook))
        mensaje = " ".join(error.exception.messages)
        self.assertIn("Datos", mensaje)
        self.assertIn("fila 2", mensaje)
        self.assertIn("fórmula", mensaje)

    def test_booleano_invalido_no_se_convierte_en_verdadero(self):
        with self.assertRaises(ValidationError) as error:
            analizar_catalogo_cuentas_desde_excel(
                self.empresa, self.estandar([["1101", "Caja", "activo", "", "quizás", "sí"]]),
            )
        self.assertIn("fila 2", " ".join(error.exception.messages))
        self.assertIn("acepta_movimientos", " ".join(error.exception.messages))
        self.assertFalse(CuentaContable.objects.filter(empresa=self.empresa).exists())

    def test_varias_hojas_exigen_seleccion_explicita(self):
        workbook = Workbook()
        workbook.active.title = "Uno"
        segundo = workbook.create_sheet("Dos")
        for sheet in (workbook.active, segundo):
            sheet.append(["codigo", "nombre", "tipo"])
            sheet.append(["1", "Activo", "activo"])
        with self.assertRaisesMessage(ValidationError, "varios catálogos"):
            analizar_catalogo_cuentas_desde_excel(self.empresa, archivo_excel(workbook))
        previa = analizar_catalogo_cuentas_desde_excel(self.empresa, archivo_excel(workbook), hoja="Dos")
        self.assertEqual(previa["hoja"], "Dos")

    def test_excel_corrupto_devuelve_error_legible(self):
        with self.assertRaisesMessage(ValidationError, "Excel válido"):
            analizar_catalogo_cuentas_desde_excel(self.empresa, BytesIO(b"no es un xlsx"))

    def test_ciclo_y_codigo_duplicado_rechazan_toda_la_propuesta(self):
        casos = (
            [["1", "Uno", "activo", "2"], ["2", "Dos", "activo", "1"]],
            [["1", "Uno", "activo", ""], ["1", "Duplicado", "activo", ""]],
            [["1", "Uno", "activo", "1"]],
        )
        for filas in casos:
            with self.subTest(filas=filas):
                with self.assertRaises(ValidationError):
                    importar_catalogo_cuentas_desde_excel(
                        self.empresa, self.estandar(filas, ["codigo", "nombre", "tipo", "codigo_padre"]),
                    )
                self.assertFalse(CuentaContable.objects.filter(empresa=self.empresa).exists())

    def test_longitudes_no_guardan_filas_validas_antes_de_error(self):
        for codigo, nombre in (("1" * 31, "Largo"), ("1", "N" * 201)):
            with self.subTest(codigo=codigo):
                with self.assertRaises(ValidationError):
                    importar_catalogo_cuentas_desde_excel(
                        self.empresa, self.estandar([["2", "Válida", "activo"], [codigo, nombre, "activo"]],
                                                   ["codigo", "nombre", "tipo"]),
                    )
                self.assertFalse(CuentaContable.objects.filter(empresa=self.empresa).exists())

    def test_aplicar_no_admite_campos_de_otra_empresa(self):
        with self.assertRaisesMessage(ValidationError, "campos no admitidos"):
            aplicar_catalogo_cuentas(
                self.empresa, [{"codigo": "1", "nombre": "Activo", "tipo": "activo", "empresa_id": self.otra.pk}],
            )
        self.assertFalse(CuentaContable.objects.filter(empresa=self.empresa).exists())

    def test_rumi_funcion_operativa_no_sustituye_tipo_y_nombre_raiz_no_se_inventa(self):
        workbook = libro_rumi()
        sheet = workbook["Catalogo_RUMI"]
        sheet["B7"] = "Clase desconocida"
        with self.assertRaisesMessage(ValidationError, "clase raíz"):
            analizar_catalogo_cuentas_desde_excel(self.empresa, archivo_excel(workbook))

    def test_rumi_padre_no_cargar_no_genera_padres_artificiales(self):
        workbook = libro_rumi()
        workbook["Catalogo_RUMI"].append(fila_rumi("991", "Hija cargable", 2, "99", "Sí", "Base propuesta"))
        with self.assertRaises(ValidationError):
            analizar_catalogo_cuentas_desde_excel(self.empresa, archivo_excel(workbook))
        self.assertFalse(CuentaContable.objects.filter(empresa=self.empresa).exists())

    def test_estandar_titulos_arriba_detecta_encabezados_y_valida_tipo_padre(self):
        workbook = Workbook()
        workbook.active.append(["Título"])
        workbook.active.append([])
        workbook.active.append(["codigo", "nombre", "tipo", "codigo_padre"])
        workbook.active.append(["1", "Activo", "activo", ""])
        workbook.active.append(["2", "Pasivo", "pasivo", "1"])
        with self.assertRaisesMessage(ValidationError, "tipo debe coincidir"):
            analizar_catalogo_cuentas_desde_excel(self.empresa, archivo_excel(workbook))

    def test_legacy_h_i_mantiene_clasificacion_anterior_y_padres(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([1, None, None, None, None, None, None, "1", "ACTIVOS"])
        sheet.append([1, 1, None, None, None, None, None, "11", "ACTIVOS CORRIENTES"])
        sheet.append([1, 1, "01", None, None, None, None, "1101", "CAJA"])
        sheet.append([4, None, None, None, None, None, None, "4", "INGRESOS"])
        previa = analizar_catalogo_cuentas_desde_excel(self.empresa, archivo_excel(workbook))
        self.assertEqual(previa["formato"], "jerarquico")
        self.assertEqual(previa["datos"][2]["codigo_padre"], "11")
        self.assertEqual(previa["datos"][3]["tipo"], "ingreso")
        self.assertFalse(previa["datos"][0]["acepta_movimientos"])
