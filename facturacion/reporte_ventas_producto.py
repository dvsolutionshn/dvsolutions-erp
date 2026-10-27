"""Consulta de líneas de un producto, siempre dentro de una sola empresa."""

from decimal import Decimal

from django import forms
from django.conf import settings
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.template.loader import render_to_string
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_safe

from contabilidad.moneda import convertir_a_lempiras, redondear_importe
from core.clinical_permissions import EMPRESAS_CLINICAS_CON_ROLES
from core.models import Empresa, Usuario
from .models import Factura, LineaFactura, Producto


class VendedorField(forms.ModelChoiceField):
    def label_from_instance(self, usuario):
        nombre = usuario.get_full_name().strip()
        return f"{nombre} ({usuario.username})" if nombre else usuario.username


class VentasProductoForm(forms.Form):
    producto = forms.ModelChoiceField(
        label="Producto", queryset=Producto.objects.none(), widget=forms.HiddenInput,
        error_messages={"required": "Selecciona un producto de la lista.",
                        "invalid_choice": "Selecciona un producto de esta empresa."},
    )
    fecha_desde = forms.DateField(
        label="Fecha inicial", input_formats=["%Y-%m-%d", "%d/%m/%Y"],
        widget=forms.DateInput(format="%Y-%m-%d", attrs={"type": "date"}),
    )
    fecha_hasta = forms.DateField(
        label="Fecha final", input_formats=["%Y-%m-%d", "%d/%m/%Y"],
        widget=forms.DateInput(format="%Y-%m-%d", attrs={"type": "date"}),
    )
    vendedor = VendedorField(
        label="Usuario / vendedor", queryset=Usuario.objects.none(),
        required=False, empty_label="Todos los vendedores",
    )
    estado = forms.ChoiceField(
        label="Estado de factura", required=False, initial="emitida",
        choices=[("emitida", "Emitida"), ("", "Todos los estados"),
                 ("borrador", "Borrador"), ("anulada", "Anulada (auditoría)")],
    )

    def __init__(self, *args, empresa, **kwargs):
        super().__init__(*args, **kwargs)
        # Inactivos y eliminados también conservan ventas históricas consultables.
        self.fields["producto"].queryset = Producto.objects.filter(empresa=empresa)
        self.fields["vendedor"].queryset = Usuario.objects.filter(
            pk__in=Factura.objects.filter(empresa=empresa).values("vendedor_id"),
        ).order_by("first_name", "last_name", "username")

    def clean(self):
        datos = super().clean()
        if hasattr(self.data, "getlist") and any(
            len(self.data.getlist(campo)) > 1 for campo in self.fields
        ):
            raise ValidationError("Cada filtro admite un solo valor.")
        desde, hasta = datos.get("fecha_desde"), datos.get("fecha_hasta")
        if desde and hasta and desde > hasta:
            self.add_error("fecha_hasta", "La fecha final debe ser igual o posterior a la inicial.")
        return datos


def _empresa_autorizada(request, empresa_slug, *, exportar=False):
    if empresa_slug not in EMPRESAS_CLINICAS_CON_ROLES:
        raise Http404
    empresa = get_object_or_404(Empresa, slug=empresa_slug, activa=True)
    usuario = request.user
    if not usuario.puede_acceder_empresa(empresa):
        raise PermissionDenied("No tienes acceso a esta empresa.")
    if not usuario.is_superuser and not empresa.licencia_operativa:
        raise PermissionDenied("La licencia de esta empresa no está operativa.")
    if not empresa.tiene_modulo_activo("facturacion"):
        raise PermissionDenied("Facturación no está habilitada para esta empresa.")
    if not usuario.tiene_permiso_erp("puede_reportes", empresa):
        raise PermissionDenied("No tienes permiso para consultar reportes de esta empresa.")
    if exportar and not usuario.tiene_permiso_erp("puede_exportar_reportes", empresa):
        raise PermissionDenied("No tienes permiso para exportar reportes de esta empresa.")
    return empresa


@login_required
@require_safe
@never_cache
def buscar_productos_ventas(request, empresa_slug):
    empresa = _empresa_autorizada(request, empresa_slug)
    consulta = (request.GET.get("q") or "").strip()[:200]
    productos = Producto.objects.filter(empresa=empresa)
    if consulta:
        productos = productos.filter(
            Q(nombre__icontains=consulta) | Q(codigo__icontains=consulta)
            | Q(descripcion__icontains=consulta)
        )
    pagina = Paginator(productos.order_by("nombre", "pk"), 30).get_page(request.GET.get("pagina"))
    resultados = []
    for producto in pagina:
        texto = producto.nombre
        if producto.codigo:
            texto += f" · {producto.codigo}"
        if producto.eliminado:
            texto += " (eliminado; histórico)"
        elif not producto.activo:
            texto += " (inactivo)"
        resultados.append({
            "id": producto.pk, "text": texto, "nombre": producto.nombre,
            "codigo": producto.codigo or "", "referencia": producto.codigo or "",
            "activo": producto.activo and not producto.eliminado,
        })
    return JsonResponse({"results": resultados, "has_more": pagina.has_next(),
                         "next_page": pagina.next_page_number() if pagina.has_next() else None})


def _lineas_reporte(empresa, filtros):
    lineas = LineaFactura.objects.filter(
        factura__empresa=empresa, producto=filtros["producto"],
        factura__fecha_emision__range=(filtros["fecha_desde"], filtros["fecha_hasta"]),
    )
    if filtros["vendedor"]:
        lineas = lineas.filter(factura__vendedor=filtros["vendedor"])
    if filtros["estado"]:
        lineas = lineas.filter(factura__estado=filtros["estado"])
    return lineas.select_related("factura__empresa", "factura__cliente", "factura__vendedor").order_by(
        "factura__fecha_emision", "factura_id", "pk",
    )


def _resumen(lineas):
    cantidad, total, facturas = Decimal("0.00"), Decimal("0.00"), set()
    for linea in lineas.exclude(factura__estado="anulada").values(
        "factura_id", "cantidad", "subtotal", "impuesto_monto",
        "factura__moneda", "factura__tipo_cambio",
    ).iterator(chunk_size=2000):
        cantidad += linea["cantidad"]
        total += convertir_a_lempiras(
            linea["subtotal"] + linea["impuesto_monto"],
            linea["factura__moneda"], linea["factura__tipo_cambio"],
        )
        facturas.add(linea["factura_id"])
    return {"cantidad": cantidad, "total": total, "documentos": len(facturas),
            "precio_promedio": redondear_importe(total / cantidad) if cantidad else None}


def _fila(linea, puede_ver_facturas):
    factura = linea.factura
    def hnl(importe):
        return convertir_a_lempiras(importe, factura.moneda, factura.tipo_cambio)
    vendedor = factura.vendedor
    return {
        "fecha": factura.fecha_emision,
        "numero_factura": factura.numero_factura or (
            f"Borrador #{factura.pk}" if factura.estado == "borrador" else f"Sin número (#{factura.pk})"
        ),
        "cliente": factura.cliente.nombre, "cantidad": linea.cantidad,
        "precio_unitario": hnl(linea.precio_unitario),
        "descuento": hnl(linea.descuento_monto),
        "total": hnl(linea.subtotal + linea.impuesto_monto),
        "vendedor": (vendedor.get_full_name().strip() or vendedor.username) if vendedor else "Sin vendedor asignado",
        "estado": factura.get_estado_display(), "anulada": factura.estado == "anulada",
        "precio_incluye_impuesto": linea.precio_incluye_impuesto,
        "factura_url": reverse("ver_factura", args=[factura.empresa.slug, factura.pk]) if puede_ver_facturas else "",
    }


def _url_filtros(request, parametros, **extra):
    consulta = parametros.copy()
    for clave, valor in extra.items():
        consulta[clave] = valor
    return f"{request.path}?{consulta.urlencode()}"


def _excel(contexto, filas):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Resumen"
    filtros, producto, resumen = contexto["filtros"], contexto["producto_seleccionado"], contexto["resumen"]
    for fila in [
        ["Ventas por Producto", contexto["empresa"].nombre],
        ["Producto", producto.nombre], ["Código / referencia", producto.codigo or ""],
        ["Fecha inicial", filtros["fecha_desde"]], ["Fecha final", filtros["fecha_hasta"]],
        ["Usuario / vendedor", str(contexto["form"].fields["vendedor"].label_from_instance(filtros["vendedor"])) if filtros["vendedor"] else "Todos"],
        ["Estado", dict(contexto["form"].fields["estado"].choices)[filtros["estado"]]],
        ["Unidades (excluye anuladas)", resumen["cantidad"]],
        ["Total HNL (excluye anuladas)", resumen["total"]],
        ["Número de facturas (excluye anuladas)", resumen["documentos"]],
        ["Precio promedio HNL, con impuesto y descuento", resumen["precio_promedio"]],
        ["Criterio", "Líneas del producto seleccionado. Total con impuestos y descuento aplicado; notas de crédito no descontadas."],
        ["Moneda", "HNL; USD convertido al tipo de cambio guardado en cada factura."],
        ["Borradores", "Se incluyen únicamente al elegir Borrador o Todos los estados; no son ventas definitivas."],
    ]:
        hoja.append(fila)
    hoja.column_dimensions["A"].width = 52
    hoja.column_dimensions["B"].width = 85
    detalle = libro.create_sheet("Detalle")
    detalle.append(["Fecha", "Factura", "Cliente / paciente", "Cantidad", "Precio unitario HNL",
                    "Descuento HNL", "Total producto HNL", "Usuario / vendedor", "Estado", "Precio incluye impuesto"])
    for fila in filas:
        detalle.append([fila["fecha"], fila["numero_factura"], fila["cliente"], fila["cantidad"],
                        fila["precio_unitario"], fila["descuento"], fila["total"], fila["vendedor"],
                        "ANULADA" if fila["anulada"] else fila["estado"],
                        "Sí" if fila["precio_incluye_impuesto"] else "No"])
    detalle.freeze_panes = "A2"
    detalle.auto_filter.ref = detalle.dimensions
    for columna, ancho in {"A": 14, "B": 26, "C": 40, "D": 16, "E": 23,
                           "F": 20, "G": 24, "H": 32, "I": 18, "J": 25}.items():
        detalle.column_dimensions[columna].width = ancho
    for hoja_actual in libro:
        for celda in hoja_actual[1]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor="174C70")
        for fila in hoja_actual:
            for celda in fila:
                # Un nombre/código de catálogo nunca se interpreta como fórmula.
                if isinstance(celda.value, str):
                    celda.data_type = "s"
                elif isinstance(celda.value, Decimal):
                    celda.number_format = "#,##0.00"
                elif hasattr(celda.value, "strftime"):
                    celda.number_format = "dd/mm/yyyy"
                celda.alignment = Alignment(vertical="top", wrap_text=True)
    respuesta = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    respuesta["Content-Disposition"] = 'attachment; filename="Ventas_por_Producto.xlsx"'
    libro.save(respuesta)
    return respuesta


def _pdf(contexto, filas):
    # Reutiliza el motor y la resolución de recursos de los PDF de Facturación.
    from weasyprint import HTML
    html = render_to_string("facturacion/reporte_ventas_producto_pdf.html", {**contexto, "filas": filas})
    respuesta = HttpResponse(
        HTML(string=html, base_url=str(settings.BASE_DIR)).write_pdf(), content_type="application/pdf",
    )
    respuesta["Content-Disposition"] = 'attachment; filename="Ventas_por_Producto.pdf"'
    return respuesta


@login_required
@require_safe
@never_cache
def reporte_ventas_producto(request, empresa_slug):
    formato = request.GET.get("formato", "")
    if formato not in {"", "excel", "pdf"}:
        raise Http404
    empresa = _empresa_autorizada(request, empresa_slug, exportar=bool(formato))
    hoy = timezone.localdate()
    parametros = request.GET.copy()
    for clave in ("pagina", "formato"):
        parametros.pop(clave, None)
    if parametros and "estado" not in parametros:
        parametros["estado"] = "emitida"
    form = VentasProductoForm(
        parametros if (parametros or formato) else None, empresa=empresa,
        initial={"fecha_desde": hoy.replace(day=1), "fecha_hasta": hoy, "estado": "emitida"},
    )
    contexto = {
        "empresa": empresa, "form": form, "producto_seleccionado": None,
        "reporte_generado": False, "pagina_obj": None,
        "puede_exportar": request.user.tiene_permiso_erp("puede_exportar_reportes", empresa),
        "search_productos_url": reverse("buscar_productos_ventas", args=[empresa.slug]),
    }
    if form.is_bound and form.is_valid():
        filtros = form.cleaned_data
        lineas = _lineas_reporte(empresa, filtros)
        contexto.update({"filtros": filtros, "producto_seleccionado": filtros["producto"],
                         "resumen": _resumen(lineas), "reporte_generado": True})
        ver_facturas = request.user.tiene_permiso_erp("puede_ver_facturas", empresa)
        if formato:
            filas = (_fila(linea, ver_facturas) for linea in lineas.iterator(chunk_size=1000))
            return _excel(contexto, filas) if formato == "excel" else _pdf(contexto, filas)
        pagina = Paginator(lineas, 50).get_page(request.GET.get("pagina"))
        pagina.object_list = [_fila(linea, ver_facturas) for linea in pagina.object_list]
        contexto.update({
            "pagina_obj": pagina, "filtros_query": parametros.urlencode(),
            "export_excel_url": _url_filtros(request, parametros, formato="excel"),
            "export_pdf_url": _url_filtros(request, parametros, formato="pdf"),
            "pagina_anterior_url": _url_filtros(request, parametros, pagina=pagina.previous_page_number()) if pagina.has_previous() else "",
            "pagina_siguiente_url": _url_filtros(request, parametros, pagina=pagina.next_page_number()) if pagina.has_next() else "",
        })
    elif form.is_bound:
        # Conserva una selección válida al corregir otra parte del formulario.
        contexto["producto_seleccionado"] = form.cleaned_data.get("producto")
    if formato:
        form.add_error(None, "Selecciona un producto y un rango de fechas válidos antes de exportar.")
    return render(request, "facturacion/reporte_ventas_producto.html", contexto, status=400 if formato else 200)
