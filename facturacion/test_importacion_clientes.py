from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import json
from unittest.mock import patch
from unittest import skipUnless
import os
from pathlib import Path
import subprocess
import tempfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, Client
from django.contrib.staticfiles.testing import StaticLiveServerTestCase
from django.urls import reverse
from openpyxl import Workbook

from . import test_clientes_contables as fixtures
from .importacion_clientes import leer_excel
from .models import Proveedor, RegistroCompraFiscal, LibroCompraMensual


def excel(filas=None, formula=False, mixto=False, total_general=None):
    libro=Workbook(); hoja=libro.active; hoja.title='COMPRAS'
    hoja['D5']='NORDIC SPS'; hoja['B6']='Libro Enero 2026'
    for col,valor in enumerate(['Fecha','beneficiario','No. De factura','Subtotal','ISV 15%',0.18,'Total'],2):
        hoja.cell(10,col,valor)
    if filas is None:
        filas=[(datetime(2025,12,25),'Proveedor histórico','6368',100,15,0,115),
               (datetime(2026,1,2),'Proveedor histórico',None,50,7.5,0,57.5)]
    for numero,fila in enumerate(filas,11):
        for col,valor in enumerate(fila,2): hoja.cell(numero,col,valor)
    if formula: hoja['F11']='=E11*0.15'
    if mixto: hoja['G11']=18
    n=len(filas)+11;hoja.cell(n,2,'TOTAL');hoja.cell(n,8,total_general if total_general is not None else sum(f[-1] or 0 for f in filas))
    destino=BytesIO();libro.save(destino);return destino.getvalue()


class ImportacionClientesTests(TestCase):
    @classmethod
    def setUpTestData(cls): fixtures.ClientesContablesTests.setUpTestData.__func__(cls)

    def setUp(self):
        self.client.force_login(self.usuario)
        self.url=reverse('importar_compras_cliente',args=[self.empresa.slug,self.nordic.pk,2026,1])

    def previo(self, contenido=None, url=None):
        archivo=SimpleUploadedFile('Enero.xlsx',contenido or excel(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response=self.client.post(url or self.url,{'archivo':archivo})
        self.assertEqual(response.status_code,200,response.content)
        self.assertIsNone(response.context['error'],response.context['error'])
        self.assertTrue(response.context['token'])
        return response

    def confirmar(self, previo, **datos):
        return self.client.post(self.url,{'token':previo.context['token'],'accion':'confirmar','revisado':'on',
            'filas':[str(f['fila']) for f in previo.context['filas']],**datos})

    def test_previo_sin_escritura_importa_y_reenvio_no_duplica(self):
        previo=self.previo()
        self.assertEqual(previo.context['total'],Decimal('172.50'))
        self.assertContains(previo,'Sin número')
        self.assertFalse(RegistroCompraFiscal.objects.exists())
        self.assertEqual(Proveedor.objects.count(),3)
        self.assertEqual(self.confirmar(previo).status_code,302)
        self.assertEqual(RegistroCompraFiscal.objects.count(),2)
        proveedor=Proveedor.objects.get(nombre='Proveedor histórico')
        self.assertEqual(proveedor.cliente_contable,self.nordic)
        self.assertFalse(proveedor.rtn)
        compra=RegistroCompraFiscal.objects.get(numero_factura='6368')
        self.assertEqual(compra.fecha_documento,date(2025,12,25))
        self.assertEqual((compra.periodo_anio,compra.periodo_mes),(2026,1))
        self.assertEqual(compra.creado_por,self.usuario)
        sin_numero=RegistroCompraFiscal.objects.get(numero_factura='')
        self.assertIsNone(sin_numero.numero_factura_normalizado)
        self.assertIsNone(sin_numero.identidad_captura)
        self.assertIn('Sin número',sin_numero.observacion)
        self.assertEqual(LibroCompraMensual.objects.get().registros.count(),2)
        response=self.confirmar(previo)
        self.assertEqual(response.status_code,200)
        self.assertEqual(RegistroCompraFiscal.objects.count(),2)
        self.assertContains(response,'ya fue importada')

    def test_excluir_incompleta_corregir_fecha_y_confirmacion_obligatoria(self):
        contenido=excel([(datetime(2026,12,25),'Histórico','001',100,15,0,115),(None,None,None,5,.75,0,5.75)])
        previo=self.previo(contenido)
        self.assertContains(previo,'Fecha posterior')
        response=self.confirmar(previo)
        self.assertContains(response,'Corrige o desmarca')
        self.assertFalse(RegistroCompraFiscal.objects.exists())
        response=self.confirmar(previo,filas=['11'],fecha_11='25/12/25',revisado='')
        self.assertContains(response,'Confirma que revisaste')
        response=self.confirmar(previo,filas=['11'],fecha_11='25/12/25')
        self.assertEqual(response.status_code,302,response.content)
        compra=RegistroCompraFiscal.objects.get()
        self.assertEqual(compra.fecha_documento,date(2025,12,25))
        self.assertIn('25/12/2026',compra.observacion)

    def test_duplicados_en_historial_y_entre_clientes(self):
        compra=RegistroCompraFiscal.objects.create(empresa=self.empresa,cliente_contable=self.nordic,proveedor=self.proveedor,
            proveedor_nombre=self.proveedor.nombre,proveedor_rtn=self.proveedor.rtn,numero_factura='000-001-01-12345678',fecha_documento=date(2025,6,1),total=115)
        previo=self.previo(excel([(datetime(2026,1,1),self.proveedor.nombre,'0000010112345678',100,15,0,115)]))
        self.assertTrue(previo.context['filas'][0]['duplicada'])
        self.confirmar(previo)
        self.assertEqual(RegistroCompraFiscal.objects.count(),1)
        self.client.force_login(self.admin)
        otra=reverse('importar_compras_cliente',args=[self.empresa.slug,self.molac.pk,2026,1])
        previo=self.previo(excel([(datetime(2026,1,1),self.proveedor.nombre,'0000010112345678',100,15,0,115)]),url=otra)
        response=self.client.post(otra,{'token':previo.context['token'],'accion':'confirmar','revisado':'on','filas':['11']})
        self.assertEqual(response.status_code,302)
        compra.refresh_from_db();self.assertEqual(compra.total,Decimal('115'))
        self.assertEqual(RegistroCompraFiscal.objects.count(),2)

    def test_permisos_tokens_y_csrf(self):
        previo=self.previo()
        otra=reverse('importar_compras_cliente',args=[self.empresa.slug,self.molac.pk,2026,1])
        self.assertEqual(self.client.get(otra).status_code,404)
        self.client.force_login(self.admin)
        self.assertEqual(self.client.post(otra,{'token':previo.context['token']}).status_code,403)
        self.assertEqual(self.client.post(self.url,{'token':previo.context['token']}).status_code,403)
        self.client.force_login(self.usuario)
        self.assertContains(self.client.post(self.url,{'token':'alterado'}),'no es válida')
        self.rol.puede_crear_compras=False;self.rol.save()
        self.assertEqual(self.client.post(self.url,{'token':previo.context['token']}).status_code,403)
        self.rol.puede_crear_compras=True;self.rol.save()
        csrf=Client(enforce_csrf_checks=True);csrf.force_login(self.usuario)
        self.assertIn(csrf.post(self.url,{'token':previo.context['token']}).status_code,(302,403))
        self.assertFalse(RegistroCompraFiscal.objects.exists())

    def test_cierre_y_proveedor_sin_permiso_revocados_despues_de_previo(self):
        previo=self.previo()
        self.rol.puede_crear_proveedores=False;self.rol.save()
        self.confirmar(previo)
        self.assertFalse(RegistroCompraFiscal.objects.exists())
        self.rol.puede_crear_proveedores=True;self.rol.save()
        LibroCompraMensual.objects.create(empresa=self.empresa,cliente_contable=self.nordic,anio=2026,mes=1,estado='finalizado')
        self.assertContains(self.confirmar(previo),'libro está finalizado')
        self.assertFalse(RegistroCompraFiscal.objects.exists())
        self.assertFalse(Proveedor.objects.filter(nombre='Proveedor histórico').exists())

    def test_archivo_formula_sin_cache_montos_invalidos_y_base_mixta(self):
        for contenido, texto in [(excel(formula=True),'fórmula sin valor guardado'),(excel(mixto=True),'ambos impuestos'),
                                 (excel([(datetime(2026,1,1),'Proveedor','1',-1,0,0,1)]),'Importe inválido')]:
            with self.subTest(texto=texto):
                previo=self.previo(contenido)
                self.assertContains(previo,texto)
                self.confirmar(previo)
                self.assertFalse(RegistroCompraFiscal.objects.exists())
        response=self.client.post(self.url,{'archivo':SimpleUploadedFile('malo.xlsx',b'no excel')})
        self.assertContains(response,'No se pudo leer')

    def test_redondeo_y_campos_monetarios_no_manipulables(self):
        previo=self.previo(excel([(datetime(2026,1,1),'Proveedor','6368',1.03,.1545,0,1.1845)]))
        self.assertEqual(previo.context['filas'][0]['total'],'1.18')
        response=self.confirmar(previo,total_11='999',base_15_11='999')
        self.assertEqual(response.status_code,302)
        self.assertEqual(RegistroCompraFiscal.objects.get().total,Decimal('1.18'))

    def test_ediciones_json_y_persistencia_en_acumulado(self):
        previo=self.previo()
        datos={'seleccionadas':['11'],'fecha_11':'5/1/26','proveedor_11':'Proveedor corregido','numero_11':'44','rtn_11':''}
        response=self.confirmar(previo,ediciones=json.dumps(datos))
        self.assertEqual(response.status_code,302)
        compra=RegistroCompraFiscal.objects.get()
        self.assertEqual(compra.numero_factura,'44')
        self.assertEqual(compra.proveedor_nombre,'Proveedor corregido')
        anual=self.client.get(reverse('acumulado_cliente_contable',args=[self.empresa.slug,self.nordic.pk]),{'anio':2026})
        self.assertEqual(anual.context['resumen']['total'],Decimal('115'))

    def test_documentos_sin_numero_se_pueden_editar_sin_colision(self):
        previo=self.previo(excel([(datetime(2026,1,1),'Proveedor',None,100,15,0,115),(datetime(2026,1,2),'Proveedor',None,50,7.5,0,57.5)]))
        self.assertEqual(self.confirmar(previo).status_code,302)
        url=reverse('captura_cliente_contable',args=[self.empresa.slug,self.nordic.pk,2026,1])
        fila=self.client.get(url,{'accion':'cuadro'}).json()['registros'][0]
        response=self.client.post(url,{'accion':'editar','registro_id':fila['id'],'version':fila['version'],'proveedor':fila['proveedor_id'],
            'numero_factura':'','fecha_documento':'1/1/26','base_15':'200'})
        self.assertEqual(response.status_code,200,response.content)
        self.assertEqual(RegistroCompraFiscal.objects.count(),2)

    def test_error_segunda_compra_revierte_todo_y_expiracion(self):
        previo=self.previo()
        real=RegistroCompraFiscal.save
        from django.core.exceptions import ValidationError
        def guardar(instancia,*args,**kwargs):
            if not instancia.numero_factura: raise ValidationError('Error simulado')
            return real(instancia,*args,**kwargs)
        with patch.object(RegistroCompraFiscal,'save',guardar):
            self.assertContains(self.confirmar(previo),'Error simulado')
        self.assertFalse(RegistroCompraFiscal.objects.exists())
        self.assertFalse(Proveedor.objects.filter(nombre='Proveedor histórico').exists())
        self.assertFalse(LibroCompraMensual.objects.exists())
        import time
        with patch('django.core.signing.time.time',return_value=time.time()+4000):
            self.assertContains(self.confirmar(previo),'venció')

    @skipUnless(os.environ.get('IMPORT_EXCEL_SAMPLE'), 'Archivo privado de referencia opcional.')
    def test_excel_referencia_en_base_temporal(self):
        previo=self.previo(Path(os.environ['IMPORT_EXCEL_SAMPLE']).read_bytes())
        self.assertEqual(len(previo.context['filas']),103)
        self.assertIsNotNone(previo.context['lote']['total_excel'])
        self.assertEqual([f['fila'] for f in previo.context['filas'] if f['errores']],[86])
        self.assertFalse(RegistroCompraFiscal.objects.exists())
        response=self.confirmar(previo,filas=[str(f['fila']) for f in previo.context['filas'] if not f['errores']])
        self.assertEqual(response.status_code,302,response.content)
        self.assertEqual(RegistroCompraFiscal.objects.count(),102)
        self.assertEqual(RegistroCompraFiscal.objects.filter(numero_factura='').count(),1)
        self.assertEqual(RegistroCompraFiscal.objects.filter(fecha_documento__month=12,fecha_documento__year=2026).count(),4)
        self.assertEqual(RegistroCompraFiscal.objects.filter(periodo_anio=2026,periodo_mes=1).count(),102)


@skipUnless(os.environ.get('PLAYWRIGHT_MODULE'), 'Configura PLAYWRIGHT_MODULE para probar navegador real.')
class ImportacionClientesBrowserTests(StaticLiveServerTestCase):
    def test_importacion_y_edicion_en_navegador(self):
        fixtures.ClientesContablesTests.setUpTestData.__func__(self.__class__)
        self.client.force_login(self.usuario)
        with tempfile.TemporaryDirectory(prefix='compras-excel-test-') as temporal:
            archivo=Path(temporal)/'Enero.xlsx'
            archivo.write_bytes(excel([(datetime(2026,12,25),'Proveedor histórico','6368',100,15,0,115),
                (datetime(2026,1,2),'Proveedor histórico',None,50,7.5,0,57.5),
                (datetime(2026,1,3),self.proveedor.nombre,'0000010112345678',100,15,0,115),
                (None,None,None,5,.75,0,5.75)]))
            diagnostico=Path(temporal)/'Diagnostico.xlsx'
            diagnostico.write_bytes(excel([(datetime(2026,1,5),'Proveedor prueba','888',1000,150,0,1149.99),
                (datetime(2026,1,6),'Proveedor prueba','889',-1,0,0,2)],total_general=1152))
            env={**os.environ,'IMPORT_DIAGNOSTIC_FILE':str(diagnostico),'IMPORT_TEST_URL':self.live_server_url+reverse('captura_cliente_contable',args=[self.empresa.slug,self.nordic.pk,2026,1]),
                 'IMPORT_TEST_FILE':str(archivo),'IMPORT_TEST_SESSION':self.client.cookies['sessionid'].value}
            result=subprocess.run(['node',str(Path(__file__).parent/'browser_tests/importacion_clientes.cjs')],env=env,capture_output=True,text=True,timeout=120)
            self.assertEqual(result.returncode,0,result.stdout+result.stderr)
        self.assertEqual(RegistroCompraFiscal.objects.count(),3)
        self.assertEqual(RegistroCompraFiscal.objects.get(numero_factura='').total,Decimal('62.50'))
