from datetime import date
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from urllib.parse import parse_qs, urlparse

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import ConfiguracionPowerBIEmpresa, Empresa, EmpresaModulo, Modulo, RolSistema, Usuario
from .models import Cliente, Factura, LineaFactura, Producto, TipoImpuesto


class ReportesDemoTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.empresa = Empresa.objects.create(nombre="Empresa de prueba", slug="demo_1", rtn="demo-reportes")
        cls.other = Empresa.objects.create(nombre="Otra empresa", slug="tralingual", rtn="otra-reportes")
        module, _ = Modulo.objects.get_or_create(codigo="facturacion", defaults={"nombre": "Facturacion"})
        for empresa in (cls.empresa, cls.other):
            EmpresaModulo.objects.create(empresa=empresa, modulo=module, activo=True)
        cls.user = Usuario.objects.create_user(username="reports-admin", password="test-only", empresa=cls.empresa, es_administrador_empresa=True)
        cls.alpha = Cliente.objects.create(empresa=cls.empresa, nombre="Cliente Alfa", rtn="08011999000001")
        cls.beta = Cliente.objects.create(empresa=cls.empresa, nombre="Cliente Beta")
        cls.foreign = Cliente.objects.create(empresa=cls.other, nombre="Cliente privado externo")
        cls.product = Producto.objects.create(empresa=cls.empresa, nombre="Servicio de prueba", precio=100)
        cls.tax = TipoImpuesto.objects.create(nombre="ISV demo reportes", porcentaje=15)
        cls.a = cls.make_invoice(cls.alpha, "001-001-01-00000001", date(2026, 8, 1))
        cls.b = cls.make_invoice(cls.beta, "001-001-01-00000002", date(2026, 8, 31), currency="USD", rate="25")
        cls.void = cls.make_invoice(cls.alpha, "001-001-01-00000003", date(2026, 8, 15), state="anulada")
        cls.old = cls.make_invoice(cls.alpha, "001-001-01-00000004", date(2024, 1, 10))

    @classmethod
    def make_invoice(cls, client, number, issued, currency="HNL", rate="1", state="emitida"):
        invoice = Factura.objects.create(empresa=cls.empresa, cliente=client, fecha_emision=issued, moneda=currency, tipo_cambio=Decimal(rate))
        LineaFactura.objects.create(factura=invoice, producto=cls.product, cantidad=1, precio_unitario=100, impuesto=cls.tax)
        invoice.calcular_totales()
        invoice.save(update_fields=["subtotal", "impuesto", "total", "total_lempiras"])
        # Fixture documents bypass issuance; no fiscal numbering is consumed.
        Factura.objects.filter(pk=invoice.pk).update(estado=state, numero_factura=number)
        invoice.refresh_from_db()
        return invoice

    def setUp(self):
        self.client.force_login(self.user)
        self.url = reverse("reportes_facturacion", args=[self.empresa.slug])

    def test_resumen_y_totales_multimoneda_no_incluyen_anulada(self):
        response = self.client.get(self.url, {"fecha_desde": "2026-08-01", "fecha_hasta": "2026-08-31"})
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "facturacion/reportes_demo.html")
        self.assertEqual(response.context["totales"]["total"], Decimal("2990.00"))
        self.assertEqual(response.context["total_saldo"], Decimal("2990.00"))
        self.assertEqual(len(response.context["facturas"]), 3)
        self.assertContains(response, "dv-solutions-oficial.png")
        self.assertNotContains(response, self.foreign.nombre)

    def test_tendencia_incluye_periodos_historicos_seleccionados(self):
        response = self.client.get(self.url, {"fecha_desde": "2024-01-01", "fecha_hasta": "2024-01-31"})
        self.assertEqual(len(response.context["report_trend"]), 1)
        self.assertEqual(response.context["report_trend"][0]["periodo"], date(2024, 1, 1))
        self.assertEqual(response.context["report_trend"][0]["total"], Decimal("115.00"))

    def test_filtros_invalidos_no_provocan_500_ni_amplian_resultados(self):
        for params in ({"fecha_desde": "no-es-fecha"}, {"fecha_desde": "2026-09-01", "fecha_hasta": "2026-08-01"},
                       {"cliente": str(self.foreign.pk)}, {"cliente": "abc"}, {"estado_pago": "otro"}, {"impuesto": "nan"}):
            with self.subTest(params=params):
                response = self.client.get(self.url, params)
                self.assertEqual(response.status_code, 200)
                self.assertFalse(response.context["report_valid"])
                self.assertEqual(len(response.context["facturas"]), 0)
                self.assertContains(response, "Revisa los filtros")
                self.assertNotContains(response, "data-report-export")
                self.assertNotContains(response, 'class="report-kpis"')

    def test_busqueda_y_excel_comparten_seleccion(self):
        params = {"q": self.b.numero_factura, "fecha_desde": "2026-08-01", "impuesto": "15", "vista": "documentos"}
        report = self.client.get(self.url, params)
        self.assertEqual([f.pk for f in report.context["facturas"]], [self.b.pk])
        excel = self.client.get(reverse("exportar_excel", args=[self.empresa.slug]), params)
        self.assertEqual(excel.status_code, 200)
        workbook = load_workbook(BytesIO(excel.content))
        details = workbook["Detalle Facturas"]
        self.assertEqual(details.max_row, 2)
        self.assertEqual(details.cell(2, 3).value, self.b.numero_factura)
        self.assertEqual(details.cell(2, 18).value, 2875)

    @patch("facturacion.views._generar_factura_pdf_bytes", return_value=b"%PDF-test")
    def test_zip_respeta_busqueda_cliente_y_fechas(self, pdf):
        params = {"cliente": self.alpha.pk, "q": self.a.numero_factura, "fecha_desde": "2026-08-01", "fecha_hasta": "2026-08-31"}
        report = self.client.get(self.url, params)
        self.assertTrue(report.context["report_can_zip"])
        response = self.client.get(reverse("descargar_facturas_filtradas_zip", args=[self.empresa.slug]), params)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["X-Total-Facturas"], "1")
        self.assertEqual(pdf.call_args.kwargs["factura"].pk if pdf.call_args.kwargs else pdf.call_args.args[1].pk, self.a.pk)

    def test_vistas_fiscal_y_documentos_y_estado_anulado(self):
        for view in ("fiscal", "documentos"):
            response = self.client.get(self.url, {"vista": view})
            self.assertEqual(response.status_code, 200)
            self.assertContains(response, "Composición fiscal")
        self.assertContains(response, "Anulada")
        self.assertContains(response, f'invoice-detail-{self.b.pk}')
        self.assertContains(response, "Moneda original: USD")

    def test_paginacion_conserva_filtros_y_totales_completos(self):
        for i in range(27):
            self.make_invoice(self.alpha, f"001-001-01-{100 + i:08d}", date(2026, 8, 20))
        response = self.client.get(self.url, {"vista": "documentos", "q": "Cliente Alfa", "pagina": "2", "orden": "mayor_total"})
        self.assertEqual(response.context["report_page"].number, 2)
        self.assertEqual(len(response.context["facturas"]), 30)
        self.assertEqual(len(response.context["report_page"]), 5)
        self.assertEqual(response.context["totales"]["total"], Decimal("3335.00"))
        query = parse_qs(urlparse(response.context["report_previous"]).query)
        self.assertEqual(query["q"], ["Cliente Alfa"])
        self.assertEqual(query["orden"], ["mayor_total"])
        chip_query = parse_qs(urlparse(response.context["report_chips"][0]["url"]).query)
        self.assertNotIn("q", chip_query)
        self.assertEqual(chip_query["vista"], ["documentos"])
        self.assertEqual(chip_query["orden"], ["mayor_total"])

    def test_estado_vacio_y_pagina_invalida(self):
        empty = self.client.get(self.url, {"q": "sin-coincidencias"})
        self.assertContains(empty, "Sin documentos en esta selección")
        self.assertFalse(empty.context["report_can_zip"])
        response = self.client.get(self.url, {"pagina": "abc", "vista": "documentos", "orden": "invalido"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["report_page"].number, 1)
        self.assertEqual(response.context["report_order"], "recientes")

    def test_sin_permiso_de_exportar_no_hay_acciones_de_descarga(self):
        role = RolSistema.objects.create(nombre="Consulta reportes", codigo="consulta-reportes-demo", puede_reportes=True)
        self.user.es_administrador_empresa = False
        self.user.rol_sistema = role
        self.user.save()
        response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertNotContains(response, "data-export-menu")
        self.assertNotContains(response, 'title="Ver factura"')
        denied = self.client.get(reverse("exportar_excel", args=[self.empresa.slug]))
        self.assertEqual(denied.status_code, 302)

    def test_empresa_ajena_y_sesion_anonima_no_acceden(self):
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)
        outsider = Usuario.objects.create_user(username="report-outsider", empresa=self.other, password="test-only")
        self.client.force_login(outsider)
        self.assertEqual(self.client.get(self.url).status_code, 302)

    def test_tralingual_y_clinicas_conservan_reportes_originales(self):
        self.user.is_superuser = True
        self.user.save()
        for slug in ("tralingual", "hospital_mia", "medical_spa", "luque_aestetic", "serviciosmedicos"):
            company, _ = Empresa.objects.get_or_create(slug=slug, defaults={"nombre": slug, "rtn": slug})
            module = Modulo.objects.get(codigo="facturacion")
            EmpresaModulo.objects.get_or_create(empresa=company, modulo=module, defaults={"activo": True})
            response = self.client.get(reverse("reportes_facturacion", args=[slug]))
            self.assertTemplateUsed(response, "facturacion/reportes_premium.html")
            self.assertNotContains(response, "reportes-demo.css")

    def test_panel_externo_se_mantiene_cuando_esta_configurado(self):
        ConfiguracionPowerBIEmpresa.objects.create(empresa=self.empresa, activo=True, mostrar_en_reportes=True,
                                                 titulo_panel="Indicadores externos", url_embed="https://app.powerbi.com/reportEmbed?reportId=test")
        response = self.client.get(self.url, {"vista": "externo"})
        self.assertContains(response, "Indicadores externos")
        self.assertContains(response, "<iframe")
