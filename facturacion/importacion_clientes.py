"""Vista previa y confirmación de Excel sobre RegistroCompraFiscal existente."""
import calendar
import hashlib
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from django import forms
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core import signing
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction, IntegrityError, OperationalError
from django.db.models import Q
from django.http import Http404
from django.shortcuts import redirect, render
from django.views.decorators.http import require_http_methods
from openpyxl import load_workbook

from core.models import Empresa
from .captura_rapida import CapturaForm, buscar_duplicada, calcular_importes, numero_normalizado
from .clientes_contables import empresa_contable, cliente_autorizado
from .importadores import _detectar_encabezado, _indice, _normalizar_texto
from .models import LibroCompraMensual, Proveedor, RegistroCompraFiscal

SALT = 'compras.importacion.cliente.v1'
MAX_FILAS = 1000
CENTAVOS = Decimal('0.01')


class ArchivoForm(forms.Form):
    archivo = forms.FileField(label='Libro de compras (.xlsx)')

    def clean_archivo(self):
        archivo = self.cleaned_data['archivo']
        if not archivo.name.lower().endswith('.xlsx') or archivo.size > 5 * 1024 * 1024:
            raise forms.ValidationError('Selecciona un archivo .xlsx de hasta 5 MB.')
        return archivo


def importe(valor):
    if valor in (None, ''):
        return Decimal('0.00')
    try:
        resultado = Decimal(str(valor))
        if not resultado.is_finite() or resultado < 0 or resultado >= Decimal('1000000000000'):
            raise ValueError
        return resultado.quantize(CENTAVOS)
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError('Importe inválido o fuera del límite permitido.')


def leer_excel(contenido, nombre):
    """Lee valores cacheados; no ejecuta fórmulas ni modifica el archivo original."""
    try:
        with ZipFile(BytesIO(contenido)) as archivo:
            if sum(i.file_size for i in archivo.infolist()) > 50 * 1024 * 1024:
                raise ValueError('El archivo descomprimido supera el límite de 50 MB.')
        valores = load_workbook(BytesIO(contenido), data_only=True, read_only=True, keep_links=False)
        formulas = load_workbook(BytesIO(contenido), data_only=False, read_only=True, keep_links=False)
    except (BadZipFile, KeyError, OSError, TypeError) as exc:
        raise ValueError('No se pudo leer el Excel. Guarda una copia válida en formato .xlsx.') from exc
    try:
        hoja = valores['COMPRAS'] if 'COMPRAS' in valores.sheetnames else valores.worksheets[0]
        origen = formulas[hoja.title]
        if hoja.max_column > 40 or hoja.max_row > 5000:
            raise ValueError('El archivo excede 40 columnas o 5000 filas. Divide el libro antes de cargarlo.')
        encabezado, headers = _detectar_encabezado(hoja)
        columnas = {
            'fecha': _indice(headers, 'fecha'), 'proveedor': _indice(headers, 'beneficiario', 'proveedor'),
            'numero': _indice(headers, 'numero_factura', 'numero_de_factura', 'factura'),
            'rtn': _indice(headers, 'rtn'), 'subtotal': _indice(headers, 'subtotal', 'sub_total'),
            'isv_15': _indice(headers, 'isv_15%', 'isv_15', '15%'),
            'isv_18': _indice(headers, '018', '18%', 'isv_18'), 'total': headers.index('total')+1 if 'total' in headers else None,
            'exento': _indice(headers, 'exenta', 'exento'),
            'base_15': _indice(headers, 'base_15%', 'base_15'), 'base_18': _indice(headers, 'base_18%', 'base_18'),
        }
        if not all(columnas[k] for k in ('fecha', 'proveedor', 'subtotal')):
            raise ValueError('Se requieren las columnas Fecha, Beneficiario/Proveedor y Subtotal.')
        filas, total_excel, notas = [], None, []
        titulo = ' · '.join(str(c.value) for row in hoja.iter_rows(max_row=encabezado-1) for c in row if c.value is not None).strip()[:400]
        for numero, (row, formulas_row) in enumerate(zip(hoja.iter_rows(min_row=encabezado+1), origen.iter_rows(min_row=encabezado+1)), encabezado+1):
            datos = {k: row[col-1].value if col else None for k, col in columnas.items()}
            if not any(v not in (None, '', ' ') for v in datos.values()):
                continue
            if str(datos['fecha'] or '').strip().upper() in ('TOTAL', 'TOTALES', 'TOTAL GENERAL'):
                if datos['total'] is not None:
                    total_excel = str(importe(datos['total']))
                break
            fila = {'fila':numero, 'fecha':datos['fecha'].strftime('%d/%m/%Y') if isinstance(datos['fecha'], (date, datetime)) else _normalizar_texto(datos['fecha']),
                    'proveedor':_normalizar_texto(datos['proveedor']), 'numero':_normalizar_texto(datos['numero']),
                    'rtn':_normalizar_texto(datos['rtn']), 'errores_origen':[]}
            for k, col in columnas.items():
                if col and formulas_row[col-1].data_type == 'f' and datos[k] is None:
                    fila['errores_origen'].append(f'{k}: fórmula sin valor guardado. Abre, recalcula y guarda el Excel.')
            try:
                subtotal = importe(datos['subtotal'])
                isv15, isv18 = importe(datos['isv_15']), importe(datos['isv_18'])
                if columnas['base_15'] or columnas['base_18'] or columnas['exento']:
                    exento, base15, base18 = (importe(datos[k]) for k in ('exento','base_15','base_18'))
                    if exento + base15 + base18 != subtotal:
                        raise ValueError('Las bases y exentas no coinciden con el subtotal.')
                elif isv15 and isv18:
                    raise ValueError('Factura con ambos impuestos: incluye columnas Base 15%, Base 18% y Exenta para identificar los importes.')
                else:
                    exento = subtotal if not isv15 and not isv18 else Decimal(0)
                    base15 = subtotal if isv15 else Decimal(0)
                    base18 = subtotal if isv18 else Decimal(0)
                calculado = calcular_importes(exento, base15, base18)
                if calculado['total'] <= 0 or calculado['total'] >= Decimal('1000000000000'):
                    raise ValueError('El total debe ser positivo y estar dentro del límite permitido.')
                fila.update({k:f'{v:.2f}' for k,v in dict(exento=exento,base_15=base15,base_18=base18,**calculado).items()})
                fila['total_excel'] = str(importe(datos['total'])) if datos['total'] is not None else None
                fila['diferencia'] = str(calculado['total']-importe(datos['total'])) if datos['total'] is not None else None
                fila['impuestos_distintos'] = (isv15 != calculado['isv_15'] or isv18 != calculado['isv_18'])
            except ValueError as exc:
                fila['errores_origen'].append(str(exc))
            filas.append(fila)
            if len(filas) > MAX_FILAS:
                raise ValueError(f'Importa como máximo {MAX_FILAS} filas por archivo.')
        if not filas:
            raise ValueError('No se encontraron filas de compras.')
        if total_excel is None:
            notas.append('No se encontró un total general guardado en el Excel.')
        return {'archivo':nombre[:150], 'sha':hashlib.sha256(contenido).hexdigest(), 'hoja':hoja.title,
                'titulo':titulo, 'filas':filas, 'total_excel':total_excel, 'notas':notas}
    finally:
        valores.close()
        formulas.close()


def clave_fila(lote, fila):
    return hashlib.sha256(f"{lote['empresa']}:{lote['cliente']}:{lote['sha']}:{lote['hoja']}:{fila['fila']}".encode()).hexdigest()


def revisar(lote, empresa, cliente, usuario, post=None):
    proveedores = list(Proveedor.objects.filter(empresa=empresa, cliente_contable=cliente).order_by('pk'))
    existentes = RegistroCompraFiscal.objects.filter(empresa=empresa, cliente_contable=cliente)
    revisadas, vistos, vistos_sin_numero = [], set(), set()
    for original in lote['filas']:
        fila = {**original, 'errores':list(original['errores_origen']), 'avisos':[], 'duplicada':None, 'obj_proveedor':None}
        n = fila['fila']
        fila['seleccionada'] = post is None or str(n) in post.getlist('filas')
        for campo in ('fecha', 'proveedor', 'rtn', 'numero'):
            if post is not None:
                fila[campo] = post.get(f'{campo}_{n}', fila[campo]).strip()
        if not fila['proveedor'] or len(fila['proveedor']) > 200:
            fila['errores'].append('Completa el proveedor (máximo 200 caracteres).')
        rtn = numero_normalizado(fila['rtn']).replace(' ', '')
        if rtn and not re.fullmatch(r'[0-9]{1,20}', rtn):
            fila['errores'].append('RTN inválido.')
        fila['rtn'] = rtn
        # Nunca inventar prefijos para referencias cortas ni un número para documentos sin él.
        if len(fila['numero']) > 120:
            fila['errores'].append('Número de factura demasiado largo.')
        if re.fullmatch(r'[0-9]{9,}', fila['numero']):
            num = fila['numero']; fila['numero'] = f'{num[:3]}-{num[3:6]}-{num[6:8]}-{num[8:]}'
        try:
            form = CapturaForm(empresa=empresa, cliente_contable=cliente)
            form.cleaned_data = {'fecha_documento':fila['fecha']}
            fecha = form.clean_fecha_documento()
            fila['fecha_obj'] = fecha
            if fecha > date(lote['anio'],lote['mes'],calendar.monthrange(lote['anio'],lote['mes'])[1]):
                fila['avisos'].append('Fecha posterior al período del libro: verifica el año. Se conservará si confirmas.')
            elif (fecha.year, fecha.month) != (lote['anio'], lote['mes']):
                fila['avisos'].append('Fecha de otro período; permanecerá en el libro seleccionado.')
        except ValidationError as exc:
            fila['errores'].extend(exc.messages)
        nombre = fila['proveedor'].casefold()
        coincidencias = [p for p in proveedores if (rtn and numero_normalizado(p.rtn or '').replace(' ','') == rtn)
                         or (not rtn and p.nombre.strip().casefold() == nombre)]
        if len(coincidencias) > 1:
            fila['errores'].append('Hay varios proveedores coincidentes. Completa el RTN para identificarlo.')
        elif coincidencias:
            fila['obj_proveedor'] = coincidencias[0]
            if not fila['rtn']:
                fila['rtn'] = numero_normalizado(coincidencias[0].rtn or '').replace(' ', '')
            if not coincidencias[0].activo:
                fila['errores'].append('El proveedor está inactivo. Revisa su ficha antes de importar.')
        else:
            mismo_nombre = [p for p in proveedores if p.nombre.strip().casefold() == nombre]
            if rtn and mismo_nombre:
                fila['errores'].append('Existe ese nombre con otro RTN o sin RTN. Completa/corrige su ficha para vincularlo sin duplicarlo.')
            if not usuario.tiene_permiso_erp('puede_crear_proveedores', empresa):
                fila['errores'].append('No tienes permiso para crear este proveedor.')
            fila['avisos'].append('Se creará un proveedor dentro de este cliente.')
        proveedor = fila['obj_proveedor'] or Proveedor(empresa=empresa,cliente_contable=cliente,nombre=fila['proveedor'],rtn=rtn)
        if not proveedor.rtn:
            fila['avisos'].append('Proveedor sin RTN: podrás completarlo en su ficha.')
        clave = clave_fila(lote,fila)
        if existentes.filter(clave_importacion=clave).exists():
            fila['duplicada'] = 'Esta fila del archivo ya fue importada.'
        elif fila['numero'] and not fila['errores']:
            duplicada = buscar_duplicada(empresa,proveedor,fila['numero'],cliente_contable=cliente)
            if duplicada:
                fila['duplicada'] = f"{duplicada['fecha']} · {duplicada['proveedor']} · {duplicada['numero']} · {duplicada['total']} · {duplicada['estado']}"
        if not fila['numero']:
            fila['avisos'].append('Sin número / número ilegible: la detección de duplicados será aproximada.')
            if not fila['errores'] and existentes.filter(fecha_documento=fila['fecha_obj'],total=Decimal(fila['total'])).filter(
                Q(proveedor=proveedor) if proveedor.pk else Q(proveedor_nombre__iexact=proveedor.nombre)).exists():
                fila['avisos'].append('Posible duplicado: ya existe una compra del proveedor con la misma fecha y total.')
            if fila['seleccionada'] and not fila['errores']:
                posible = (proveedor.pk or rtn or nombre, fila['fecha_obj'], fila['total'])
                if posible in vistos_sin_numero:
                    fila['avisos'].append('Posible duplicado dentro del archivo: proveedor, fecha y total coinciden en otra fila sin número.')
                vistos_sin_numero.add(posible)
        identidad = (f'p:{proveedor.pk}' if proveedor.pk else (rtn or nombre),numero_normalizado(fila['numero']))
        if fila['seleccionada'] and fila['numero'] and not fila['errores'] and not fila['duplicada']:
            if identidad in vistos:
                fila['duplicada'] = 'Número repetido para el mismo proveedor dentro del archivo.'
            vistos.add(identidad)
        if fila.get('diferencia') not in (None,'0.00') or fila.get('impuestos_distintos'):
            fila['avisos'].append('Hay diferencias de cálculo/redondeo respecto al Excel. Revisa los importes.')
        revisadas.append(fila)
    return revisadas


def datos_revision(request, lote):
    post = request.POST.copy()
    if 'ediciones' in post:
        try:
            datos = json.loads(post['ediciones'])
            if not isinstance(datos, dict) or not isinstance(datos.get('seleccionadas'), list):
                raise ValueError
            post.setlist('filas', [str(n) for n in datos['seleccionadas']])
            for fila in lote['filas']:
                for campo in ('fecha','proveedor','numero','rtn'):
                    nombre = f"{campo}_{fila['fila']}"
                    valor = datos.get(nombre, fila[campo])
                    if not isinstance(valor,str) or len(valor) > 500:
                        raise ValueError
                    post[nombre] = valor
        except (ValueError, TypeError):
            raise ValueError('Los datos de revisión no son válidos. Recarga el archivo.')
    return post


@login_required
@require_http_methods(['GET','POST'])
def importar_cliente(request, empresa_slug, cliente_id, anio, mes):
    empresa = empresa_contable(request,empresa_slug)
    cliente = cliente_autorizado(request,empresa,cliente_id)
    if not (1 <= anio <= 9999 and 1 <= mes <= 12):
        raise Http404
    if not cliente.activo or not request.user.tiene_permiso_erp('puede_crear_compras',empresa):
        raise PermissionDenied
    form, lote, filas, error = ArchivoForm(), None, [], None
    token = None
    if request.method == 'POST':
        try:
            if request.POST.get('token'):
                lote = signing.loads(request.POST['token'],salt=SALT,max_age=3600)
                if (lote['empresa'],lote['cliente'],lote['usuario'],lote['anio'],lote['mes']) != (empresa.pk,cliente.pk,request.user.pk,anio,mes):
                    raise PermissionDenied
            else:
                form = ArchivoForm(request.POST,request.FILES)
                if form.is_valid():
                    archivo = form.cleaned_data['archivo']
                    lote = leer_excel(archivo.read(),archivo.name)
                    lote.update(empresa=empresa.pk,cliente=cliente.pk,usuario=request.user.pk,anio=anio,mes=mes)
            if lote:
                token = signing.dumps(lote,salt=SALT,compress=True)
                post = datos_revision(request,lote) if request.POST.get('token') else None
                filas = revisar(lote,empresa,cliente,request.user,post)
                if request.POST.get('accion') == 'confirmar':
                    if not request.POST.get('revisado'):
                        raise ValueError('Confirma que revisaste el cliente, período, fechas, proveedores e importes.')
                    with transaction.atomic():
                        Empresa.objects.select_for_update().get(pk=empresa.pk)
                        cliente = cliente_autorizado(request,empresa,cliente_id)
                        if not cliente.activo:
                            raise PermissionDenied
                        filas = revisar(lote,empresa,cliente,request.user,post)
                        elegidas = [f for f in filas if f['seleccionada'] and not f['duplicada']]
                        if any(f['errores'] for f in elegidas):
                            raise ValueError('Corrige o desmarca las filas con errores. No se guardó ninguna compra.')
                        if not elegidas:
                            raise ValueError('No hay filas nuevas seleccionadas para importar.')
                        libro, _ = LibroCompraMensual.objects.get_or_create(empresa=empresa,cliente_contable=cliente,anio=anio,mes=mes)
                        if libro.estado == 'finalizado':
                            raise ValueError('El libro está finalizado. Reábrelo para importar.')
                        creadas, nuevos = 0, {}
                        for fila in elegidas:
                            proveedor = fila['obj_proveedor']
                            if proveedor is None:
                                # Otra fila del lote puede haber creado este proveedor dentro de la transacción.
                                identidad_proveedor = ('rtn',fila['rtn']) if fila['rtn'] else ('nombre',fila['proveedor'].casefold())
                                qs = Proveedor.objects.filter(empresa=empresa,cliente_contable=cliente)
                                proveedor = nuevos.get(identidad_proveedor)
                                if proveedor is None:
                                    proveedor = qs.filter(rtn=fila['rtn']).first() if fila['rtn'] else qs.filter(nombre__iexact=fila['proveedor']).first()
                                if proveedor is None:
                                    proveedor = Proveedor.objects.create(empresa=empresa,cliente_contable=cliente,nombre=fila['proveedor'],rtn=fila['rtn'])
                                nuevos[identidad_proveedor] = proveedor
                            normalizado = numero_normalizado(fila['numero']) or None
                            if normalizado and buscar_duplicada(empresa,proveedor,fila['numero'],cliente_contable=cliente):
                                continue
                            rtn = numero_normalizado(proveedor.rtn or '').replace(' ','')
                            identidad = (f'rtn:{rtn}' if rtn else f'proveedor:{proveedor.pk}') if normalizado else None
                            observacion = f"Importación Excel · hoja {lote['hoja']} · fila {fila['fila']}. Total Excel: {fila.get('total_excel')}. Diferencia: {fila.get('diferencia')}."
                            if not normalizado:
                                observacion += ' Sin número / número ilegible.'
                            observacion += f" Fecha original: {next(f['fecha'] for f in lote['filas'] if f['fila']==fila['fila'])}."
                            RegistroCompraFiscal.objects.create(empresa=empresa,cliente_contable=cliente,proveedor=proveedor,
                                proveedor_nombre=proveedor.nombre,proveedor_rtn=proveedor.rtn,numero_factura=fila['numero'],
                                numero_factura_normalizado=normalizado,identidad_captura=identidad,clave_importacion=clave_fila(lote,fila),
                                fecha_documento=fila['fecha_obj'],periodo_anio=anio,periodo_mes=mes,creado_por=request.user,
                                origen_importacion=lote['archivo'],observacion=observacion,
                                **{k:Decimal(fila[k]) for k in ('exento','base_15','base_18','subtotal','isv_15','isv_18','total')})
                            creadas += 1
                        libro.actualizado_por=request.user
                        libro.save()
                    messages.success(request,f'Importación completada: {creadas} compras. Las duplicadas se omitieron.')
                    return redirect('captura_cliente_contable',empresa_slug=empresa_slug,cliente_id=cliente.pk,anio=anio,mes=mes)
        except signing.BadSignature:
            error = 'La vista previa venció o no es válida. Vuelve a cargar el archivo.'
        except (ValueError, ValidationError) as exc:
            error = '; '.join(exc.messages) if isinstance(exc,ValidationError) else str(exc)
        except (IntegrityError, OperationalError):
            error = 'Hubo otro guardado en curso. No se guardó este lote; revisa la vista previa y reintenta.'
    campos_totales = ('exento', 'base_15', 'base_18', 'isv_15', 'isv_18', 'total')
    validas = [f for f in filas if f.get('total') and f['seleccionada'] and not f['errores'] and not f['duplicada']]
    totales = {campo: sum((Decimal(f[campo]) for f in validas), Decimal('0.00')) for campo in campos_totales}
    total = totales['total']
    return render(request,'facturacion/importar_cliente.html',dict(empresa=empresa,cliente=cliente,anio=anio,mes=mes,
        form=form,lote=lote,filas=filas,token=token,error=error,total=total,totales=totales,
        diferencia=total-Decimal(lote['total_excel']) if lote and lote['total_excel'] is not None else None))
