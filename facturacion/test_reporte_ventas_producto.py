"""Ventas de un producto: importes de sus líneas y aislamiento por empresa."""

from datetime import date
from decimal import Decimal
from io import BytesIO

from django.core.exceptions import PermissionDenied
from django.db import connection
from django.http import Http404
from django.test import RequestFactory, TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from openpyxl import load_workbook
from pypdf import PdfReader

from core.models import Empresa, EmpresaModulo, Modulo, RolSistema, Usuario
from .models import Cliente, Factura, LineaFactura, Producto, TipoImpuesto
from .reporte_ventas_producto import buscar_productos_ventas, reporte_ventas_producto
from .services_clientes_compartidos import suspender_sincronizacion_clientes_compartidos


@override_settings(EMAIL_BACKEND="django.core.mail.backends.locmem.EmailBackend")
class ReporteVentasProductoTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.empresas = []
        modulo, _ = Modulo.objects.get_or_create(codigo="facturacion", defaults={"nombre": "Facturación"})
        for indice, slug in enumerate(("hospital_mia", "medical_spa", "luque_aestetic", "serviciosmedicos"), 1):
            empresa = Empresa.objects.create(nombre=f"Empresa producto {indice}", slug=slug,
                rtn=f"RVP-{indice}", tipo_solucion="clinica", estado_licencia="activa")
            EmpresaModulo.objects.create(empresa=empresa, modulo=modulo)
            cls.empresas.append(empresa)
        cls.empresa, cls.otra = cls.empresas[:2]
        cls.fuera = Empresa.objects.create(nombre="Empresa ajena", slug="empresa_ajena_rvp", rtn="RVP-AJENA")
        cls.rol = RolSistema.objects.create(nombre="Reportes de producto", codigo="rvp-reportes",
            puede_reportes=True, puede_exportar_reportes=True, puede_ver_facturas=True)
        cls.usuario = Usuario.objects.create_user(username="rvp-lector", empresa=cls.empresa, rol_sistema=cls.rol)
        cls.vendedor = Usuario.objects.create_user(username="rvp-vendedor", empresa=cls.empresa,
            first_name="Vendedor", last_name="Histórico")
        cls.vendedor_ajeno = Usuario.objects.create_user(username="rvp-vendedor-ajeno", empresa=cls.otra)
        cls.impuesto = TipoImpuesto.objects.create(nombre="ISV producto 15", porcentaje=15)
        cls.exento = TipoImpuesto.objects.create(nombre="Exento producto", porcentaje=0)
        cls.producto = Producto.objects.create(empresa=cls.empresa, nombre="Suero de prueba",
            codigo="RVP-SUERO", descripcion="Referencia botella azul", precio=100,
            controla_inventario=False, impuesto_predeterminado=cls.impuesto)
        cls.no_seleccionado = Producto.objects.create(empresa=cls.empresa, nombre="Producto que no se suma",
            codigo="RVP-OTRO", precio=999, controla_inventario=False)
        cls.producto_ajeno = Producto.objects.create(empresa=cls.otra, nombre=cls.producto.nombre,
            codigo=cls.producto.codigo, precio=100, controla_inventario=False)
        cls.historico = Producto.objects.create(empresa=cls.empresa, nombre="Suero antiguo",
            codigo="RVP-ANTIGUO", descripcion="Referencia histórica", precio=15,
            activo=False, eliminado=True, controla_inventario=False)
        with suspender_sincronizacion_clientes_compartidos():
            cls.cliente = Cliente.objects.create(empresa=cls.empresa, nombre="Cliente producto", rtn="0801199000501")
            cls.cliente_ajeno = Cliente.objects.create(empresa=cls.otra, nombre="Cliente ajeno", rtn="0801199000502")

    def setUp(self):
        self.client.force_login(self.usuario)
        self.rf = RequestFactory()
        self.url = reverse("reporte_ventas_producto", args=[self.empresa.slug])
        self.search_url = reverse("buscar_productos_ventas", args=[self.empresa.slug])

    def parametros(self, **cambios):
        return {"producto": str(self.producto.pk), "fecha_desde": "2026-02-01",
            "fecha_hasta": "2026-02-28", "estado": "emitida", **cambios}

    def factura(self, *, fecha=date(2026, 2, 15), estado="emitida", empresa=None,
                vendedor=None, moneda="HNL", tipo_cambio="1", cliente=None):
        empresa = empresa or self.empresa
        factura = Factura.objects.create(empresa=empresa,
            cliente=cliente or (self.cliente if empresa == self.empresa else self.cliente_ajeno),
            vendedor=vendedor or self.vendedor, fecha_emision=fecha,
            moneda=moneda, tipo_cambio=Decimal(tipo_cambio), estado="borrador")
        Factura.objects.filter(pk=factura.pk).update(estado=estado)
        factura.estado = estado
        return factura

    def linea(self, factura, *, producto=None, cantidad="1", precio="10", descuento="0", impuesto=None):
        return LineaFactura.objects.create(factura=factura, producto=producto or self.producto,
            cantidad=Decimal(cantidad), precio_unitario=Decimal(precio),
            descuento_porcentaje=Decimal(descuento), impuesto=impuesto or self.exento)

    def test_suma_solo_lineas_del_producto_fechas_inclusivas_descuentos_y_dolares(self):
        inicial = self.factura(fecha=date(2026, 2, 1))
        self.linea(inicial, cantidad="2", precio="100", descuento="10", impuesto=self.impuesto)
        self.linea(inicial, precio="50")
        self.linea(inicial, producto=self.no_seleccionado, cantidad="9", precio="999")
        final = self.factura(fecha=date(2026, 2, 28), moneda="USD", tipo_cambio="25")
        self.linea(final, cantidad="1.5", precio="10", impuesto=self.impuesto)
        for fecha in (date(2026, 1, 31), date(2026, 3, 1)):
            self.linea(self.factura(fecha=fecha), cantidad="90", precio="999")
        self.linea(self.factura(empresa=self.otra), producto=self.producto_ajeno,
            cantidad="99", precio="100")

        response = self.client.get(self.url, self.parametros())

        self.assertEqual(response.status_code, 200)
        resumen = response.context["resumen"]
        self.assertEqual(resumen["cantidad"], Decimal("4.50"))
        self.assertEqual(resumen["total"], Decimal("688.25"))
        self.assertEqual(resumen["documentos"], 2)
        self.assertEqual(resumen["precio_promedio"].quantize(Decimal("0.01")), Decimal("152.94"))
        filas = list(response.context["pagina_obj"])
        self.assertEqual(len(filas), 3)
        self.assertEqual(sum((fila["total"] for fila in filas), Decimal("0")), resumen["total"])
        self.assertTrue(all(fila["factura_url"] for fila in filas))

    def test_anuladas_se_muestran_sin_sumar_y_borradores_solo_al_solicitarlos(self):
        self.linea(self.factura(), cantidad="2", precio="10")
        self.linea(self.factura(estado="borrador"), cantidad="3", precio="10")
        self.linea(self.factura(estado="anulada"), cantidad="90", precio="10")
        response = self.client.get(self.url, self.parametros())
        self.assertEqual(response.context["resumen"]["total"], Decimal("20.00"))
        self.assertEqual(len(response.context["pagina_obj"]), 1)
        response = self.client.get(self.url, self.parametros(estado=""))
        self.assertEqual(response.context["resumen"]["total"], Decimal("50.00"))
        self.assertEqual(response.context["resumen"]["documentos"], 2)
        self.assertEqual(sum(fila["anulada"] for fila in response.context["pagina_obj"]), 1)
        response = self.client.get(self.url, self.parametros(estado="anulada"))
        self.assertEqual(response.context["resumen"]["cantidad"], Decimal("0"))
        self.assertEqual(response.context["resumen"]["total"], Decimal("0"))
        self.assertEqual(response.context["resumen"]["documentos"], 0)
        self.assertTrue(all(fila["anulada"] for fila in response.context["pagina_obj"]))
        response = self.client.get(self.url, self.parametros(estado="borrador"))
        self.assertEqual(response.context["resumen"]["total"], Decimal("30.00"))

    def test_filtra_vendedor_y_rechaza_vendedores_ajenos_a_las_facturas_locales(self):
        self.linea(self.factura(), cantidad="2", precio="10")
        self.linea(self.factura(vendedor=self.usuario), cantidad="3", precio="10")
        self.linea(self.factura(empresa=self.otra, vendedor=self.vendedor_ajeno),
            producto=self.producto_ajeno)
        response = self.client.get(self.url, self.parametros(vendedor=str(self.vendedor.pk)))
        self.assertEqual(response.context["resumen"]["total"], Decimal("20.00"))
        response = self.client.get(self.url, self.parametros(vendedor=str(self.vendedor_ajeno.pk)))
        self.assertTrue(response.context["form"].errors)
        self.assertFalse(response.context["reporte_generado"])

    def test_producto_de_otra_empresa_no_consulta_ni_revela_ventas(self):
        self.linea(self.factura(empresa=self.otra), producto=self.producto_ajeno, precio="999")
        response = self.client.get(self.url, self.parametros(producto=str(self.producto_ajeno.pk)))
        self.assertEqual(response.status_code, 200)
        self.assertIn("producto", response.context["form"].errors)
        self.assertFalse(response.context["reporte_generado"])

    def test_errores_de_fecha_estado_e_identificadores_no_ejecutan_consulta_de_ventas(self):
        for cambios in ({"fecha_desde": "2026-02-30"}, {"fecha_hasta": "2026-01-01"},
                {"fecha_desde": ""}, {"producto": "no-numerico"}, {"estado": "inventado"}):
            with self.subTest(cambios=cambios), CaptureQueriesContext(connection) as consultas:
                response = self.client.get(self.url, self.parametros(**cambios))
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.context["form"].errors)
            self.assertFalse(response.context["reporte_generado"])
            self.assertFalse(any('FROM "facturacion_lineafactura"' in consulta["sql"]
                for consulta in consultas.captured_queries))

    def test_visita_inicial_no_suma_todo_el_catalogo_y_no_escribe(self):
        self.linea(self.factura(), precio="999")
        with CaptureQueriesContext(connection) as consultas:
            response = self.client.get(self.url)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.context["reporte_generado"])
        self.assertEqual(response.context["form"]["estado"].value(), "emitida")
        mutaciones = [consulta["sql"] for consulta in consultas.captured_queries
            if consulta["sql"].lstrip().upper().startswith(("INSERT", "UPDATE", "DELETE"))]
        self.assertEqual(mutaciones, [])

    def test_busqueda_inicial_y_por_nombre_codigo_o_referencia_incluye_historicos(self):
        response = self.client.get(self.search_url)
        self.assertEqual(response.status_code, 200)
        ids = {item["id"] for item in response.json()["results"]}
        self.assertIn(self.historico.pk, ids)
        self.assertNotIn(self.producto_ajeno.pk, ids)
        for busqueda in ("suero de", "RVP-SUERO", "botella azul"):
            with self.subTest(busqueda=busqueda):
                response = self.client.get(self.search_url, {"q": busqueda})
                self.assertEqual([item["id"] for item in response.json()["results"]], [self.producto.pk])
                self.assertTrue({"id", "text", "nombre", "codigo", "referencia", "activo"}
                    <= set(response.json()["results"][0]))

    def test_busqueda_pagina_todo_el_catalogo_sin_repeticiones_ni_fugas(self):
        Producto.objects.bulk_create([Producto(empresa=self.empresa, nombre=f"Paginar producto {i:03d}",
            codigo=f"RVP-PAG-{i:03d}", precio=1, controla_inventario=False) for i in range(65)])
        ids = []
        pagina = 1
        while pagina is not None:
            response = self.client.get(self.search_url, {"q": "Paginar producto", "pagina": pagina})
            datos = response.json()
            self.assertLessEqual(len(datos["results"]), 30)
            ids.extend(item["id"] for item in datos["results"])
            self.assertEqual(datos["has_more"], datos["next_page"] is not None)
            pagina = datos["next_page"]
        self.assertEqual(len(ids), 65)
        self.assertEqual(len(set(ids)), 65)

    def test_consulta_de_una_empresa_no_exige_acceso_a_las_otras_tres(self):
        self.assertFalse(self.usuario.puede_acceder_empresa(self.otra))
        response = self.client.get(self.url, self.parametros())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["resumen"]["total"], Decimal("0"))

    def test_autorizacion_de_vista_y_busqueda_no_depende_solo_del_middleware(self):
        sin_rol = Usuario.objects.create_user(username="rvp-sin-rol", empresa=self.empresa)
        admin_ajeno = Usuario.objects.create_user(username="rvp-admin-ajeno", empresa=self.otra,
            es_administrador_empresa=True)
        for vista in (reporte_ventas_producto, buscar_productos_ventas):
            for usuario in (sin_rol, admin_ajeno):
                request = self.rf.get(self.url, self.parametros())
                request.user = usuario
                with self.subTest(vista=vista.__name__, usuario=usuario.username), self.assertRaises(PermissionDenied):
                    vista(request, empresa_slug=self.empresa.slug)
            request = self.rf.get(self.url)
            request.user = self.usuario
            with self.assertRaises(Http404):
                vista(request, empresa_slug=self.fuera.slug)

    def test_lectura_no_concede_exportacion_ni_enlace_a_factura(self):
        self.linea(self.factura(), precio="10")
        self.rol.puede_exportar_reportes = False
        self.rol.puede_ver_facturas = False
        self.rol.save(update_fields=["puede_exportar_reportes", "puede_ver_facturas"])
        response = self.client.get(self.url, self.parametros())
        self.assertEqual(response.status_code, 200)
        self.assertTrue(all(not fila["factura_url"] for fila in response.context["pagina_obj"]))
        for formato in ("excel", "pdf"):
            request = self.rf.get(self.url, self.parametros(formato=formato))
            request.user = Usuario.objects.get(pk=self.usuario.pk)
            with self.subTest(formato=formato), self.assertRaises(PermissionDenied):
                reporte_ventas_producto(request, empresa_slug=self.empresa.slug)

    def test_excel_respeta_filtros_y_preserva_nombres_como_texto(self):
        Cliente.objects.filter(pk=self.cliente.pk).update(nombre="=1+1")
        self.linea(self.factura(), cantidad="2", precio="12")
        self.linea(self.factura(fecha=date(2026, 3, 1)), cantidad="90", precio="999")
        self.linea(self.factura(), producto=self.no_seleccionado, cantidad="80", precio="999")
        response = self.client.get(self.url, self.parametros(formato="excel"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        libro = load_workbook(BytesIO(response.content))
        celdas = [celda for hoja in libro for fila in hoja for celda in fila]
        self.assertTrue(any(str(celda.value).lstrip("'") == "=1+1" for celda in celdas))
        self.assertFalse(any(celda.data_type == "f" for celda in celdas))
        self.assertFalse(any(celda.value == self.no_seleccionado.nombre for celda in celdas))
        self.assertFalse(any(celda.value in (Decimal("79920"), Decimal("89910")) for celda in celdas))
        self.assertTrue(any(celda.value == 24 for celda in celdas))

    def test_resumen_y_excel_incluyen_ventas_fuera_de_la_pagina_actual(self):
        factura = self.factura()
        LineaFactura.objects.bulk_create([LineaFactura(factura=factura, producto=self.producto,
            cantidad=1, precio_unitario=1, subtotal=1, impuesto=self.exento, impuesto_monto=0)
            for _ in range(65)])
        response = self.client.get(self.url, self.parametros(pagina="2"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["resumen"]["cantidad"], Decimal("65"))
        self.assertEqual(response.context["resumen"]["total"], Decimal("65"))
        self.assertEqual(response.context["resumen"]["documentos"], 1)
        self.assertLess(len(response.context["pagina_obj"]), 65)
        self.assertEqual(response.context["pagina_obj"].paginator.count, 65)
        response = self.client.get(self.url, self.parametros(pagina="2", formato="excel"))
        libro = load_workbook(BytesIO(response.content))
        self.assertEqual(libro["Detalle"].max_row, 66)

    def test_vistas_directas_respetan_modulo_y_licencia(self):
        for cambio in ("modulo", "licencia"):
            with self.subTest(cambio=cambio):
                if cambio == "modulo":
                    EmpresaModulo.objects.filter(empresa=self.empresa).update(activo=False)
                else:
                    EmpresaModulo.objects.filter(empresa=self.empresa).update(activo=True)
                    Empresa.objects.filter(pk=self.empresa.pk).update(estado_licencia="suspendida")
                for vista in (reporte_ventas_producto, buscar_productos_ventas):
                    request = self.rf.get(self.url, self.parametros())
                    request.user = self.usuario
                    with self.assertRaises(PermissionDenied):
                        vista(request, empresa_slug=self.empresa.slug)

    def test_exportar_sin_filtros_devuelve_error_de_validacion(self):
        for formato in ("excel", "pdf"):
            with self.subTest(formato=formato):
                response = self.client.get(self.url, {"formato": formato})
                self.assertEqual(response.status_code, 400)
                self.assertFalse(response.context["reporte_generado"])
                self.assertTrue(response.context["form"].errors)

    def test_pdf_real_respeta_filtros_y_marca_anuladas_sin_sumarlas(self):
        with suspender_sincronizacion_clientes_compartidos():
            clientes = [Cliente.objects.create(empresa=self.empresa, nombre=nombre,
                rtn=f"08011990006{indice:02d}") for indice, nombre in enumerate((
                    "PDFClienteElegido", "PDFClienteAnulado", "PDFClienteFueraFecha",
                    "PDFClienteOtroVendedor"), 1)]
        elegida = self.factura(cliente=clientes[0], fecha=date(2026, 2, 1))
        self.linea(elegida, cantidad="2", precio="10")
        self.linea(elegida, producto=self.no_seleccionado, cantidad="3", precio="999")
        self.linea(self.factura(cliente=clientes[1], estado="anulada", fecha=date(2026, 2, 28)),
            cantidad="7", precio="111")
        self.linea(self.factura(cliente=clientes[2], fecha=date(2026, 3, 1)), precio="995")
        self.linea(self.factura(cliente=clientes[3], vendedor=self.usuario), precio="994")
        self.linea(self.factura(empresa=self.otra), producto=self.producto_ajeno, precio="993")

        response = self.client.get(self.url, self.parametros(
            estado="", vendedor=str(self.vendedor.pk), formato="pdf"))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))
        lector = PdfReader(BytesIO(response.content))
        texto = " ".join(" ".join(pagina.extract_text().split()) for pagina in lector.pages)
        self.assertIn("PDFClienteElegido", texto)
        self.assertIn("PDFClienteAnulado", texto)
        self.assertIn("ANULADA", texto)
        self.assertNotIn("PDFClienteFueraFecha", texto)
        self.assertNotIn("PDFClienteOtroVendedor", texto)
        self.assertNotIn(self.cliente_ajeno.nombre, texto)
        self.assertNotIn("2,997.00", texto)
        self.assertRegex(texto, r"Total vendido\s+L\.\s*20[.,]00")
        self.assertRegex(texto, r"Facturas\s+1\s*Excluye anuladas")

    def test_endpoint_solo_lectura(self):
        for url in (self.url, self.search_url):
            response = self.client.post(url, self.parametros())
            self.assertEqual(response.status_code, 405)
